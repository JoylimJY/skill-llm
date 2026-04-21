from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

# Create a simple presentation with multiple slides
prs = Presentation()

# Slide 1: Title slide
slide1 = prs.slides.add_slide(prs.slide_layouts[0])
title = slide1.shapes.title
subtitle = slide1.placeholders[1]
title.text = 'Sample Presentation'
subtitle.text = 'Testing Brand Application'

# Slide 2: Content slide
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
title2 = slide2.shapes.title
content2 = slide2.placeholders[1]
title2.text = 'Main Content'
content2.text = 'This is body text that should use Lora font.\nThis presentation needs Anthropic branding applied.'

# Slide 3: Bullet points
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
title3 = slide3.shapes.title
content3 = slide3.placeholders[1]
title3.text = 'Key Points'
content3.text = 'First important point\nSecond key message\nThird critical item'

prs.save('presentation.pptx')
print('Generated presentation.pptx with 3 slides')