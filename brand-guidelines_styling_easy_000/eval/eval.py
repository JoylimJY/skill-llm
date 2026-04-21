import sys
import os
import json
from pptx import Presentation

def check_presentation(workspace_dir):
    checks = []
    output_file = os.path.join(workspace_dir, 'branded_presentation.pptx')
    
    if not os.path.exists(output_file):
        return [{"name": "Output file exists", "passed": False, "detail": "branded_presentation.pptx not found"}]
    
    checks.append({"name": "Output file exists", "passed": True, "detail": "branded_presentation.pptx found"})
    
    try:
        prs = Presentation(output_file)
        
        # 1. Slide Count
        if len(prs.slides) < 3:
            checks.append({"name": "Slide count preserved", "passed": False, "detail": f"Expected 3, found {len(prs.slides)}"})
        else:
            checks.append({"name": "Slide count preserved", "passed": True, "detail": "All 3 slides present"})
        
        # Anthropic Colors (Tolerance within 5)
        brand_colors = [(20, 20, 19), (250, 249, 245), (217, 119, 87), (106, 155, 204), (120, 140, 93)]
        
        font_applied = False
        color_applied = False
        
        for slide in prs.slides:
            # 检查背景色 (修复属性访问)
            try:
                bg_rgb = slide.background.fill.fore_color.rgb
                # 使用索引访问 [0],[1],[2]
                if any(all(abs(bg_rgb[i] - c[i]) <= 5 for i in range(3)) for c in brand_colors):
                    color_applied = True
            except: pass

            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            # 检查字体
                            fname = (run.font.name or "").lower()
                            if any(f in fname for f in ['poppins', 'lora', 'arial', 'georgia']):
                                font_applied = True
                            # 检查文字颜色 (修复属性访问)
                            try:
                                c_rgb = run.font.color.rgb
                                if any(all(abs(c_rgb[i] - c[i]) <= 5 for i in range(3)) for c in brand_colors):
                                    color_applied = True
                            except: pass

        checks.append({"name": "Typography applied", "passed": font_applied, "detail": "Brand fonts detected"})
        checks.append({"name": "Brand colors applied", "passed": color_applied, "detail": "Anthropic colors detected"})
        
        # 2. Content Preservation
        content_found = any('sample presentation' in s.text.lower() for slide in prs.slides for s in slide.shapes if hasattr(s, 'text'))
        checks.append({"name": "Content preserved", "passed": content_found, "detail": "Original content maintained"})
        
    except Exception as e:
        checks.append({"name": "File processing", "passed": False, "detail": str(e)})
    
    return checks

if __name__ == '__main__':
    workspace = sys.argv[1]
    all_checks = check_presentation(workspace)
    score = sum(1 for c in all_checks if c['passed']) / len(all_checks)
    
    # 使用真正的 json.dumps ！！！
    print(json.dumps({
        "passed": score >= 0.75,
        "score": score,
        "checks": all_checks
    }))