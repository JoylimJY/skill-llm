import json
import os
import sys
from pathlib import Path
import subprocess
from openpyxl import load_workbook

def main():
    workspace = Path(sys.argv[1])
    checks = []
    
    # Check 1: Financial model file exists
    model_file = workspace / 'financial_model.xlsx'
    if model_file.exists():
        checks.append({"name": "financial_model_xlsx_exists", "passed": True, "detail": "Found financial_model.xlsx"})
    else:
        checks.append({"name": "financial_model_xlsx_exists", "passed": False, "detail": "financial_model.xlsx not found"})
        score = 0.0
        result = {"passed": False, "score": score, "checks": checks}
        print(json.dumps(result))
        return
    
    try:
        wb = load_workbook(str(model_file), data_only=False)
        sheet_names = [name.lower() for name in wb.sheetnames]
        
        # Check 2: Required sheets exist
        required_sheets = ['assumptions', 'income statement', 'balance sheet', 'cash flow', 'summary']
        sheets_found = []
        for req_sheet in required_sheets:
            found = any(req_sheet in sheet_name for sheet_name in sheet_names)
            sheets_found.append(found)
            checks.append({
                "name": f"sheet_{req_sheet.replace(' ', '_')}_exists",
                "passed": found,
                "detail": f"{'Found' if found else 'Missing'} {req_sheet} sheet"
            })
        
        # Check 3: Formulas present (not hardcoded values)
        formula_count = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                        formula_count += 1
        
        has_formulas = formula_count >= 20
        checks.append({
            "name": "contains_formulas",
            "passed": has_formulas,
            "detail": f"Found {formula_count} formulas (expected 20+)"
        })
        
        # Check 4: Financial years 2024-2028 present
        years_found = 0
        target_years = ['2024', '2025', '2026', '2027', '2028']
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and str(cell.value).strip() in target_years:
                        years_found += 1
                        break
        
        has_years = years_found >= 3
        checks.append({
            "name": "contains_projection_years",
            "passed": has_years,
            "detail": f"Found {years_found} projection year references (expected 3+)"
        })
        
        # Check 5: Currency formatting applied
        currency_formatted = False
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.number_format and ('$' in cell.number_format or '#,##0' in cell.number_format):
                        currency_formatted = True
                        break
                if currency_formatted:
                    break
            if currency_formatted:
                break
        
        checks.append({
            "name": "currency_formatting_applied",
            "passed": currency_formatted,
            "detail": "Currency formatting found" if currency_formatted else "No currency formatting detected"
        })
        
        # Check 6: Professional color coding (blue for inputs)
        has_blue_formatting = False
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.font and cell.font.color and hasattr(cell.font.color, 'rgb'):
                        if cell.font.color.rgb and '0000ff' in str(cell.font.color.rgb).lower():
                            has_blue_formatting = True
                            break
                if has_blue_formatting:
                    break
            if has_blue_formatting:
                break
        
        checks.append({
            "name": "blue_text_formatting",
            "passed": has_blue_formatting,
            "detail": "Blue text formatting found" if has_blue_formatting else "No blue text formatting detected"
        })
        
        # Check 7: Financial statement keywords present
        keywords_found = 0
        financial_keywords = ['revenue', 'expenses', 'assets', 'liabilities', 'equity', 'cash flow', 'net income']
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        cell_text = cell.value.lower()
                        for keyword in financial_keywords:
                            if keyword in cell_text:
                                keywords_found += 1
                                break
        
        has_financial_content = keywords_found >= 5
        checks.append({
            "name": "financial_statement_content",
            "passed": has_financial_content,
            "detail": f"Found {keywords_found} financial statement keywords (expected 5+)"
        })
        
        wb.close()
        
    except Exception as e:
        checks.append({
            "name": "file_processing_error",
            "passed": False,
            "detail": f"Error processing Excel file: {str(e)}"
        })
    
    # Calculate score
    passed_checks = sum(1 for check in checks if check['passed'])
    total_checks = len(checks)
    score = passed_checks / total_checks if total_checks > 0 else 0.0
    overall_passed = score >= 0.8
    
    result = {
        "passed": overall_passed,
        "score": score,
        "checks": checks
    }
    
    print(json.dumps(result))

if __name__ == '__main__':
    main()