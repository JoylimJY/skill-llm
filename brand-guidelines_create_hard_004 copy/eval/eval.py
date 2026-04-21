import sys
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt
import json

def check_presentation(workspace_dir):
    checks = []
    
    # Check if presentation file exists
    pptx_path = os.path.join(workspace_dir, 'anthropic_brand_demo.pptx')
    if not os.path.exists(pptx_path):
        checks.append({"name": "file_exists", "passed": False, "detail": "anthropic_brand_demo.pptx not found"})
        return checks
    
    checks.append({"name": "file_exists", "passed": True, "detail": "Presentation file found"})
    
    try:
        prs = Presentation(pptx_path)
    except Exception as e:
        checks.append({"name": "file_readable", "passed": False, "detail": f"Cannot read presentation: {str(e)}"})
        return checks
    
    checks.append({"name": "file_readable", "passed": True, "detail": "Presentation is readable"})
    
    # Check slide count
    slide_count = len(prs.slides)
    checks.append({
        "name": "slide_count", 
        "passed": slide_count == 4, 
        "detail": f"Found {slide_count} slides (expected 4)"
    })
    
    if slide_count < 4:
        return checks
    
    # Check title slide content
    title_slide = prs.slides[0]
    title_found = False
    subtitle_found = False
    
    for shape in title_slide.shapes:
        if hasattr(shape, 'text'):
            text_lower = shape.text.lower()
            if 'anthropic brand guidelines' in text_lower:
                title_found = True
            if 'visual identity showcase' in text_lower:
                subtitle_found = True
    
    checks.append({
        "name": "title_slide_title", 
        "passed": title_found, 
        "detail": "Title 'Anthropic Brand Guidelines' found" if title_found else "Title not found"
    })
    
    checks.append({
        "name": "title_slide_subtitle", 
        "passed": subtitle_found, 
        "detail": "Subtitle 'Visual Identity Showcase' found" if subtitle_found else "Subtitle not found"
    })
    
    # Check for brand colors usage
    brand_colors = {
        'dark': (20, 20, 19),
        'light': (250, 249, 245),
        'orange': (217, 119, 87),
        'blue': (106, 155, 204),
        'green': (120, 140, 93),
        'mid_gray': (176, 174, 165),
        'light_gray': (232, 230, 220)
    }
    
    colors_used = set()
    shapes_with_colors = 0
    font_usage = {'poppins': 0, 'lora': 0, 'arial': 0, 'georgia': 0}
    
    for slide in prs.slides:
        for shape in slide.shapes:
            # Check fill colors
            if hasattr(shape, 'fill') and shape.fill.type == 1:  # SOLID fill
                try:
                    rgb = shape.fill.fore_color.rgb
                    color_tuple = (rgb.red, rgb.green, rgb.blue)
                    for color_name, brand_rgb in brand_colors.items():
                        if abs(color_tuple[0] - brand_rgb[0]) <= 5 and \
                           abs(color_tuple[1] - brand_rgb[1]) <= 5 and \
                           abs(color_tuple[2] - brand_rgb[2]) <= 5:
                            colors_used.add(color_name)
                            shapes_with_colors += 1
                            break
                except:
                    pass
            
            # Check text formatting
            if hasattr(shape, 'text_frame'):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if hasattr(run, 'font') and run.font.name:
                            font_name = run.font.name.lower()
                            if 'poppins' in font_name:
                                font_usage['poppins'] += 1
                            elif 'lora' in font_name:
                                font_usage['lora'] += 1
                            elif 'arial' in font_name:
                                font_usage['arial'] += 1
                            elif 'georgia' in font_name:
                                font_usage['georgia'] += 1
    
    # Check brand colors usage
    accent_colors_used = len([c for c in colors_used if c in ['orange', 'blue', 'green']])
    checks.append({
        "name": "brand_colors_used", 
        "passed": len(colors_used) >= 3, 
        "detail": f"Found {len(colors_used)} brand colors used (need at least 3)"
    })
    
    checks.append({
        "name": "accent_colors_used", 
        "passed": accent_colors_used >= 1, 
        "detail": f"Found {accent_colors_used} accent colors used (need at least 1)"
    })
    
    # Check font usage
    brand_fonts_used = font_usage['poppins'] + font_usage['lora'] > 0
    fallback_fonts_used = font_usage['arial'] + font_usage['georgia'] > 0
    
    checks.append({
        "name": "typography_applied", 
        "passed": brand_fonts_used or fallback_fonts_used, 
        "detail": f"Brand/fallback fonts used: Poppins({font_usage['poppins']}), Lora({font_usage['lora']}), Arial({font_usage['arial']}), Georgia({font_usage['georgia']})"
    })
    
    # Check for colored shapes
    checks.append({
        "name": "colored_shapes", 
        "passed": shapes_with_colors >= 3, 
        "detail": f"Found {shapes_with_colors} shapes with brand colors (need at least 3)"
    })
    
    # Check for colors slide content
    colors_slide_found = False
    for slide in prs.slides:
        slide_text = ''
        for shape in slide.shapes:
            if hasattr(shape, 'text'):
                slide_text += shape.text.lower() + ' '
        
        if 'color' in slide_text and any(color in slide_text for color in ['141413', 'faf9f5', 'd97757', '6a9bcc', '788c5d']):
            colors_slide_found = True
            break
    
    checks.append({
        "name": "colors_slide_content", 
        "passed": colors_slide_found, 
        "detail": "Colors slide with hex values found" if colors_slide_found else "Colors slide content not found"
    })
    
    # Check for typography slide
    typography_slide_found = False
    for slide in prs.slides:
        slide_text = ''
        for shape in slide.shapes:
            if hasattr(shape, 'text'):
                slide_text += shape.text.lower() + ' '
        
        if 'typography' in slide_text or ('poppins' in slide_text and 'lora' in slide_text) or ('heading' in slide_text and 'body' in slide_text):
            typography_slide_found = True
            break
    
    checks.append({
        "name": "typography_slide_content", 
        "passed": typography_slide_found, 
        "detail": "Typography slide found" if typography_slide_found else "Typography slide not found"
    })
    
    return checks

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "error", "passed": False, "detail": "Invalid arguments"}]}))
        sys.exit(1)
    
    workspace_dir = sys.argv[1]
    checks = check_presentation(workspace_dir)
    
    passed_count = sum(1 for check in checks if check['passed'])
    total_count = len(checks)
    score = passed_count / total_count if total_count > 0 else 0.0
    overall_passed = score >= 0.8
    
    result = {
        "passed": overall_passed,
        "score": score,
        "checks": checks
    }
    
    print(json.dumps(result))