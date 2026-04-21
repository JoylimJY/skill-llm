import sys
import os
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_styled_presentation(workspace_dir):
    checks = []
    score = 0
    
    # Check if output file exists
    output_file = os.path.join(workspace_dir, 'styled_presentation.pptx')
    if not os.path.exists(output_file):
        checks.append({"name": "output_file_exists", "passed": False, "detail": "styled_presentation.pptx not found"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    checks.append({"name": "output_file_exists", "passed": True, "detail": "styled_presentation.pptx found"})
    
    try:
        # Load the styled presentation
        prs = Presentation(output_file)
        
        # Define Anthropic brand colors for comparison
        anthropic_colors = {
            'dark': (20, 20, 19),     # #141413
            'light': (250, 249, 245), # #faf9f5
            'mid_gray': (176, 174, 165), # #b0aea5
            'light_gray': (232, 230, 220), # #e8e6dc
            'orange': (217, 119, 87), # #d97757
            'blue': (106, 155, 204),  # #6a9bcc
            'green': (120, 140, 93)   # #788c5d
        }
        
        def color_matches_brand(rgb_color, tolerance=10):
            """Check if a color matches any Anthropic brand color within tolerance"""
            if rgb_color is None:
                return False
            
            r, g, b = rgb_color.rgb
            for brand_color in anthropic_colors.values():
                br, bg, bb = brand_color
                if (abs(r - br) <= tolerance and 
                    abs(g - bg) <= tolerance and 
                    abs(b - bb) <= tolerance):
                    return True
            return False
        
        # Check for brand color usage
        brand_colors_found = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if hasattr(run.font, 'color') and run.font.color.rgb:
                                if color_matches_brand(run.font.color):
                                    brand_colors_found = True
                                    break
                
                # Check shape fill colors
                if hasattr(shape, 'fill') and hasattr(shape.fill, 'fore_color'):
                    try:
                        if hasattr(shape.fill.fore_color, 'rgb') and shape.fill.fore_color.rgb:
                            if color_matches_brand(shape.fill.fore_color):
                                brand_colors_found = True
                    except:
                        pass
        
        checks.append({"name": "brand_colors_applied", "passed": brand_colors_found, "detail": "Anthropic brand colors found" if brand_colors_found else "No Anthropic brand colors detected"})
        
        # Check for font styling (look for any font changes)
        font_styling_found = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if hasattr(run.font, 'name') and run.font.name:
                                font_name = run.font.name.lower()
                                if 'poppins' in font_name or 'lora' in font_name or 'arial' in font_name or 'georgia' in font_name:
                                    font_styling_found = True
                                    break
        
        checks.append({"name": "typography_applied", "passed": font_styling_found, "detail": "Typography styling found" if font_styling_found else "No typography styling detected"})
        
        # Check that presentation has content (slides exist)
        has_content = len(prs.slides) > 0
        checks.append({"name": "content_preserved", "passed": has_content, "detail": f"Presentation has {len(prs.slides)} slides" if has_content else "No slides found in presentation"})
        
    except Exception as e:
        checks.append({"name": "file_readable", "passed": False, "detail": f"Error reading presentation: {str(e)}"})
    
    # Calculate score
    passed_checks = sum(1 for check in checks if check['passed'])
    total_checks = len(checks)
    score = passed_checks / total_checks if total_checks > 0 else 0.0
    
    return {
        "passed": score >= 0.75,
        "score": score,
        "checks": checks
    }

if __name__ == '__main__':
    workspace_dir = sys.argv[1] if len(sys.argv) > 1 else '.'
    result = check_styled_presentation(workspace_dir)
    print(result)