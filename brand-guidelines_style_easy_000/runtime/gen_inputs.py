from pptx import Presentation
from pptx.util import Inches, Pt

# Create a simple presentation with text content
prs = Presentation()

# Title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Company Overview"
subtitle.text = "Our Mission and Vision"

# Content slide
content_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Key Points"
content.text = "• Innovation in AI technology\n• Customer-focused solutions\n• Sustainable growth"

# Save the presentation
prs.save('presentation.pptx')
print("Generated presentation.pptx with sample content")