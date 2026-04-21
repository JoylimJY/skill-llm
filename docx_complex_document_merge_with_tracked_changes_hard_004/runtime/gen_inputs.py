import os
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_section1():
    doc = Document()
    
    # Add title
    title = doc.add_heading('Executive Summary', level=1)
    
    # Add content with specific markers
    p1 = doc.add_paragraph('This executive summary provides an overview of our quarterly performance.')
    p2 = doc.add_paragraph('Key highlights include:')
    
    # Add bullet points
    doc.add_paragraph('Revenue growth of 15% year-over-year', style='List Bullet')
    doc.add_paragraph('Market expansion into three new regions', style='List Bullet')
    doc.add_paragraph('Successful product launch of Widget Pro 2024', style='List Bullet')
    
    p3 = doc.add_paragraph('The organization has demonstrated strong resilience and adaptability in challenging market conditions.')
    
    doc.save('section1.docx')

def create_section2():
    doc = Document()
    
    # Add title
    title = doc.add_heading('Financial Analysis', level=1)
    
    # Add subheading
    subtitle = doc.add_heading('Revenue Breakdown', level=2)
    
    # Add table with financial data
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Quarter'
    hdr_cells[1].text = 'Revenue ($M)'
    hdr_cells[2].text = 'Growth (%)'
    
    # Add data rows
    quarters = [('Q1 2024', '125.6', '12%'), ('Q2 2024', '138.2', '15%'), ('Q3 2024', '142.8', '18%')]
    for quarter, revenue, growth in quarters:
        row_cells = table.add_row().cells
        row_cells[0].text = quarter
        row_cells[1].text = revenue
        row_cells[2].text = growth
    
    # Add analysis paragraph
    p1 = doc.add_paragraph('The financial performance shows consistent upward trajectory with strong momentum in Q3.')
    p2 = doc.add_paragraph('Cost optimization initiatives have contributed significantly to improved margins.')
    
    doc.save('section2.docx')

def create_section3():
    doc = Document()
    
    # Add title
    title = doc.add_heading('Strategic Recommendations', level=1)
    
    # Add subheadings and content
    subtitle1 = doc.add_heading('Short-term Initiatives', level=2)
    p1 = doc.add_paragraph('We recommend the following immediate actions:')
    doc.add_paragraph('Expand marketing efforts in high-growth regions', style='List Number')
    doc.add_paragraph('Increase production capacity by 25%', style='List Number')
    doc.add_paragraph('Launch customer retention program', style='List Number')
    
    subtitle2 = doc.add_heading('Long-term Strategy', level=2)
    p2 = doc.add_paragraph('Our long-term strategic vision includes:')
    p3 = doc.add_paragraph('Investment in research and development will drive innovation and maintain competitive advantage.')
    p4 = doc.add_paragraph('Partnership opportunities should be explored to accelerate market penetration.')
    
    doc.save('section3.docx')

if __name__ == '__main__':
    create_section1()
    create_section2()
    create_section3()
    print('Generated section1.docx, section2.docx, and section3.docx')