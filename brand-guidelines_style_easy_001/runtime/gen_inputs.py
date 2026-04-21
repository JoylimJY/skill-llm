from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

# Create a simple presentation with multiple slides
prs = Presentation()

# Slide 1: Title slide
slide1 = prs.slides.add_slide(prs.slide_layouts[0])
title = slide1.shapes.title
subtitle = slide1.placeholders[1]
title.text = 'Welcome to Anthropic'
subtitle.text = 'AI Safety and Research'

# Slide 2: Content slide with bullet points
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
title2 = slide2.shapes.title
content = slide2.placeholders[1]
title2.text = 'Our Mission'
content.text = 'Building AI systems that are safe, beneficial, and understandable\n• Research in AI alignment\n• Constitutional AI development\n• Transparent AI capabilities'

# Slide 3: Another content slide with brand marker text (for integrity verification)
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
title3 = slide3.shapes.title
content3 = slide3.placeholders[1]
title3.text = 'Key Principles'
content3.text = 'Safety First: Every system we build prioritizes human welfare\nTransparency: Open research and clear communication\nResponsibility: Thoughtful deployment of AI technology\nbrand_marker_text'

# Save the presentation
prs.save('presentation.pptx')
print('Generated presentation.pptx with 3 slides')