import os
import sys
import json
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_presentation(workspace_dir):
    checks = []
    
    # Check if the presentation file exists
    pptx_path = os.path.join(workspace_dir, 'claude_capabilities.pptx')
    file_exists = os.path.exists(pptx_path)
    checks.append({
        'name': 'presentation_file_exists',
        'passed': file_exists,
        'detail': f'PowerPoint file exists at {pptx_path}' if file_exists else 'PowerPoint file not found'
    })
    
    if not file_exists:
        return checks
    
    try:
        prs = Presentation(pptx_path)
        
        # Check slide count
        slide_count_correct = len(prs.slides) == 4
        checks.append({
            'name': 'slide_count',
            'passed': slide_count_correct,
            'detail': f'Found {len(prs.slides)} slides' + (' (correct)' if slide_count_correct else ' (expected 4)')
        })
        
        # Check slide 1 - Title slide
        if len(prs.slides) >= 1:
            slide1 = prs.slides[0]
            title_found = False
            subtitle_found = False
            
            for shape in slide1.shapes:
                if hasattr(shape, 'text'):
                    text_lower = shape.text.lower()
                    if 'claude ai capabilities' in text_lower:
                        title_found = True
                    if 'anthropic' in text_lower and 'advanced' in text_lower and 'assistant' in text_lower:
                        subtitle_found = True
            
            checks.append({
                'name': 'title_slide_content',
                'passed': title_found and subtitle_found,
                'detail': f'Title found: {title_found}, Subtitle found: {subtitle_found}'
            })
        
        # Check slide 2 - Natural Language Understanding
        if len(prs.slides) >= 2:
            slide2 = prs.slides[1]
            nlu_content_found = False
            
            for shape in slide2.shapes:
                if hasattr(shape, 'text'):
                    text_lower = shape.text.lower()
                    if 'natural language understanding' in text_lower or ('natural language' in text_lower and 'understanding' in text_lower):
                        nlu_content_found = True
                        break
            
            checks.append({
                'name': 'nlu_slide_content',
                'passed': nlu_content_found,
                'detail': f'Natural Language Understanding slide content found: {nlu_content_found}'
            })
        
        # Check slide 3 - Code Generation & Analysis
        if len(prs.slides) >= 3:
            slide3 = prs.slides[2]
            code_content_found = False
            
            for shape in slide3.shapes:
                if hasattr(shape, 'text'):
                    text_lower = shape.text.lower()
                    if ('code' in text_lower and 'generation' in text_lower) or ('code' in text_lower and 'analysis' in text_lower):
                        code_content_found = True
                        break
            
            checks.append({
                'name': 'code_slide_content',
                'passed': code_content_found,
                'detail': f'Code Generation & Analysis slide content found: {code_content_found}'
            })
        
        # Check slide 4 - Safety & Alignment
        if len(prs.slides) >= 4:
            slide4 = prs.slides[3]
            safety_content_found = False
            
            for shape in slide4.shapes:
                if hasattr(shape, 'text'):
                    text_lower = shape.text.lower()
                    if ('safety' in text_lower and 'alignment' in text_lower) or ('safety' in text_lower or 'alignment' in text_lower):
                        safety_content_found = True
                        break
            
            checks.append({
                'name': 'safety_slide_content',
                'passed': safety_content_found,
                'detail': f'Safety & Alignment slide content found: {safety_content_found}'
            })
        
        # Check for Anthropic brand colors usage
        brand_colors_found = False
        anthropic_dark = RGBColor(0x14, 0x14, 0x13)  # #141413
        anthropic_light = RGBColor(0xfa, 0xf9, 0xf5)  # #faf9f5
        anthropic_orange = RGBColor(0xd9, 0x77, 0x57)  # #d97757
        
        color_usage_count = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'fill') and hasattr(shape.fill, 'fore_color'):
                    try:
                        color = shape.fill.fore_color.rgb
                        if color in [anthropic_dark, anthropic_light, anthropic_orange]:
                            color_usage_count += 1
                    except:
                        pass
                
                if hasattr(shape, 'text_frame'):
                    try:
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if hasattr(run.font, 'color') and hasattr(run.font.color, 'rgb'):
                                    color = run.font.color.rgb
                                    if color in [anthropic_dark, anthropic_light, anthropic_orange]:
                                        color_usage_count += 1
                    except:
                        pass
        
        brand_colors_found = color_usage_count > 0
        checks.append({
            'name': 'brand_colors_applied',
            'passed': brand_colors_found,
            'detail': f'Anthropic brand colors detected in {color_usage_count} elements'
        })
        
        # Check for font application attempts (Poppins/Lora or fallbacks)
        font_application_found = False
        font_count = 0
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    try:
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if hasattr(run.font, 'name') and run.font.name:
                                    font_name = run.font.name.lower()
                                    if any(font in font_name for font in ['poppins', 'lora', 'arial', 'georgia']):
                                        font_count += 1
                    except:
                        pass
        
        font_application_found = font_count > 0
        checks.append({
            'name': 'typography_applied',
            'passed': font_application_found,
            'detail': f'Brand typography or fallback fonts detected in {font_count} text elements'
        })
        
    except Exception as e:
        checks.append({
            'name': 'presentation_parsing',
            'passed': False,
            'detail': f'Error parsing presentation: {str(e)}'
        })
    
    return checks

def main():
    workspace_dir = sys.argv[1]
    checks = check_presentation(workspace_dir)
    
    passed_checks = sum(1 for check in checks if check['passed'])
    total_checks = len(checks)
    score = passed_checks / total_checks if total_checks > 0 else 0.0
    
    result = {
        'passed': score >= 0.8,
        'score': score,
        'checks': checks
    }
    
    print(json.dumps(result, indent=2))

if __name__ == '__main__':
    main()