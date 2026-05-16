import os
import random
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from openpyxl import Workbook

random.seed(42)

WORKSPACE = Path("/workspace")

# --- Directory structure ---
dirs = [
    "submissions/contracts/2024/Q1",
    "submissions/contracts/2024/Q2",
    "submissions/contracts/2024/Q3",
    "submissions/amendments/pending",
    "submissions/amendments/approved",
    "submissions/exhibits/A",
    "submissions/exhibits/B",
    "internal/templates",
    "internal/archive/2023",
    "internal/drafts",
    "output/pdf_archive",  # Target output dir (empty, agent must populate it)
    "logs",
    "metadata",
]

for d in dirs:
    (WORKSPACE / d).mkdir(parents=True, exist_ok=True)

# ── Helper: write a .docx ─────────────────────────────────────────────────────
def make_docx(path: Path, title: str, body: str):
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(body)
    doc.save(str(path))

# ── Helper: write a .pptx ─────────────────────────────────────────────────────
def make_pptx(path: Path, title: str, content: str):
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = content
    prs.save(str(path))

# ── Helper: write a .xlsx ─────────────────────────────────────────────────────
def make_xlsx(path: Path, sheet_title: str, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    for row in rows:
        ws.append(row)
    wb.save(str(path))

# ════════════════════════════════════════════════════════════════════════════
# TARGET files — these MUST be converted to PDF
# ════════════════════════════════════════════════════════════════════════════

make_docx(
    WORKSPACE / "submissions/contracts/2024/Q1/NDA_Client_Acme_Corp.docx",
    "Non-Disclosure Agreement",
    "This Non-Disclosure Agreement ('Agreement') is entered into as of January 15, 2024 "
    "between Acme Corp ('Disclosing Party') and Initech LLC ('Receiving Party')."
)

make_docx(
    WORKSPACE / "submissions/contracts/2024/Q2/ServiceAgreement_GlobalTech.docx",
    "Master Service Agreement",
    "This Master Service Agreement governs the relationship between GlobalTech Inc and "
    "the firm for the provision of legal consultancy services starting April 1, 2024."
)

make_pptx(
    WORKSPACE / "submissions/amendments/pending/Amendment_3_TermSheet.pptx",
    "Amendment No. 3 – Term Sheet",
    "Key changes: (1) Extend exclusivity window to 180 days. "
    "(2) Revise indemnification clause 7.2. (3) Add arbitration rider."
)

make_xlsx(
    WORKSPACE / "submissions/exhibits/A/ExhibitA_FeeSchedule.xlsx",
    "Fee Schedule",
    [
        ["Service", "Rate (USD/hr)", "Estimated Hours", "Total"],
        ["Contract Drafting", 450, 20, 9000],
        ["Due Diligence", 380, 35, 13300],
        ["Litigation Support", 500, 15, 7500],
        ["Regulatory Filing", 320, 10, 3200],
    ]
)

make_docx(
    WORKSPACE / "submissions/exhibits/B/ExhibitB_Definitions.docx",
    "Exhibit B – Defined Terms",
    "'Affiliate' means any entity that directly or indirectly controls, is controlled by, "
    "or is under common control with the subject entity. 'Confidential Information' means "
    "any non-public information disclosed by either party."
)

make_pptx(
    WORKSPACE / "submissions/contracts/2024/Q3/PartnershipDeck_Meridian.pptx",
    "Meridian Partners – Joint Venture Overview",
    "Slide 1: Proposed equity split 60/40. "
    "Slide 2: Governance structure. "
    "Slide 3: Exit provisions."
)

# ════════════════════════════════════════════════════════════════════════════
# DISTRACTOR files — various non-target files to increase noise
# ════════════════════════════════════════════════════════════════════════════

# Plain text files
(WORKSPACE / "submissions/contracts/2024/Q1/checklist.txt").write_text(
    "TODO:\n- [ ] Get signatures\n- [ ] File with county clerk\n"
)
(WORKSPACE / "submissions/amendments/approved/approval_log.txt").write_text(
    "Amendment 1 approved 2023-11-01\nAmendment 2 approved 2024-01-15\n"
)
(WORKSPACE / "logs/conversion_log_2023.txt").write_text(
    "2023-12-01 10:00:00 INFO Conversion complete\n"
    "2023-12-02 11:23:00 INFO 3 files processed\n"
)

# CSV distractors
(WORKSPACE / "metadata/document_index.csv").write_text(
    "filename,type,status,uploaded_by\n"
    "NDA_Client_Acme_Corp.docx,contract,pending,j.smith\n"
    "ServiceAgreement_GlobalTech.docx,contract,pending,r.jones\n"
    "Amendment_3_TermSheet.pptx,amendment,pending,k.lee\n"
    "ExhibitA_FeeSchedule.xlsx,exhibit,pending,j.smith\n"
    "ExhibitB_Definitions.docx,exhibit,pending,r.jones\n"
    "PartnershipDeck_Meridian.pptx,contract,pending,k.lee\n"
)

# A .pdf already in archive (should NOT confuse agent)
(WORKSPACE / "internal/archive/2023/OldContract_2023.pdf").write_text(
    "%PDF-1.4 fake archive pdf content"
)

# Internal template docs (should NOT be converted)
make_docx(
    WORKSPACE / "internal/templates/TemplateMSA.docx",
    "TEMPLATE – Do Not Use for Filing",
    "[TEMPLATE PLACEHOLDER] This document is for internal reference only."
)

make_xlsx(
    WORKSPACE / "internal/drafts/DraftBudget_Q4.xlsx",
    "Draft",
    [["Item", "Budget"], ["Legal Research", 5000], ["Filings", 1200]]
)

# JSON metadata sidecar
import json
(WORKSPACE / "metadata/case_metadata.json").write_text(json.dumps({
    "case_id": "2024-LIT-0042",
    "client": "Acme Corp",
    "matter": "Corporate Advisory",
    "documents": 6,
    "status": "pending_archival"
}, indent=2))

# Shell scripts (distractors)
(WORKSPACE / "internal/archive/2023/reindex.sh").write_text(
    "#!/bin/bash\necho 'Reindexing archive...'\nfind . -name '*.pdf' | wc -l\n"
)

print("Workspace generation complete.")
print("\nTarget files to convert:")
target_files = [
    "submissions/contracts/2024/Q1/NDA_Client_Acme_Corp.docx",
    "submissions/contracts/2024/Q2/ServiceAgreement_GlobalTech.docx",
    "submissions/amendments/pending/Amendment_3_TermSheet.pptx",
    "submissions/exhibits/A/ExhibitA_FeeSchedule.xlsx",
    "submissions/exhibits/B/ExhibitB_Definitions.docx",
    "submissions/contracts/2024/Q3/PartnershipDeck_Meridian.pptx",
]
for f in target_files:
    p = WORKSPACE / f
    print(f"  {'✓' if p.exists() else '✗'} {f}")