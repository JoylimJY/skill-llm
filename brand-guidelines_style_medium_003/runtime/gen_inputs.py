import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Create a sample presentation with various elements
prs = Presentation()

# Slide 1: Title slide
slide1_layout = prs.slide_layouts[0]
slide1 = prs.slides.add_slide(slide1_layout)
title1 = slide1.shapes.title
subtitle1 = slide1.placeholders[1]

title1.text = "SAMPLE COMPANY PRESENTATION"
subtitle1.text = "Marketing Strategy Overview"

# Slide 2: Content slide with text and shapes
slide2_layout = prs.slide_layouts[1]
slide2 = prs.slides.add_slide(slide2_layout)
title2 = slide2.shapes.title
content2 = slide2.placeholders[1]

title2.text = "Key Performance Metrics"
content2.text = "Our quarterly results show significant growth across all major product lines. Customer satisfaction has increased by 25% while operational costs have decreased by 15%."

# Add some shapes to slide 2
rect_shape = slide2.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(1), Inches(4), Inches(2), Inches(1)
)
circle_shape = slide2.shapes.add_shape(
    MSO_SHAPE.OVAL, Inches(4), Inches(4), Inches(2), Inches(1)
)
triangle_shape = slide2.shapes.add_shape(
    MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(7), Inches(4), Inches(1.5), Inches(1)
)

# Slide 3: Another content slide
slide3_layout = prs.slide_layouts[1]
slide3 = prs.slides.add_slide(slide3_layout)
title3 = slide3.shapes.title
content3 = slide3.placeholders[1]

title3.text = "Future Roadmap"
content3.text = "Looking ahead to next quarter, we plan to expand into three new markets and launch two innovative product features that will drive customer engagement."

# Add a shape to slide 3
square_shape = slide3.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(4.5), Inches(2), Inches(1.5)
)

# Save the presentation
prs.save('company_presentation.pptx')
print("Generated company_presentation.pptx with 3 slides containing titles, content, and shapes")