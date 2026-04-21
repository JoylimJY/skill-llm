import sys
import os
import json
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "Import Error", "passed": False, "detail": "openpyxl not available"}]}))
    sys.exit(1)

def eval_budget_file(workspace_dir):
    checks = []
    workspace_path = Path(workspace_dir)
    
    # Check if budget.xlsx exists
    budget_file = workspace_path / 'budget.xlsx'
    if not budget_file.exists():
        checks.append({"name": "File Exists", "passed": False, "detail": "budget.xlsx not found"})
        return checks
    
    checks.append({"name": "File Exists", "passed": True, "detail": "budget.xlsx found"})
    
    try:
        # Load workbook with formulas
        wb_formulas = load_workbook(str(budget_file), data_only=False)
        ws_formulas = wb_formulas.active
        
        # Load workbook with calculated values
        wb_values = load_workbook(str(budget_file), data_only=True)
        ws_values = wb_values.active
        
        # Check headers (case-insensitive)
        expected_headers = ['category', 'budget', 'actual', 'variance']
        actual_headers = []
        for col in range(1, 5):
            cell_val = ws_values.cell(row=1, column=col).value
            if cell_val:
                actual_headers.append(str(cell_val).lower().strip())
        
        headers_match = all(header in actual_headers for header in expected_headers)
        checks.append({"name": "Headers Present", "passed": headers_match, "detail": f"Expected headers found: {headers_match}. Got: {actual_headers}"})
        
        # Check data rows (flexible matching)
        expected_data = [
            ('office supplies', 500, 450),
            ('marketing', 2000, 1800),
            ('travel', 1500, 1650),
            ('utilities', 800, 825)
        ]
        
        data_found = 0
        for row in range(2, 10):  # Check multiple rows to find data
            category = ws_values.cell(row=row, column=1).value
            budget = ws_values.cell(row=row, column=2).value
            actual = ws_values.cell(row=row, column=3).value
            
            if category and budget and actual:
                category_str = str(category).lower().strip()
                for exp_cat, exp_budget, exp_actual in expected_data:
                    if exp_cat in category_str and abs(float(budget) - exp_budget) < 1 and abs(float(actual) - exp_actual) < 1:
                        data_found += 1
                        break
        
        data_complete = data_found >= 4
        checks.append({"name": "Data Rows", "passed": data_complete, "detail": f"Found {data_found}/4 expected data rows"})
        
        # Check variance formulas (look for formulas in column 4)
        variance_formulas = 0
        for row in range(2, 10):
            cell_formula = ws_formulas.cell(row=row, column=4).value
            if cell_formula and isinstance(cell_formula, str) and cell_formula.startswith('='):
                # Check if formula references columns 3 and 2 (Actual - Budget)
                if 'c' in cell_formula.lower() and 'b' in cell_formula.lower():
                    variance_formulas += 1
        
        formulas_present = variance_formulas >= 3
        checks.append({"name": "Variance Formulas", "passed": formulas_present, "detail": f"Found {variance_formulas} variance formulas"})
        
        # Check for total row with SUM formulas
        total_formulas = 0
        for row in range(6, 12):  # Look for total row
            for col in range(2, 5):  # Budget, Actual, Variance columns
                cell_formula = ws_formulas.cell(row=row, column=col).value
                if cell_formula and isinstance(cell_formula, str) and 'sum' in cell_formula.lower():
                    total_formulas += 1
        
        totals_present = total_formulas >= 2
        checks.append({"name": "Total Formulas", "passed": totals_present, "detail": f"Found {total_formulas} SUM formulas for totals"})
        
        # Check calculated variance values are reasonable
        variance_correct = 0
        expected_variances = [-50, -200, 150, 25]  # Actual - Budget
        for row in range(2, 6):
            variance = ws_values.cell(row=row, column=4).value
            if variance is not None:
                try:
                    var_val = float(variance)
                    if any(abs(var_val - exp) < 1 for exp in expected_variances):
                        variance_correct += 1
                except:
                    pass
        
        variances_calculated = variance_correct >= 3
        checks.append({"name": "Variance Calculations", "passed": variances_calculated, "detail": f"Found {variance_correct}/4 correct variance calculations"})
        
        wb_formulas.close()
        wb_values.close()
        
    except Exception as e:
        checks.append({"name": "File Processing", "passed": False, "detail": f"Error processing file: {str(e)}"})
    
    return checks

def main():
    if len(sys.argv) != 2:
        print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "Usage", "passed": False, "detail": "Workspace directory required"}]}))
        return
    
    workspace_dir = sys.argv[1]
    checks = eval_budget_file(workspace_dir)
    
    score = sum(1 for c in checks if c['passed']) / len(checks) if checks else 0.0
    passed = score >= 0.8
    
    result = {
        "passed": passed,
        "score": score,
        "checks": checks
    }
    
    print(json.dumps(result, indent=2))

if __name__ == "__main__":
    main()