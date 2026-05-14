#!/usr/bin/env python3
"""
Evaluation script for the NexaPay investor pitch deck task.
Tests PptxGenJS-specific constraints, content completeness, and proprietary pitfall avoidance.
"""

import sys
import json
import re
import subprocess
import zipfile
from pathlib import Path

def run_markitdown(pptx_path: Path) -> str:
    try:
        result = subprocess.run(
            ["python3", "-m", "markitdown", str(pptx_path)],
            capture_output=True, text=True, timeout=60
        )
        return result.stdout + result.stderr
    except Exception as e:
        return f"ERROR: {e}"

def find_pptx(workspace: Path) -> Path | None:
    candidates = list(workspace.rglob("nexapay_pitch.pptx"))
    if candidates:
        return candidates[0]
    # Also accept any .pptx in workspace root
    candidates = list(workspace.glob("*.pptx"))
    if candidates:
        return candidates[0]
    return None

def find_js_script(workspace: Path) -> Path | None:
    """Find the main JS generation script."""
    candidates = list(workspace.rglob("*.js"))
    # Exclude node_modules
    candidates = [c for c in candidates if "node_modules" not in str(c)]
    if candidates:
        return candidates[0]
    return None

def check_js_source(workspace: Path) -> dict:
    """Check the JS source for proprietary pitfalls."""
    issues = []
    warnings = []
    
    js_files = [f for f in workspace.rglob("*.js") if "node_modules" not in str(f)]
    if not js_files:
        return {"found": False, "issues": ["No JS file found"], "warnings": []}
    
    all_js = ""
    for f in js_files:
        try:
            all_js += f.read_text(errors="replace")
        except:
            pass
    
    # Check 1: No # prefix on hex colors
    # Look for color patterns with # prefix
    hash_hex_pattern = re.compile(r'color\s*:\s*["\']#[0-9A-Fa-f]{6}["\']')
    if hash_hex_pattern.search(all_js):
        issues.append("Found '#' prefix in hex color strings (causes file corruption per SKILL.md)")
    
    # Check 2: No 8-char hex in shadow color (opacity encoded in color)
    eight_char_hex = re.compile(r'color\s*:\s*["\'][0-9A-Fa-f]{8}["\']')
    if eight_char_hex.search(all_js):
        issues.append("Found 8-character hex color string (opacity must not be encoded in color per SKILL.md)")
    
    # Check 3: shadow object not reused — check if shadow objects are defined as factory functions
    # or as separate inline objects. If a variable is assigned once and used in multiple addShape calls,
    # that's a violation.
    # Heuristic: if there's a const/let/var shadow = { and it's used 2+ times via the same variable name
    shadow_var_pattern = re.compile(r'(?:const|let|var)\s+(\w*[Ss]hadow\w*)\s*=\s*\{')
    shadow_vars = shadow_var_pattern.findall(all_js)
    for var in shadow_vars:
        # Count usages
        usages = len(re.findall(r'\b' + re.escape(var) + r'\b', all_js))
        if usages > 2:  # defined once + used 2+ times = reuse
            issues.append(f"Shadow object '{var}' appears to be reused across multiple calls (causes corruption per SKILL.md). Use a factory function.")
    
    # Check 4: charSpacing not letterSpacing
    if re.search(r'letterSpacing', all_js):
        issues.append("Used 'letterSpacing' instead of 'charSpacing' (letterSpacing is silently ignored per SKILL.md)")
    
    # Check 5: No unicode bullets
    if '•' in all_js or '&#8226;' in all_js:
        issues.append("Unicode bullet '•' found in JS source (creates double bullets per SKILL.md)")
    
    # Check 6: breakLine used for multi-line text arrays
    # If addText is used with array and multiple text items, breakLine should be present
    # This is a soft check — we look for multi-item arrays without breakLine
    
    return {"found": True, "issues": issues, "warnings": warnings, "js_content": all_js}

def evaluate(workspace_str: str):
    workspace = Path(workspace_str)
    checks = []
    
    # ── CHECK 1: Output file exists ─────────────────────────────────────────
    pptx_path = find_pptx(workspace)
    checks.append({
        "name": "output_file_exists",
        "passed": pptx_path is not None,
        "detail": f"Found: {pptx_path}" if pptx_path else "nexapay_pitch.pptx not found in workspace"
    })
    
    if pptx_path is None:
        result = {"passed": False, "score": 0.0, "checks": checks}
        print(json.dumps(result))
        return
    
    # ── CHECK 2: Valid PPTX structure ───────────────────────────────────────
    try:
        from pptx import Presentation
        prs = Presentation(str(pptx_path))
        slide_count = len(prs.slides)
        checks.append({
            "name": "valid_pptx_structure",
            "passed": True,
            "detail": f"Valid PPTX with {slide_count} slides"
        })
    except Exception as e:
        checks.append({
            "name": "valid_pptx_structure",
            "passed": False,
            "detail": f"PPTX is corrupted or invalid: {e}"
        })
        result = {"passed": False, "score": 0.0, "checks": checks}
        print(json.dumps(result))
        return
    
    # ── CHECK 3: Slide count (5 slides required) ────────────────────────────
    checks.append({
        "name": "slide_count_5",
        "passed": slide_count == 5,
        "detail": f"Expected 5 slides, found {slide_count}"
    })
    
    # ── CHECK 4-8: Content via markitdown ────────────────────────────────────
    md_content = run_markitdown(pptx_path)
    md_lower = md_content.lower()
    
    # Check 4: Title slide content
    has_nexapay = "nexapay" in md_lower
    checks.append({
        "name": "title_slide_company_name",
        "passed": has_nexapay,
        "detail": "Found 'NexaPay' in extracted text" if has_nexapay else "Company name 'NexaPay' not found in presentation"
    })
    
    # Check 5: Problem slide content
    problem_keywords = ["3-5 days" in md_content or "3 to 5 days" in md_lower or 
                       "settlement" in md_lower and "days" in md_lower,
                       "2.8" in md_content or "fee" in md_lower,
                       "47%" in md_content or "47 percent" in md_lower or "cfos" in md_lower,
                       "$120b" in md_lower or "120b" in md_lower or "120 billion" in md_lower or "failed" in md_lower]
    problem_score = sum(1 for k in problem_keywords if k)
    checks.append({
        "name": "problem_slide_content",
        "passed": problem_score >= 2,
        "detail": f"Problem slide content checks: {problem_score}/4 key data points found"
    })
    
    # Check 6: Solution slide with features
    solution_keywords = [
        "3-second" in md_lower or "3 second" in md_lower or "sub-3" in md_lower or "settlement" in md_lower,
        "soc 2" in md_lower or "bank-grade" in md_lower or "security" in md_lower,
        "190" in md_content or "countries" in md_lower,
        "0.4%" in md_content or "0.4 percent" in md_lower or "flat fee" in md_lower
    ]
    solution_score = sum(1 for k in solution_keywords if k)
    checks.append({
        "name": "solution_slide_features",
        "passed": solution_score >= 3,
        "detail": f"Solution slide feature checks: {solution_score}/4 features found"
    })
    
    # Check 7: Traction stats
    traction_keywords = [
        "$42m" in md_lower or "42m" in md_lower or "42 million" in md_lower,
        "380" in md_content,
        "3.2x" in md_lower or "3.2" in md_content
    ]
    traction_score = sum(1 for k in traction_keywords if k)
    checks.append({
        "name": "traction_slide_stats",
        "passed": traction_score >= 2,
        "detail": f"Traction stats checks: {traction_score}/3 key stats found"
    })
    
    # Check 8: Closing slide with Series A ask
    closing_keywords = [
        "$12m" in md_lower or "12m" in md_lower or "12 million" in md_lower or "series a" in md_lower,
        "invest@nexapay.io" in md_lower or "nexapay.io" in md_lower,
        "40%" in md_content or "35%" in md_content or "25%" in md_content
    ]
    closing_score = sum(1 for k in closing_keywords if k)
    checks.append({
        "name": "closing_slide_content",
        "passed": closing_score >= 2,
        "detail": f"Closing slide checks: {closing_score}/3 key elements found"
    })
    
    # ── CHECK 9: No placeholder text left ───────────────────────────────────
    placeholder_patterns = ["xxxx", "lorem", "ipsum", "placeholder", "[your", "todo", "tbd"]
    found_placeholders = [p for p in placeholder_patterns if p in md_lower]
    checks.append({
        "name": "no_placeholder_text",
        "passed": len(found_placeholders) == 0,
        "detail": f"Placeholder text found: {found_placeholders}" if found_placeholders else "No placeholder text found"
    })
    
    # ── CHECK 10-14: JS source code proprietary checks ───────────────────────
    js_analysis = check_js_source(workspace)
    
    if js_analysis["found"]:
        # Check 10: No # prefix in hex colors
        no_hash_issue = not any("'#'" in i or "hash" in i.lower() for i in js_analysis["issues"] 
                                 if "prefix" in i or "#" in i)
        hash_issues = [i for i in js_analysis["issues"] if "#" in i]
        checks.append({
            "name": "no_hash_prefix_in_hex_colors",
            "passed": len(hash_issues) == 0,
            "detail": "No '#' prefix found in hex colors" if len(hash_issues) == 0 else f"Issues: {hash_issues}"
        })
        
        # Check 11: No 8-char hex for opacity
        eight_char_issues = [i for i in js_analysis["issues"] if "8-character" in i or "8-char" in i]
        checks.append({
            "name": "no_8char_hex_opacity",
            "passed": len(eight_char_issues) == 0,
            "detail": "No 8-char hex opacity encoding found" if len(eight_char_issues) == 0 else f"Issues: {eight_char_issues}"
        })
        
        # Check 12: No letterSpacing
        letter_spacing_issues = [i for i in js_analysis["issues"] if "letterSpacing" in i]
        checks.append({
            "name": "charspacing_not_letterspacing",
            "passed": len(letter_spacing_issues) == 0,
            "detail": "charSpacing used correctly (no letterSpacing found)" if len(letter_spacing_issues) == 0 else f"Issues: {letter_spacing_issues}"
        })
        
        # Check 13: No unicode bullets
        bullet_issues = [i for i in js_analysis["issues"] if "bullet" in i.lower() or "•" in i]
        checks.append({
            "name": "no_unicode_bullets",
            "passed": len(bullet_issues) == 0,
            "detail": "No unicode bullets found" if len(bullet_issues) == 0 else f"Issues: {bullet_issues}"
        })
        
        # Check 14: Shadow object not reused
        shadow_reuse_issues = [i for i in js_analysis["issues"] if "shadow" in i.lower() and "reused" in i.lower()]
        checks.append({
            "name": "shadow_object_not_reused",
            "passed": len(shadow_reuse_issues) == 0,
            "detail": "Shadow objects appear to be fresh per call" if len(shadow_reuse_issues) == 0 else f"Issues: {shadow_reuse_issues}"
        })
    else:
        # JS not found — check if pptx was generated some other way (still valid if content is correct)
        for check_name in ["no_hash_prefix_in_hex_colors", "no_8char_hex_opacity", 
                           "charspacing_not_letterspacing", "no_unicode_bullets", "shadow_object_not_reused"]:
            checks.append({
                "name": check_name,
                "passed": True,  # Can't verify without JS source
                "detail": "No JS source found — assuming correct (PPTX created by other means)"
            })
    
    # ── CHECK 15: PPTX XML — no unicode bullets in actual output ─────────────
    try:
        bullet_found_in_xml = False
        with zipfile.ZipFile(str(pptx_path), 'r') as z:
            slide_files = [n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml', n)]
            for sf in slide_files:
                xml_content = z.read(sf).decode('utf-8', errors='replace')
                if '•' in xml_content or '\u2022' in xml_content:
                    bullet_found_in_xml = True
                    break
        checks.append({
            "name": "no_unicode_bullets_in_pptx_xml",
            "passed": not bullet_found_in_xml,
            "detail": "No unicode bullet characters in PPTX XML" if not bullet_found_in_xml else "Unicode bullet '•' found in PPTX XML slides"
        })
    except Exception as e:
        checks.append({
            "name": "no_unicode_bullets_in_pptx_xml",
            "passed": False,
            "detail": f"Could not inspect PPTX XML: {e}"
        })
    
    # ── CHECK 16: Chart data present (bar chart with traction data) ──────────
    try:
        from pptx.util import Inches
        has_chart = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_chart:
                    has_chart = True
                    break
        checks.append({
            "name": "bar_chart_present",
            "passed": has_chart,
            "detail": "Bar chart found in presentation" if has_chart else "No chart found — traction slide should have a bar chart"
        })
    except Exception as e:
        checks.append({
            "name": "bar_chart_present",
            "passed": False,
            "detail": f"Could not check for charts: {e}"
        })
    
    # ── CHECK 17: At least one image/icon in slides ──────────────────────────
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        has_image = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    has_image = True
                    break
        checks.append({
            "name": "icons_or_images_present",
            "passed": has_image,
            "detail": "Image/icon found in presentation" if has_image else "No images/icons found — solution slide should have icons per brief"
        })
    except Exception as e:
        checks.append({
            "name": "icons_or_images_present",
            "passed": False,
            "detail": f"Could not check for images: {e}"
        })
    
    # ── CHECK 18: 16x9 layout ────────────────────────────────────────────────
    try:
        from pptx.util import Emu
        width_inches = prs.slide_width.inches
        height_inches = prs.slide_height.inches
        is_16x9 = abs(width_inches - 10.0) < 0.1 and abs(height_inches - 5.625) < 0.1
        checks.append({
            "name": "layout_16x9",
            "passed": is_16x9,
            "detail": f"Slide dimensions: {width_inches:.2f}\" × {height_inches:.2f}\" (expected 10\" × 5.625\")"
        })
    except Exception as e:
        checks.append({
            "name": "layout_16x9",
            "passed": False,
            "detail": f"Could not check slide layout: {e}"
        })
    
    # ── CHECK 19: Tagline text present (wide character spacing) ──────────────
    tagline_present = (
        "instant" in md_lower and "global" in md_lower and "payments" in md_lower
    )
    # Also check if charSpacing is used in JS
    char_spacing_used = False
    if js_analysis["found"] and "js_content" in js_analysis:
        char_spacing_used = "charSpacing" in js_analysis["js_content"]
    checks.append({
        "name": "tagline_with_char_spacing",
        "passed": tagline_present and (char_spacing_used or not js_analysis["found"]),
        "detail": (
            f"Tagline present: {tagline_present}, charSpacing used in JS: {char_spacing_used}"
        )
    })
    
    # ── Score computation ────────────────────────────────────────────────────
    total = len(checks)
    passed_count = sum(1 for c in checks if c["passed"])
    score = passed_count / total
    
    # Must pass critical checks to be considered "passed" overall
    critical_checks = [
        "output_file_exists",
        "valid_pptx_structure", 
        "slide_count_5",
        "title_slide_company_name",
        "solution_slide_features",
        "traction_slide_stats",
        "bar_chart_present",
        "no_hash_prefix_in_hex_colors",
        "no_unicode_bullets_in_pptx_xml",
    ]
    critical_passed = all(
        next((c["passed"] for c in checks if c["name"] == cn), False)
        for cn in critical_checks
    )
    
    overall_passed = critical_passed and score >= 0.72
    
    result = {
        "passed": overall_passed,
        "score": round(score, 3),
        "checks": checks
    }
    print(json.dumps(result, indent=2))

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "args", "passed": False, "detail": "Usage: eval.py <workspace>"}]}))
        sys.exit(1)
    evaluate(sys.argv[1])