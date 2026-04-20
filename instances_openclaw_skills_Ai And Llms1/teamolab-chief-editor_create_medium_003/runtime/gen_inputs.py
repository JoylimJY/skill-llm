from pathlib import Path
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

# Deterministic content with markers for evaluation
base = Path('.')

# Text source with embedded URL and marker
(base / 'source_notes.txt').write_text(
    'Chief editor source packet\n'
    'MARKER_ALPHA: editorial-core-2025\n'
    'Reference URL: https://example.com/editorial-guide\n'
    'This material explains the desired tone, structure, and final deliverable.\n',
    encoding='utf-8'
)

# DOCX source with marker content
_doc = Document()
_doc.add_heading('Editorial Brief', level=1)
_doc.add_paragraph('MARKER_BETA: structure-check-8842')
_doc.add_paragraph('The final article should end with a references section.')
_doc.add_paragraph('Secondary URL: https://www.example.org/style-manual')
_doc.save(base / 'brief.docx')

# PDF source with marker content
pdf_path = base / 'report.pdf'
c = canvas.Canvas(str(pdf_path), pagesize=letter)
c.drawString(72, 720, 'MARKER_GAMMA: pdf-anchor-551')
c.drawString(72, 700, 'Use a professional chief editor style.')
c.drawString(72, 680, 'Additional URL: https://example.net/research-note')
c.save()
