from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
import json
import random

random.seed(1337)
base = Path('.')

# Text input with markers
(base / 'input_notes.txt').write_text(
    'Token Alert Task\n'
    'MARKER_ALPHA: preserve this phrase\n'
    'Threshold targets: 75 percent, 90 percent, 95 percent\n'
    'MARKER_BETA: final summary required\n',
    encoding='utf-8'
)

# JSON input with deterministic values
payload = {
    'session': 'clawdbot-demo',
    'limit': 200000,
    'used': 156000,
    'markers': ['MARKER_ALPHA', 'MARKER_BETA'],
    'thresholds': [75, 90, 95]
}
(base / 'input_data.json').write_text(json.dumps(payload, indent=2), encoding='utf-8')

# CSV input
(base / 'usage.csv').write_text(
    'label,value\n'
    'current,156000\n'
    'limit,200000\n'
    'remaining,44000\n',
    encoding='utf-8'
)

# DOCX input with a marker
doc = Document()
doc.add_heading('Token Alert Reference', level=1)
doc.add_paragraph('MARKER_GAMMA: include this in the final output if relevant.')
doc.add_paragraph('High warning begins at 75%. Critical warning begins at 90%. Emergency begins at 95%.')
doc.save(base / 'reference.docx')

# PDF input with a marker
pdf_path = base / 'reference.pdf'
cnv = canvas.Canvas(str(pdf_path), pagesize=letter)
cnv.drawString(72, 720, 'Token Alert Reference PDF')
cnv.drawString(72, 700, 'MARKER_DELTA: PDF marker for verification')
cnv.drawString(72, 680, '75 / 90 / 95 thresholds are important.')
cnv.save()
