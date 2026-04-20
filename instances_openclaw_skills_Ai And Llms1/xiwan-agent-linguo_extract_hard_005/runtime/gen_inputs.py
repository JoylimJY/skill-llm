from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
from openpyxl import Workbook
import json

base = Path('.')

# marker text used across files
marker = 'MARKER-ALPHA-7QX'
second_marker = 'SPEC-DRIFT-BETA-19'

# txt file
(base / 'incident_note.txt').write_text(
    f"Protocol draft recovered from a mixed bundle.\nPrimary marker: {marker}\nSecondary marker: {second_marker}\nObserved note: agent-lingua handshake uses signature suffixes.\n",
    encoding='utf-8'
)

# json file
json_data = {
    'name': 'agent-lingua',
    'version': '0.4.0',
    'capabilities': ['1', '7', 'A'],
    'security': ['P', 'B', 'E'],
    'marker': marker,
    'note': 'protocol draft with a typo in canonical URL spelling'
}
(base / 'bundle.json').write_text(json.dumps(json_data, indent=2), encoding='utf-8')

# csv file
(base / 'matrix.csv').write_text(
    'section,value,marker\n'
    f'domain,MSG,{marker}\n'
    f'action,HSK,{second_marker}\n'
    'security,E,encrypted session\n',
    encoding='utf-8'
)

# docx file
_doc = Document()
_doc.add_heading('Agent Lingua Recovery', level=1)
_doc.add_paragraph(f'This document mentions {marker} and emphasizes handshake negotiation.')
_doc.add_paragraph('A source typo appears in the canonical URL string.')
_doc.save(base / 'notes.docx')

# xlsx file
wb = Workbook()
ws = wb.active
ws.title = 'evidence'
ws['A1'] = 'field'
ws['B1'] = 'value'
ws['A2'] = 'marker'
ws['B2'] = marker
ws['A3'] = 'secondary'
ws['B3'] = second_marker
wb.save(base / 'evidence.xlsx')

# pdf file
pdf_path = base / 'scan.pdf'
c = canvas.Canvas(str(pdf_path), pagesize=letter)
c.drawString(72, 720, f'Incident scan: {marker}')
c.drawString(72, 700, f'Handshake path references {second_marker}')
c.drawString(72, 680, 'Canonical spec appears once with a misspelled path component.')
c.showPage()
c.save()

# metadata file for deterministic checks
(base / 'expected_markers.txt').write_text(
    f'{marker}\n{second_marker}\n',
    encoding='utf-8'
)
