from pathlib import Path
from openpyxl import Workbook
from docx import Document
from PyPDF2 import PdfWriter

base = Path('.')
# Text brief with marker content
(base / 'brief.txt').write_text(
    'MARKER::REVENUEMODEL::2025\n'
    'Business: solo founder building a niche SaaS for freelance designers.\n'
    'Audience: freelance designers and small design studios.\n'
    'Pain: recurring client work is chaotic; need predictable monthly cash flow.\n'
    'Existing assets: newsletter, template library, and occasional consulting calls.\n'
    'Constraints: low churn is important; international customers expected.\n',
    encoding='utf-8'
)

# CSV-like data in xlsx for projection markers
wb = Workbook()
ws = wb.active
ws.title = 'inputs'
ws.append(['marker', 'value'])
ws.append(['MARKER_CUSTOMERS_M1', 35])
ws.append(['MARKER_GROWTH_RATE', 0.18])
ws.append(['MARKER_ARPC', 24])
ws.append(['MARKER_CHURN', 0.05])
ws.append(['MARKER_SECONDARY_REVENUE', 1200])
ws.append(['MARKER_OPP_REV', 400])
wb.save(base / 'metrics.xlsx')

# DOCX with a note
 doc = Document()
doc.add_heading('Founder Notes', level=1)
doc.add_paragraph('MARKER::STACK::SUBSCRIPTION+COURSE+CONSULTING')
doc.add_paragraph('The founder wants predictable monthly revenue and wants to avoid marketplace complexity.')
doc.save(base / 'notes.docx')

# Tiny PDF marker
writer = PdfWriter()
writer.add_blank_page(width=72, height=72)
writer.add_metadata({'/Title': 'MARKER_PDF_REVENUE'})
with open(base / 'attachment.pdf', 'wb') as f:
    writer.write(f)
