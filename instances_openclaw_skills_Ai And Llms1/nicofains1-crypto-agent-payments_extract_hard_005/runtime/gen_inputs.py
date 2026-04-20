from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
import json

base = Path('.')
(base / 'data').mkdir(exist_ok=True)

# Deterministic marker contents
markers = {
    'primary_id': 'ALPHA-7842',
    'secondary_id': 'BETA-1930',
    'code_phrase': 'ORANGE LANTERN',
    'checksum': 'M7Q-55-X9',
}

# text file with decoys
(base / 'data' / 'notes.txt').write_text(
    'Record A: BETA-1930\n'
    'Record B: ALPHA-7842\n'
    'Primary record is the one tagged PRIMARY.\n'
    'PRIMARY => ALPHA-7842 | ORANGE LANTERN | M7Q-55-X9\n',
    encoding='utf-8'
)

# json file with structure
(base / 'data' / 'manifest.json').write_text(json.dumps({
    'records': [
        {'label': 'secondary', 'id': markers['secondary_id']},
        {'label': 'primary', 'id': markers['primary_id'], 'phrase': markers['code_phrase'], 'checksum': markers['checksum']}
    ]
}, indent=2), encoding='utf-8')

# docx with more decoys

doc = Document()
doc.add_paragraph('Internal memo')
doc.add_paragraph('Candidate: BETA-1930')
doc.add_paragraph('Primary record: ALPHA-7842')
doc.add_paragraph('Phrase: ORANGE LANTERN')
doc.add_paragraph('Checksum: M7Q-55-X9')
doc.save(base / 'data' / 'memo.docx')

# minimal pdf containing marker text in an accessible way
writer = PdfWriter()
page = writer.add_blank_page(width=300, height=300)
# Embed metadata for extraction-friendly marker content
writer.add_metadata({
    '/Title': 'Primary Record',
    '/Subject': 'ALPHA-7842 | ORANGE LANTERN | M7Q-55-X9'
})
with open(base / 'data' / 'brief.pdf', 'wb') as f:
    writer.write(f)

# spreadsheet with rows of markers
wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'
ws.append(['type', 'value'])
ws.append(['secondary', markers['secondary_id']])
ws.append(['primary_id', markers['primary_id']])
ws.append(['phrase', markers['code_phrase']])
ws.append(['checksum', markers['checksum']])
wb.save(base / 'data' / 'table.xlsx')
