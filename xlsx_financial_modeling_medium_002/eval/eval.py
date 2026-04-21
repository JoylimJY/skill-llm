import json
import sys
import os
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font

def main():
    workspace = Path(sys.argv[1])
    checks = []
    
    # Check if file exists
    excel_file = workspace / 'startup_valuation.xlsx'
    if not excel_file.exists():
        checks.append({"name": "file_exists", "passed": False, "detail": "startup_valuation.xlsx not found"})
        result = {"passed": False, "score": 0.0, "checks": checks}
        print(json.dumps(result))
        return
    
    checks.append({"name": "file_exists", "passed": True, "detail": "startup_valuation.xlsx found"})
    
    try:
        wb = load_workbook(excel_file, data_only=False)
        
        # Check worksheets exist
        required_sheets = ['assumptions', 'projections', 'valuation']
        sheet_names_lower = [name.lower() for name in wb.sheetnames]
        
        for sheet in required_sheets:
            if sheet in sheet_names_lower:
                checks.append({"name": f"sheet_{sheet}", "passed": True, "detail": f"{sheet.title()} sheet found"})
            else:
                checks.append({"name": f"sheet_{sheet}", "passed": False, "detail": f"{sheet.title()} sheet missing"})
        
        # Find actual sheet objects
        assumptions_sheet = None
        projections_sheet = None
        valuation_sheet = None
        
        for name in wb.sheetnames:
            if 'assumption' in name.lower():
                assumptions_sheet = wb[name]
            elif 'projection' in name.lower():
                projections_sheet = wb[name]
            elif 'valuation' in name.lower():
                valuation_sheet = wb[name]
        
        # Check for key content in assumptions
        has_growth_rates = False
        has_margins = False
        if assumptions_sheet:
            for row in assumptions_sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        text = cell.value.lower()
                        if any(term in text for term in ['growth', 'rate']):
                            has_growth_rates = True
                        if any(term in text for term in ['margin', 'ebitda']):
                            has_margins = True
        
        checks.append({"name": "assumptions_content", "passed": has_growth_rates and has_margins, "detail": f"Growth rates found: {has_growth_rates}, Margins found: {has_margins}"})
        
        # Check for formulas in projections
        formula_count = 0
        has_revenue_projection = False
        has_ebitda_projection = False
        
        if projections_sheet:
            for row in projections_sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        if isinstance(cell.value, str):
                            text = cell.value.lower()
                            if 'revenue' in text:
                                has_revenue_projection = True
                            if 'ebitda' in text:
                                has_ebitda_projection = True
                            if cell.value.startswith('='):
                                formula_count += 1
        
        checks.append({"name": "projections_content", "passed": has_revenue_projection and has_ebitda_projection, "detail": f"Revenue projection: {has_revenue_projection}, EBITDA projection: {has_ebitda_projection}"})
        checks.append({"name": "formulas_present", "passed": formula_count >= 10, "detail": f"Found {formula_count} formulas"})
        
        # Check for valuation content
        has_enterprise_value = False
        has_multiples = False
        
        if valuation_sheet:
            for row in valuation_sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        text = cell.value.lower()
                        if any(term in text for term in ['enterprise', 'ev']):
                            has_enterprise_value = True
                        if any(term in text for term in ['multiple', 'ev/']):
                            has_multiples = True
        
        checks.append({"name": "valuation_content", "passed": has_enterprise_value and has_multiples, "detail": f"Enterprise value: {has_enterprise_value}, Multiples: {has_multiples}"})
        
        # Check for proper color coding (look for blue inputs)
        has_blue_inputs = False
        if assumptions_sheet:
            for row in assumptions_sheet.iter_rows():
                for cell in row:
                    if cell.font and cell.font.color:
                        color_rgb = cell.font.color.rgb if hasattr(cell.font.color, 'rgb') else None
                        if color_rgb and ('0000ff' in str(color_rgb).lower() or 'blue' in str(color_rgb).lower()):
                            has_blue_inputs = True
                            break
                if has_blue_inputs:
                    break
        
        checks.append({"name": "color_coding", "passed": has_blue_inputs, "detail": f"Blue input formatting found: {has_blue_inputs}"})
        
        # Check for currency formatting indicators
        has_currency_format = False
        for sheet in wb.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        if any(term in cell.value.lower() for term in ['$mm', 'millions', '$m']):
                            has_currency_format = True
                            break
                if has_currency_format:
                    break
            if has_currency_format:
                break
        
        checks.append({"name": "currency_formatting", "passed": has_currency_format, "detail": f"Currency format indicators found: {has_currency_format}"})
        
        wb.close()
        
    except Exception as e:
        checks.append({"name": "file_processing", "passed": False, "detail": f"Error processing file: {str(e)}"})
    
    # Calculate score
    passed_count = sum(1 for check in checks if check['passed'])
    total_checks = len(checks)
    score = passed_count / total_checks if total_checks > 0 else 0.0
    passed = score >= 0.8
    
    result = {
        "passed": passed,
        "score": score,
        "checks": checks
    }
    
    print(json.dumps(result))

if __name__ == '__main__':
    main()