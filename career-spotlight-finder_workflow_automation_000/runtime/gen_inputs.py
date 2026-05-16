#!/usr/bin/env python3
"""
Generates the workspace for the career-spotlight-finder evaluation.
Creates:
1. Project source materials (biotech/bioinformatics domain)
2. A pre-existing ~/.career-spotlight/ with stale state (to force archival)
3. The skill's guide and template files at their documented paths
4. Distractor files to test contextual awareness
"""

import os
import random
import json
from pathlib import Path
from datetime import datetime, timedelta
import subprocess

random.seed(42)

WORKSPACE = Path("/workspace")

# ─────────────────────────────────────────────
# 1. Create skill infrastructure (guides + templates)
#    These "already exist in the workspace" per SKILL.md
# ─────────────────────────────────────────────

guides_dir = WORKSPACE / "guides"
templates_dir = WORKSPACE / "templates"
guides_dir.mkdir(parents=True, exist_ok=True)
templates_dir.mkdir(parents=True, exist_ok=True)

# ── input-collection-guide.md ──
(guides_dir / "input-collection-guide.md").write_text("""# Input Collection Guide

## Section 1 — Init

Create the following directory structure if it does not exist:

```
~/.career-spotlight/
├── analyses/
├── copies/
└── history/
```

Write a config file at `~/.career-spotlight/config.json` with:
```json
{
  "initialized": true,
  "version": "1.0.0"
}
```

## Section 2 — Source Collection

Ask the user for project sources. Accepted types:
- Local file paths (`.md`, `.txt`, `.py`, `.docx`)
- Local directory paths (auto-expand to all readable files within)
- URLs (fetch content via HTTP)

## Section 3 — Source Validation

For each source:
- Verify it exists (or is fetchable)
- For `.docx` files, convert to plain text using `pandoc` before analysis
- For directories, recursively collect all `.md`, `.txt`, `.py`, `.ipynb` files

## Section 4 — Priority Assignment

For each project source group, ask the user:
> "Should this project be `highlight` (primary) or `supporting` (secondary)?"

Record priority in the analysis front-matter.

## Section 5 — Staleness Check

For each source that has a prior analysis in `~/.career-spotlight/analyses/`:
- If source is a local file: compare current `mtime` against the `analyzed_at` timestamp in the analysis front-matter
- If source is a git repo: compare current `HEAD` hash against `git_hash` in front-matter
- If stale (source newer than analysis), mark for re-analysis

## Section 6 — Analysis Execution

For each source marked for (re-)analysis, run the analysis per `guides/project-analysis-guide.md`.

## Section 7 — Write Analyses

Write each analysis to `~/.career-spotlight/analyses/<project-slug>.md`.

The filename slug must be derived from the project name: lowercase, spaces→hyphens, strip special chars.

## Section 8 — Confirmation

List all analyses written. Ask the user to confirm before proceeding to Step 2.
""")

# ── project-analysis-guide.md ──
(guides_dir / "project-analysis-guide.md").write_text("""# Project Analysis Guide

For each project source, produce a structured analysis in Markdown with the following front-matter and sections.

## Front-matter (YAML block at top)

```yaml
---
project: <Project Name>
slug: <project-slug>
priority: <highlight|supporting>
analyzed_at: <ISO8601 timestamp>
source_path: <original source path or URL>
git_hash: <SHA if git repo, else null>
mtime: <file mtime ISO8601 if local file, else null>
---
```

## Required Sections

### 1. Project Overview
2–4 sentence summary of what the project does/did.

### 2. Technical Contributions
Bullet list of concrete technical contributions (tools, algorithms, architectures used).

### 3. Hidden Strengths
3–5 transferable skills or capabilities that are non-obvious from the project title alone.

### 4. Industry Buzzwords
10–15 keywords relevant to the domain and modern hiring filters.

### 5. Impact Signals
Quantified or qualifiable outcomes (performance improvements, scale, business value).

### 6. Career Narrative Hook
One sentence framing this project as a career differentiator.
""")

# ── domain-positioning-guide.md ──
(guides_dir / "domain-positioning-guide.md").write_text("""# Domain Positioning Guide

## Section 1 — Purpose
Identify the single most compelling expert framing for the candidate's career brand.

## Section 2 — Evidence Aggregation
Review all analyses. Count frequency of:
- Technical domains (e.g., ML, genomics, cloud infra)
- Role signals (e.g., research, engineering, product)
- Seniority indicators (e.g., led, architected, designed)

## Section 3 — Framing Candidates
Generate 2–3 positioning candidates, e.g.:
- "Computational Biologist bridging wet-lab and ML pipelines"
- "Bioinformatics Engineer specializing in scalable genomics data infrastructure"
- "Research Software Engineer in precision medicine"

## Section 4 — Distinctiveness Thesis
For the recommended framing, write a one-sentence distinctiveness thesis:
> "Unlike most [role], [candidate] uniquely combines [X] with [Y] to deliver [Z]."

Present the top framing with its thesis. Keep alternatives as wrappers.
Ask the user to confirm before proceeding.
""")

# ── narrative-synthesis-guide.md ──
(guides_dir / "narrative-synthesis-guide.md").write_text("""# Narrative Synthesis Guide

## Methodology

1. Read all project analyses from `~/.career-spotlight/analyses/`.
2. Cluster hidden strengths across projects into 3–5 meta-themes.
3. Identify the top 3 "hidden capabilities" — skills that appear across multiple projects but are not in job titles.
4. Extract the 20 most recurring industry buzzwords across all analyses.
5. Build a career arc narrative: early work → inflection point → current focus.
6. Write the report using `templates/aggregated-report.md` as the structural template.

## Aggregation Rules

- `highlight` projects contribute 2× weight to theme clustering
- `supporting` projects contribute 1× weight
- Buzzwords appearing in ≥2 projects get promoted to "Core Keywords"
""")

# ── copywriting-guide.md ──
(guides_dir / "copywriting-guide.md").write_text("""# Copywriting Guide

## Methodology

Read `~/.career-spotlight/report.md` and produce four distinct copy variants.

### resume-bullets.md
- 6–10 bullet points starting with strong action verbs
- Each bullet: action verb + technical context + quantified/qualifiable outcome
- Integrate core keywords naturally

### elevator-pitch.md
- 3–5 sentences, spoken-word register
- Opens with the distinctiveness thesis
- Ends with a clear value proposition

### linkedin-summary.md
- 150–300 words
- First-person, professional but personable
- Structured: hook → expertise → proof points → call to action

### casual-intro.md
- 2–3 sentences, conversational register
- Avoids jargon
- Answers "What do you do?" at a dinner party

## Quality Rules
- Never repeat the exact same phrase across two variants
- Each variant must reference at least one specific project by name
- Buzzwords from the report's "Core Keywords" must appear in at least 3 of the 4 variants
""")

# ── templates/project-analysis.md ──
(templates_dir / "project-analysis.md").write_text("""---
project: PROJECT_NAME
slug: project-slug
priority: highlight
analyzed_at: TIMESTAMP
source_path: SOURCE_PATH
git_hash: null
mtime: null
---

## Project Overview

[2–4 sentence summary]

## Technical Contributions

- [contribution 1]
- [contribution 2]

## Hidden Strengths

- [strength 1]
- [strength 2]
- [strength 3]

## Industry Buzzwords

[keyword1, keyword2, keyword3, ...]

## Impact Signals

- [impact 1]

## Career Narrative Hook

[One differentiator sentence]
""")

# ── templates/aggregated-report.md ──
(templates_dir / "aggregated-report.md").write_text("""# Career Brand Report

**Generated:** TIMESTAMP
**Positioning:** POSITIONING_STATEMENT

---

## Distinctiveness Thesis

DISTINCTIVENESS_THESIS

---

## Meta-Themes

### Theme 1: THEME_NAME
THEME_DESCRIPTION

### Theme 2: THEME_NAME
THEME_DESCRIPTION

### Theme 3: THEME_NAME
THEME_DESCRIPTION

---

## Top 3 Hidden Capabilities

1. HIDDEN_CAP_1
2. HIDDEN_CAP_2
3. HIDDEN_CAP_3

---

## Core Keywords

KEYWORD_LIST

---

## Career Arc Narrative

NARRATIVE

---

## Project Highlights

### highlight Projects
PROJECT_HIGHLIGHTS

### supporting Projects
PROJECT_SUPPORTING
""")

# ── templates/copywriting-variants.md ──
(templates_dir / "copywriting-variants.md").write_text("""# Copywriting Variants Template

## resume-bullets.md
```
- [Action Verb] [technical context] resulting in [outcome]
```

## elevator-pitch.md
```
[Distinctiveness thesis]. [2–3 sentences of proof]. [Value proposition].
```

## linkedin-summary.md
```
[Hook sentence]. [150–300 words of expertise + proof + CTA]
```

## casual-intro.md
```
[2–3 conversational sentences answering "What do you do?"]
```
""")

# ─────────────────────────────────────────────
# 2. Create realistic project source materials
#    (biotech / bioinformatics domain)
# ─────────────────────────────────────────────

projects_dir = WORKSPACE / "my_projects"
projects_dir.mkdir(exist_ok=True)

# ── Project A: Genomics Pipeline (Python code) ──
proj_a = projects_dir / "variant-calling-pipeline"
proj_a.mkdir(exist_ok=True)

(proj_a / "README.md").write_text("""# Variant Calling Pipeline

A high-throughput somatic variant calling pipeline built with Snakemake and GATK4.
Processes whole-exome sequencing (WES) data from raw FASTQ to annotated VCF.

## Architecture
- Snakemake workflow engine for parallelization
- GATK4 HaplotypeCaller + MuTect2 for variant detection
- ANNOVAR for functional annotation
- Runs on SLURM HPC clusters

## Performance
- Processes 200-sample cohorts in <6 hours on a 32-node cluster
- 98.7% concordance with GIAB benchmark variants
""")

(proj_a / "pipeline.py").write_text("""#!/usr/bin/env python3
\"\"\"
Somatic variant calling orchestrator.
Wraps GATK4 MuTect2 with tumor-normal matched pairs.
\"\"\"
import subprocess
import pathlib
import argparse
import json

def run_mutect2(tumor_bam, normal_bam, ref_genome, output_vcf, intervals=None):
    cmd = [
        "gatk", "Mutect2",
        "-I", tumor_bam, "--tumor-sample", "TUMOR",
        "-I", normal_bam, "--normal-sample", "NORMAL",
        "-R", ref_genome,
        "-O", output_vcf,
    ]
    if intervals:
        cmd += ["-L", intervals]
    subprocess.run(cmd, check=True)
    return output_vcf

def filter_variants(raw_vcf, stats_file, output_vcf):
    cmd = [
        "gatk", "FilterMutectCalls",
        "-V", raw_vcf,
        "--stats", stats_file,
        "-O", output_vcf,
    ]
    subprocess.run(cmd, check=True)

def annotate(vcf, output_dir):
    cmd = [
        "perl", "annovar/table_annovar.pl", vcf,
        "humandb/", "-buildver", "hg38",
        "-out", str(output_dir / "annotated"),
        "-protocol", "refGene,cosmic70",
        "-operation", "g,f",
        "-nastring", ".",
    ]
    subprocess.run(cmd, check=True)

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--tumor", required=True)
    parser.add_argument("--normal", required=True)
    parser.add_argument("--ref", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    raw = run_mutect2(args.tumor, args.normal, args.ref, args.out + ".raw.vcf")
    filter_variants(raw, args.out + ".stats", args.out + ".filtered.vcf")
    annotate(pathlib.Path(args.out + ".filtered.vcf"), pathlib.Path(args.out))
""")

(proj_a / "Snakefile").write_text("""# Snakemake workflow for variant calling
configfile: "config.yaml"

rule all:
    input: expand("results/{sample}.annotated.vcf", sample=config["samples"])

rule bwa_mem:
    input: r1="data/{sample}_R1.fastq.gz", r2="data/{sample}_R2.fastq.gz"
    output: "aligned/{sample}.bam"
    shell: "bwa mem -t 8 {config[ref]} {input.r1} {input.r2} | samtools sort -o {output}"

rule mutect2:
    input: tumor="aligned/{sample}_tumor.bam", normal="aligned/{sample}_normal.bam"
    output: "variants/{sample}.vcf"
    shell: "gatk Mutect2 -I {input.tumor} -I {input.normal} -R {config[ref]} -O {output}"
""")

(proj_a / "config.yaml").write_text("""ref: /data/hg38/GRCh38.fa
samples:
  - TCGA_LUAD_001
  - TCGA_LUAD_002
  - TCGA_BRCA_001
threads: 8
""")

(proj_a / "benchmarks.txt").write_text("""Sample          Precision  Recall  F1
GIAB_NA12878    0.9991     0.9823  0.9906
GIAB_NA24385    0.9987     0.9811  0.9898
TCGA_LUAD_001   0.9934     0.9756  0.9844
""")

# ── Project B: CRISPR Off-Target Analysis Articles ──
proj_b = projects_dir / "crispr-offtarget-analysis"
proj_b.mkdir(exist_ok=True)

(proj_b / "method_description.md").write_text("""# CRISPR Off-Target Detection: Computational Approach

## Background
Guide RNA (gRNA) off-target effects remain a major safety concern for therapeutic CRISPR applications.
Existing tools (Cas-OFFinder, CRISPOR) rely on mismatch counting but miss complex structural variants.

## Our Approach
We developed a graph-based alignment scoring model that accounts for:
- DNA bulge positions
- RNA secondary structure of the gRNA
- Chromatin accessibility (ATAC-seq signal integration)

## Implementation
- Python 3.10, NetworkX for graph construction
- Integration with ENCODE chromatin data via REST API
- Benchmarked on 47 published gRNA datasets
- 23% improvement in specificity over Cas-OFFinder baseline

## Clinical Relevance
Deployed as an internal screening tool at a gene therapy startup.
Reduced false-positive off-target calls by 41%, accelerating IND-enabling studies.
""")

(proj_b / "analysis_notes.txt").write_text("""Off-target analysis run log:
- Dataset: Listgarten 2018 benchmark (47 gRNAs, GUIDE-seq validated)
- Model: graph-based alignment + chromatin accessibility
- Metric: AUROC = 0.93 vs 0.81 baseline (Cas-OFFinder)
- Runtime: 8 min per gRNA on 16-core workstation
- Dependencies: Python 3.10, NetworkX 3.1, pysam 0.21

Key finding: bulge positions at seed region (positions 1-12) 
most predictive of high-risk off-targets.
""")

(proj_b / "presentation_abstract.md").write_text("""## Abstract — CRISPR Symposium 2023

Title: Graph-Based Off-Target Prediction Integrating Chromatin Context

We present a novel computational framework for CRISPR off-target prediction that
models gRNA-DNA interactions as weighted graphs, incorporating:
1. Sequence mismatch penalties with position-dependent weights
2. DNA/RNA bulge structural penalties derived from MD simulation data
3. Chromatin accessibility scores from cell-type-specific ATAC-seq

Evaluated on 847 GUIDE-seq validated off-target sites across 47 gRNAs,
our method achieves AUROC 0.93 (±0.02) versus 0.81 for the best prior tool.

Clinical deployment at BioEdit Therapeutics reduced off-target false-positives
by 41%, compressing regulatory review timelines by an estimated 3 months.
""")

# ── Project C: Lab Notebook (.docx) — will be converted via pandoc ──
proj_c = projects_dir / "protein-folding-ml"
proj_c.mkdir(exist_ok=True)

# Create .docx programmatically using python-docx
docx_content = """
Project: ML-Assisted Protein Folding for Rare Disease Targets

Summary:
I led a 6-month project to fine-tune AlphaFold2 on rare disease protein targets
where PDB coverage is sparse. We curated a training set of 1,200 homology-modeled
structures from SwissProt + manual curation, then fine-tuned the EvoFormer trunk
using LoRA adapters to preserve base model generalization.

Technical Work:
- Fine-tuned AlphaFold2 EvoFormer with LoRA (rank=16) on 1,200 rare disease proteins
- Built a data curation pipeline: UniProt API → structure quality filtering → MSA generation
- Achieved TM-score improvement of +0.08 over base AlphaFold2 on held-out test set (n=120)
- Deployed as REST microservice on internal Kubernetes cluster
- Integrated with medicinal chemistry team's docking workflow (AutoDock Vina)

Impact:
- 3 drug discovery programs unblocked (FKRP, DYSF, COL6A3 targets)
- Reduced structure prediction turnaround from 2 weeks to 4 hours
- Co-authored internal technical report, cited in 2 patent applications

Technologies:
Python, PyTorch, JAX, AlphaFold2, Kubernetes, Docker, LoRA, UniProt API, AutoDock Vina
"""

try:
    from docx import Document
    doc = Document()
    doc.add_heading("ML-Assisted Protein Folding for Rare Disease Targets", 0)
    for para in docx_content.strip().split('\n\n'):
        if para.strip():
            if para.strip().endswith(':') or para.strip().startswith('Project:') or para.strip().startswith('Summary:'):
                doc.add_heading(para.strip(), level=2)
            else:
                doc.add_paragraph(para.strip())
    doc.save(str(proj_c / "lab_notebook.docx"))
    print("Created lab_notebook.docx")
except Exception as e:
    print(f"Warning: could not create .docx ({e}), creating plain text fallback")
    (proj_c / "lab_notebook.txt").write_text(docx_content)

(proj_c / "experiment_log.md").write_text("""# Experiment Log — Protein Folding ML

## 2024-01-15
- Baseline TM-score (AlphaFold2 base): 0.71 on 120 test proteins
- After LoRA fine-tuning (rank=16, lr=1e-4): TM-score 0.79

## 2024-01-22
- Ablation: LoRA rank 8 → TM-score 0.77, rank 32 → 0.79 (diminishing returns)
- Selected rank=16 as production config

## 2024-02-01
- Deployed to Kubernetes (3 replicas, GPU-enabled nodes)
- Latency: 4 hours end-to-end vs 2 weeks manual pipeline
""")

# ─────────────────────────────────────────────
# 3. Pre-existing stale ~/.career-spotlight/ state
#    Forces the agent to: check staleness, archive report, archive copies
# ─────────────────────────────────────────────

career_dir = Path.home() / ".career-spotlight"
(career_dir / "analyses").mkdir(parents=True, exist_ok=True)
(career_dir / "copies").mkdir(parents=True, exist_ok=True)
(career_dir / "history").mkdir(parents=True, exist_ok=True)

# Stale analysis (old timestamp, predates the source files)
old_ts = (datetime.now() - timedelta(days=90)).isoformat()

(career_dir / "analyses" / "variant-calling-pipeline.md").write_text(f"""---
project: Variant Calling Pipeline
slug: variant-calling-pipeline
priority: highlight
analyzed_at: {old_ts}
source_path: /workspace/my_projects/variant-calling-pipeline
git_hash: null
mtime: {old_ts}
---

## Project Overview

An old analysis that is now stale and should be regenerated.

## Technical Contributions

- GATK4 usage

## Hidden Strengths

- Pipeline engineering

## Industry Buzzwords

bioinformatics, NGS, variant calling

## Impact Signals

- Processed samples

## Career Narrative Hook

Old hook that needs refreshing.
""")

# Old report.md (must be archived before writing new one)
(career_dir / "report.md").write_text("""# Career Brand Report

**Generated:** 2024-01-01T00:00:00
**Positioning:** OUTDATED - Generic Bioinformatics Engineer

This is an old report that must be archived to history/ before a new one is written.

## Distinctiveness Thesis

An outdated thesis that no longer reflects current work.
""")

# Old copy files (must be archived before writing new ones)
(career_dir / "copies" / "resume-bullets.md").write_text("""# Resume Bullets (OLD)

- Developed bioinformatics pipelines
- Analyzed genomic data
""")

(career_dir / "copies" / "elevator-pitch.md").write_text("""# Elevator Pitch (OLD)

I work in bioinformatics. I analyze genetic data.
""")

(career_dir / "copies" / "linkedin-summary.md").write_text("""# LinkedIn Summary (OLD)

Bioinformatics professional with experience in genomics.
""")

(career_dir / "copies" / "casual-intro.md").write_text("""# Casual Intro (OLD)

I work with DNA data on computers.
""")

# ─────────────────────────────────────────────
# 4. Distractor files (noise to test focus)
# ─────────────────────────────────────────────

distractor_dir = WORKSPACE / "misc"
distractor_dir.mkdir(exist_ok=True)

(distractor_dir / "old_cv.txt").write_text("""John Smith
PhD Computational Biology, MIT 2019
Skills: Python, R, Bash, GATK, BWA
""")

(distractor_dir / "meeting_notes.txt").write_text("""2024-03-10 meeting notes
- Discussed Q2 hiring plan
- Need to update resume before conference
- Ask HR about LinkedIn profile review
""")

(distractor_dir / "random_ideas.md").write_text("""Ideas for side projects:
- Build a drug-drug interaction predictor
- Write blog post about AlphaFold
- Contribute to Biopython
""")

backup_dir = WORKSPACE / "backups" / "2023"
backup_dir.mkdir(parents=True, exist_ok=True)
(backup_dir / "old_resume_v1.pdf.txt").write_text("placeholder for old resume binary")
(backup_dir / "cover_letter_draft.txt").write_text("Dear Hiring Manager, I am writing to...")
(backup_dir / "references.txt").write_text("Prof. Jane Doe, MIT\nDr. Bob Chen, Broad Institute")

papers_dir = WORKSPACE / "papers"
papers_dir.mkdir(exist_ok=True)
(papers_dir / "crispr_review_2022.md").write_text("""# CRISPR Therapeutic Applications Review

A survey of off-target detection methods...
This is a background reading file, not a project source.
""")
(papers_dir / "alphafold_notes.txt").write_text("Personal notes on AlphaFold2 architecture...")
(papers_dir / "TODO.txt").write_text("TODO: Read Jumper et al. 2021 supplementary\n")

temp_dir = WORKSPACE / "tmp"
temp_dir.mkdir(exist_ok=True)
(temp_dir / "scratch.py").write_text("# scratch work\nx = [1,2,3]\nprint(x)")
(temp_dir / "data_test.csv").write_text("sample,value\nA,1\nB,2\n")

print("Workspace generation complete.")
print(f"Projects created: {list(projects_dir.iterdir())}")
print(f"Career dir pre-populated: {list(career_dir.iterdir())}")