import sys
import os
import json
from pathlib import Path
from openpyxl import load_workbook

def evaluate_budget_file(workspace_dir):
    checks = []
    workspace_path = Path(workspace_dir)
    
    # Check 1: File exists with correct name
    target_file = workspace_path / 'monthly_budget.xlsx'
    if target_file.exists():
        checks.append({"name": "File exists", "passed": True, "detail": "monthly_budget.xlsx found"})
    else:
        checks.append({"name": "File exists", "passed": False, "detail": "monthly_budget.xlsx not found"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    try:
        wb = load_workbook(str(target_file), data_only=True)
        
        # Check 2: Sheet name is 'Budget'
        if 'Budget' in wb.sheetnames:
            checks.append({"name": "Sheet name", "passed": True, "detail": "Sheet named 'Budget' found"})
            sheet = wb['Budget']
        else:
            checks.append({"name": "Sheet name", "passed": False, "detail": f"Sheet 'Budget' not found. Found: {wb.sheetnames}"})
            sheet = wb.active
        
        # Check 3: Headers in row 1
        expected_headers = ['category', 'budgeted', 'actual', 'difference']
        actual_headers = [str(sheet.cell(1, i).value or '').lower().strip() for i in range(1, 5)]
        headers_match = all(expected in ' '.join(actual_headers) for expected in expected_headers)
        if headers_match:
            checks.append({"name": "Headers present", "passed": True, "detail": "All required headers found"})
        else:
            checks.append({"name": "Headers present", "passed": False, "detail": f"Headers mismatch. Expected: {expected_headers}, Got: {actual_headers}"})
        
        # Check 4: Expense categories in column A
        expected_categories = ['rent', 'groceries', 'utilities', 'transportation', 'entertainment']
        actual_categories = [str(sheet.cell(i, 1).value or '').lower().strip() for i in range(2, 7)]
        categories_match = all(cat in actual_categories for cat in expected_categories)
        if categories_match:
            checks.append({"name": "Categories present", "passed": True, "detail": "All expense categories found"})
        else:
            checks.append({"name": "Categories present", "passed": False, "detail": f"Categories mismatch. Expected: {expected_categories}, Got: {actual_categories}"})
        
        # Check 5: Budgeted amounts in column B
        expected_budgeted = [1200, 400, 150, 300, 200]
        actual_budgeted = []
        for i in range(2, 7):
            val = sheet.cell(i, 2).value
            if isinstance(val, (int, float)):
                actual_budgeted.append(val)
            else:
                actual_budgeted.append(0)
        budgeted_match = actual_budgeted == expected_budgeted
        if budgeted_match:
            checks.append({"name": "Budgeted amounts", "passed": True, "detail": "Budgeted amounts correct"})
        else:
            checks.append({"name": "Budgeted amounts", "passed": False, "detail": f"Budgeted amounts mismatch. Expected: {expected_budgeted}, Got: {actual_budgeted}"})
        
        # Check 6: Actual amounts in column C
        expected_actual = [1200, 380, 165, 285, 250]
        actual_actual = []
        for i in range(2, 7):
            val = sheet.cell(i, 3).value
            if isinstance(val, (int, float)):
                actual_actual.append(val)
            else:
                actual_actual.append(0)
        actual_match = actual_actual == expected_actual
        if actual_match:
            checks.append({"name": "Actual amounts", "passed": True, "detail": "Actual amounts correct"})
        else:
            checks.append({"name": "Actual amounts", "passed": False, "detail": f"Actual amounts mismatch. Expected: {expected_actual}, Got: {actual_actual}"})
        
        # Check 7: Formulas in column D (difference)
        wb_formulas = load_workbook(str(target_file), data_only=False)
        sheet_formulas = wb_formulas['Budget'] if 'Budget' in wb_formulas.sheetnames else wb_formulas.active
        
        formula_count = 0
        for i in range(2, 7):
            cell_val = sheet_formulas.cell(i, 4).value
            if isinstance(cell_val, str) and cell_val.startswith('='):
                formula_count += 1
        
        if formula_count >= 3:
            checks.append({"name": "Difference formulas", "passed": True, "detail": f"Found {formula_count} formulas in difference column"})
        else:
            checks.append({"name": "Difference formulas", "passed": False, "detail": f"Expected formulas in difference column, found only {formula_count}"})
        
        # Check 8: Total row with SUM formulas
        total_formulas = 0
        for col in [2, 3, 4]:
            cell_val = sheet_formulas.cell(7, col).value
            if isinstance(cell_val, str) and 'sum' in cell_val.lower():
                total_formulas += 1
        
        if total_formulas >= 2:
            checks.append({"name": "Total formulas", "passed": True, "detail": f"Found {total_formulas} SUM formulas in total row"})
        else:
            checks.append({"name": "Total formulas", "passed": False, "detail": f"Expected SUM formulas in total row, found only {total_formulas}"})
        
        wb.close()
        wb_formulas.close()
        
    except Exception as e:
        checks.append({"name": "File processing", "passed": False, "detail": f"Error processing file: {str(e)}"})
    
    score = sum(1 for c in checks if c['passed']) / len(checks)
    passed = score >= 0.8
    
    return {"passed": passed, "score": score, "checks": checks}

if __name__ == "__main__":
    result = evaluate_budget_file(sys.argv[1])
    print(json.dumps(result))