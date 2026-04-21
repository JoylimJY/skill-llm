import os
from pptx import Presentation
from pptx.util import Inches, Pt

# Create a simple presentation with basic content
prs = Presentation()

# Title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Sample Presentation"
subtitle.text = "A test presentation for brand styling"

# Content slide
content_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Main Content"
content.text = "This is the main content of our presentation\n• First point\n• Second point\n• Third point"

# Add a text box with specific content for verification
textbox = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(8), Inches(1))
tf = textbox.text_frame
tf.text = "BRAND_MARKER_TEXT"

# Save the presentation
prs.save('input_presentation.pptx')
print("Generated input_presentation.pptx")