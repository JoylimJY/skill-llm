import sys
import os
import json
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_presentation_styling(workspace_dir):
    checks = []
    
    # Check if output file exists
    output_file = os.path.join(workspace_dir, 'styled_presentation.pptx')
    if not os.path.exists(output_file):
        checks.append({"name": "output_file_exists", "passed": False, "detail": "styled_presentation.pptx not found"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    checks.append({"name": "output_file_exists", "passed": True, "detail": "styled_presentation.pptx found"})
    
    try:
        prs = Presentation(output_file)
        
        # Check if presentation can be opened
        checks.append({"name": "presentation_readable", "passed": True, "detail": "Presentation can be opened successfully"})
        
        # Check if slides exist
        if len(prs.slides) >= 2:
            checks.append({"name": "slides_preserved", "passed": True, "detail": f"Found {len(prs.slides)} slides as expected"})
        else:
            checks.append({"name": "slides_preserved", "passed": False, "detail": f"Expected at least 2 slides, found {len(prs.slides)}"})
        
        # Check for brand colors in shapes
        brand_colors_found = False
        anthropic_colors = [
            (20, 20, 19),    # Dark #141413
            (250, 249, 245), # Light #faf9f5
            (176, 174, 165), # Mid Gray #b0aea5
            (232, 230, 220), # Light Gray #e8e6dc
            (217, 119, 87),  # Orange #d97757
            (106, 155, 204), # Blue #6a9bcc
            (120, 140, 93)   # Green #788c5d
        ]
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'fill') and shape.fill.type == 1:  # solid fill
                    try:
                        color = shape.fill.fore_color.rgb
                        color_tuple = (color.red, color.green, color.blue)
                        if color_tuple in anthropic_colors:
                            brand_colors_found = True
                            break
                    except:
                        pass
            if brand_colors_found:
                break
        
        if brand_colors_found:
            checks.append({"name": "brand_colors_applied", "passed": True, "detail": "Anthropic brand colors detected in presentation"})
        else:
            checks.append({"name": "brand_colors_applied", "passed": False, "detail": "No Anthropic brand colors detected"})
        
        # Check if marker text is preserved
        marker_found = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    if 'brand_marker_text' in shape.text.lower():
                        marker_found = True
                        break
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        if 'brand_marker_text' in paragraph.text.lower():
                            marker_found = True
                            break
            if marker_found:
                break
        
        if marker_found:
            checks.append({"name": "content_preserved", "passed": True, "detail": "Original content markers preserved"})
        else:
            checks.append({"name": "content_preserved", "passed": False, "detail": "Original content markers not found"})
        
    except Exception as e:
        checks.append({"name": "presentation_readable", "passed": False, "detail": f"Error reading presentation: {str(e)}"})
    
    # Calculate score
    passed_count = sum(1 for check in checks if check['passed'])
    score = passed_count / len(checks)
    overall_passed = score >= 0.75
    
    return {
        "passed": overall_passed,
        "score": score,
        "checks": checks
    }

if __name__ == "__main__":
    workspace_dir = sys.argv[1]
    result = check_presentation_styling(workspace_dir)
    print(json.dumps(result))