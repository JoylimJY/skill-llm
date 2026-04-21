import sys
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
import json

def check_branded_presentation(workspace_dir):
    checks = []
    
    # Check if the presentation file exists
    pptx_path = os.path.join(workspace_dir, 'branded_presentation.pptx')
    if not os.path.exists(pptx_path):
        checks.append({"name": "file_exists", "passed": False, "detail": "branded_presentation.pptx not found"})
        return checks
    
    checks.append({"name": "file_exists", "passed": True, "detail": "branded_presentation.pptx found"})
    
    try:
        prs = Presentation(pptx_path)
        
        # Check slide count (at least 3 slides)
        slide_count = len(prs.slides)
        checks.append({"name": "slide_count", "passed": slide_count >= 3, "detail": f"Found {slide_count} slides (expected >= 3)"})
        
        # Define Anthropic brand colors
        brand_colors = {
            'dark': (20, 20, 19),      # #141413
            'light': (250, 249, 245),  # #faf9f5
            'mid_gray': (176, 174, 165), # #b0aea5
            'light_gray': (232, 230, 220), # #e8e6dc
            'orange': (217, 119, 87),  # #d97757
            'blue': (106, 155, 204),   # #6a9bcc
            'green': (120, 140, 93)    # #788c5d
        }
        
        # Check for brand colors usage
        brand_colors_found = set()
        typography_applied = False
        
        for slide in prs.slides:
            # Check slide background
            if hasattr(slide.background, 'fill') and slide.background.fill.type:
                try:
                    bg_color = slide.background.fill.fore_color.rgb
                    bg_rgb = (bg_color.red, bg_color.green, bg_color.blue)
                    for color_name, color_rgb in brand_colors.items():
                        if abs(bg_rgb[0] - color_rgb[0]) <= 5 and abs(bg_rgb[1] - color_rgb[1]) <= 5 and abs(bg_rgb[2] - color_rgb[2]) <= 5:
                            brand_colors_found.add(color_name)
                except:
                    pass
            
            # Check shapes and text for brand colors and fonts
            for shape in slide.shapes:
                # Check shape colors
                if hasattr(shape, 'fill') and shape.fill.type:
                    try:
                        shape_color = shape.fill.fore_color.rgb
                        shape_rgb = (shape_color.red, shape_color.green, shape_color.blue)
                        for color_name, color_rgb in brand_colors.items():
                            if abs(shape_rgb[0] - color_rgb[0]) <= 5 and abs(shape_rgb[1] - color_rgb[1]) <= 5 and abs(shape_rgb[2] - color_rgb[2]) <= 5:
                                brand_colors_found.add(color_name)
                    except:
                        pass
                
                # Check text formatting
                if hasattr(shape, 'text_frame') and shape.text_frame.text.strip():
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            # Check font
                            font_name = run.font.name
                            if font_name and ('poppins' in font_name.lower() or 'lora' in font_name.lower() or 'arial' in font_name.lower() or 'georgia' in font_name.lower()):
                                typography_applied = True
                            
                            # Check text color
                            try:
                                if run.font.color.type:
                                    text_color = run.font.color.rgb
                                    text_rgb = (text_color.red, text_color.green, text_color.blue)
                                    for color_name, color_rgb in brand_colors.items():
                                        if abs(text_rgb[0] - color_rgb[0]) <= 5 and abs(text_rgb[1] - color_rgb[1]) <= 5 and abs(text_rgb[2] - color_rgb[2]) <= 5:
                                            brand_colors_found.add(color_name)
                            except:
                                pass
        
        # Check brand colors usage
        brand_colors_used = len(brand_colors_found) >= 2
        checks.append({"name": "brand_colors", "passed": brand_colors_used, "detail": f"Found {len(brand_colors_found)} brand colors: {list(brand_colors_found)}"})
        
        # Check typography application
        checks.append({"name": "typography", "passed": typography_applied, "detail": "Typography guidelines applied" if typography_applied else "No brand fonts detected"})
        
        # Check for visual elements (shapes beyond text boxes)
        visual_elements_found = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if not hasattr(shape, 'text_frame') or not shape.text_frame.text.strip():
                    visual_elements_found = True
                    break
            if visual_elements_found:
                break
        
        checks.append({"name": "visual_elements", "passed": visual_elements_found, "detail": "Visual elements like shapes found" if visual_elements_found else "No non-text visual elements detected"})
        
    except Exception as e:
        checks.append({"name": "presentation_parsing", "passed": False, "detail": f"Error parsing presentation: {str(e)}"})
    
    return checks

if __name__ == '__main__':
    workspace_dir = sys.argv[1]
    checks = check_branded_presentation(workspace_dir)
    
    score = sum(1 for c in checks if c['passed']) / len(checks)
    passed = score >= 0.8
    
    result = {
        "passed": passed,
        "score": score,
        "checks": checks
    }
    
    print(json.dumps(result))