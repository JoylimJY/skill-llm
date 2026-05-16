import sys
import json
import traceback
from pathlib import Path

def rgb_close(actual_rgb, expected_hex, tolerance=2):
    """Check if an RGBColor matches expected hex within tolerance."""
    try:
        er = int(expected_hex[0:2], 16)
        eg = int(expected_hex[2:4], 16)
        eb = int(expected_hex[4:6], 16)
        return (abs(actual_rgb[0] - er) <= tolerance and
                abs(actual_rgb[1] - eg) <= tolerance and
                abs(actual_rgb[2] - eb) <= tolerance)
    except Exception:
        return False

def get_rgb_tuple(color):
    try:
        rgb = color.rgb
        return (rgb[0], rgb[1], rgb[2])
    except Exception:
        return None

def main():
    workspace = sys.argv[1]
    checks = []
    
    # Anthropic brand colors
    DARK = "141413"
    LIGHT = "faf9f5"
    MID_GRAY = "b0aea5"
    LIGHT_GRAY = "e8e6dc"
    ORANGE = "d97757"
    BLUE = "6a9bcc"
    GREEN = "788c5d"
    
    BRAND_BACKGROUNDS = [DARK, LIGHT, MID_GRAY, LIGHT_GRAY]
    ACCENT_COLORS = [ORANGE, BLUE, GREEN]
    HEADING_FONTS = ["Poppins", "Arial"]
    BODY_FONTS = ["Lora", "Georgia"]
    
    # Find the output file
    output_path = None
    try:
        candidates = list(Path(workspace).rglob("qbr_deck_branded.pptx"))
        if not candidates:
            # Also accept if agent renamed/saved in place with a branded suffix
            candidates = list(Path(workspace).rglob("*branded*.pptx"))
        if not candidates:
            candidates = list(Path(workspace).rglob("*brand*.pptx"))
        if candidates:
            output_path = candidates[0]
    except Exception as e:
        pass

    if output_path is None:
        checks.append({"name": "output_file_exists", "passed": False,
                        "detail": "Could not find qbr_deck_branded.pptx anywhere in workspace."})
        result = {"passed": False, "score": 0.0, "checks": checks}
        print(json.dumps(result))
        return

    checks.append({"name": "output_file_exists", "passed": True,
                    "detail": f"Found output at {output_path}"})

    try:
        from pptx import Presentation
        from pptx.util import Pt
        prs = Presentation(str(output_path))
    except Exception as e:
        checks.append({"name": "pptx_readable", "passed": False,
                        "detail": f"Could not open PPTX: {e}"})
        result = {"passed": False, "score": 0.0, "checks": checks}
        print(json.dumps(result))
        return

    checks.append({"name": "pptx_readable", "passed": True, "detail": "PPTX opened successfully."})

    # ---- CHECK: Slide backgrounds use brand colors ----
    bg_pass_count = 0
    bg_details = []
    for i, slide in enumerate(prs.slides):
        try:
            bg_fill = slide.background.fill
            bg_rgb = get_rgb_tuple(bg_fill.fore_color)
            if bg_rgb is not None:
                matched = any(rgb_close(bg_rgb, c) for c in BRAND_BACKGROUNDS)
                if matched:
                    bg_pass_count += 1
                    bg_details.append(f"Slide {i+1}: bg RGB{bg_rgb} OK")
                else:
                    bg_details.append(f"Slide {i+1}: bg RGB{bg_rgb} NOT brand color")
            else:
                bg_details.append(f"Slide {i+1}: bg color not readable")
        except Exception as ex:
            bg_details.append(f"Slide {i+1}: error - {ex}")

    bg_ok = bg_pass_count >= 2  # At least 2 of 3 slides must have brand bg
    checks.append({"name": "slide_backgrounds_brand_colors",
                   "passed": bg_ok,
                   "detail": f"{bg_pass_count}/3 slides have brand backgrounds. " + " | ".join(bg_details)})

    # ---- CHECK: Heading fonts (>=24pt) are Poppins or Arial ----
    heading_font_ok_count = 0
    heading_font_fail = []
    body_font_ok_count = 0
    body_font_fail = []
    
    heading_color_ok_count = 0
    body_color_ok_count = 0
    
    BRAND_ALL_COLORS = [DARK, LIGHT, MID_GRAY, LIGHT_GRAY, ORANGE, BLUE, GREEN]
    
    for slide_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    try:
                        size_pt = run.font.size
                        if size_pt is None:
                            continue
                        size_val = size_pt / 12700  # EMU to pt
                        font_name = run.font.name or ""
                        
                        if size_val >= 24:
                            # Should be Poppins or Arial
                            if any(hf.lower() in font_name.lower() for hf in HEADING_FONTS):
                                heading_font_ok_count += 1
                            else:
                                heading_font_fail.append(
                                    f"Slide {slide_idx+1}: '{run.text[:30]}' @ {size_val}pt has font '{font_name}' (expected Poppins/Arial)")
                            # Color check
                            try:
                                rgb = get_rgb_tuple(run.font.color)
                                if rgb and any(rgb_close(rgb, c) for c in BRAND_ALL_COLORS):
                                    heading_color_ok_count += 1
                            except Exception:
                                pass
                        else:
                            # Should be Lora or Georgia
                            if any(bf.lower() in font_name.lower() for bf in BODY_FONTS):
                                body_font_ok_count += 1
                            else:
                                body_font_fail.append(
                                    f"Slide {slide_idx+1}: '{run.text[:30]}' @ {size_val}pt has font '{font_name}' (expected Lora/Georgia)")
                            # Color check
                            try:
                                rgb = get_rgb_tuple(run.font.color)
                                if rgb and any(rgb_close(rgb, c) for c in BRAND_ALL_COLORS):
                                    body_color_ok_count += 1
                            except Exception:
                                pass
                    except Exception as ex:
                        pass

    heading_font_pass = (heading_font_ok_count >= 3) and (len(heading_font_fail) == 0)
    checks.append({
        "name": "heading_font_poppins_or_arial",
        "passed": heading_font_pass,
        "detail": f"{heading_font_ok_count} heading runs correct. Failures: {heading_font_fail[:5]}"
    })

    body_font_pass = (body_font_ok_count >= 3) and (len(body_font_fail) == 0)
    checks.append({
        "name": "body_font_lora_or_georgia",
        "passed": body_font_pass,
        "detail": f"{body_font_ok_count} body runs correct. Failures: {body_font_fail[:5]}"
    })

    # ---- CHECK: Text colors are brand colors ----
    text_color_total = heading_color_ok_count + body_color_ok_count
    text_color_pass = text_color_total >= 5
    checks.append({
        "name": "text_colors_are_brand_colors",
        "passed": text_color_pass,
        "detail": f"{text_color_total} text runs have brand colors (headings: {heading_color_ok_count}, body: {body_color_ok_count})"
    })

    # ---- CHECK: Non-text shapes use accent colors (orange, blue, green cycling) ----
    accent_rgb_tuples = []
    for c in ACCENT_COLORS:
        accent_rgb_tuples.append((int(c[0:2],16), int(c[2:4],16), int(c[4:6],16)))
    
    shape_accent_ok = 0
    shape_accent_fail = []
    shape_accent_colors_found = []
    
    for slide_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                continue
            # Non-text shape
            try:
                fill = shape.fill
                rgb = get_rgb_tuple(fill.fore_color)
                if rgb is not None:
                    if any(rgb_close(rgb, c) for c in ACCENT_COLORS):
                        shape_accent_ok += 1
                        shape_accent_colors_found.append(rgb)
                    else:
                        shape_accent_fail.append(f"Slide {slide_idx+1} shape: RGB{rgb} not accent color")
            except Exception as ex:
                pass

    accent_shape_pass = shape_accent_ok >= 3 and len(shape_accent_fail) == 0
    checks.append({
        "name": "non_text_shapes_use_accent_colors",
        "passed": accent_shape_pass,
        "detail": f"{shape_accent_ok} shapes have accent colors. Colors: {shape_accent_colors_found[:6]}. Failures: {shape_accent_fail[:5]}"
    })

    # ---- CHECK: Accent colors cycle orange->blue->green ----
    cycle_pass = False
    cycle_detail = "Not enough shapes to verify cycling."
    if len(shape_accent_colors_found) >= 3:
        expected_cycle = [
            (int(ORANGE[0:2],16), int(ORANGE[2:4],16), int(ORANGE[4:6],16)),
            (int(BLUE[0:2],16), int(BLUE[2:4],16), int(BLUE[4:6],16)),
            (int(GREEN[0:2],16), int(GREEN[2:4],16), int(GREEN[4:6],16)),
        ]
        matches = 0
        for idx, found_rgb in enumerate(shape_accent_colors_found[:6]):
            exp = expected_cycle[idx % 3]
            if (abs(found_rgb[0]-exp[0])<=2 and abs(found_rgb[1]-exp[1])<=2 and abs(found_rgb[2]-exp[2])<=2):
                matches += 1
        cycle_pass = matches >= 3
        cycle_detail = f"Cycle check: {matches}/{min(len(shape_accent_colors_found),6)} shapes match expected orange->blue->green cycle. Found: {shape_accent_colors_found[:6]}"
    
    checks.append({
        "name": "accent_colors_cycle_orange_blue_green",
        "passed": cycle_pass,
        "detail": cycle_detail
    })

    # ---- CHECK: Exact brand hex values used (not approximations) ----
    # Spot-check that the brand dark color is EXACTLY #141413 and light is #faf9f5
    exact_dark_found = False
    exact_light_found = False
    for slide in prs.slides:
        try:
            bg_rgb = get_rgb_tuple(slide.background.fill.fore_color)
            if bg_rgb and rgb_close(bg_rgb, DARK, tolerance=0):
                exact_dark_found = True
            if bg_rgb and rgb_close(bg_rgb, LIGHT, tolerance=0):
                exact_light_found = True
        except Exception:
            pass
    
    exact_colors_pass = exact_dark_found or exact_light_found
    checks.append({
        "name": "exact_brand_hex_values_used",
        "passed": exact_colors_pass,
        "detail": f"Exact #141413 found: {exact_dark_found}, Exact #faf9f5 found: {exact_light_found}"
    })

    # ---- FINAL SCORING ----
    critical_checks = [
        "output_file_exists",
        "pptx_readable",
        "slide_backgrounds_brand_colors",
        "heading_font_poppins_or_arial",
        "body_font_lora_or_georgia",
        "non_text_shapes_use_accent_colors",
    ]
    bonus_checks = [
        "text_colors_are_brand_colors",
        "accent_colors_cycle_orange_blue_green",
        "exact_brand_hex_values_used",
    ]
    
    check_map = {c["name"]: c["passed"] for c in checks}
    
    critical_passed = sum(1 for c in critical_checks if check_map.get(c, False))
    bonus_passed = sum(1 for c in bonus_checks if check_map.get(c, False))
    
    total_checks = len(critical_checks) + len(bonus_checks)
    score = (critical_passed * 1.0 + bonus_passed * 0.5) / (len(critical_checks) * 1.0 + len(bonus_checks) * 0.5)
    score = round(min(score, 1.0), 3)
    
    overall_passed = critical_passed == len(critical_checks)
    
    result = {
        "passed": overall_passed,
        "score": score,
        "checks": checks
    }
    print(json.dumps(result, indent=2))

if __name__ == "__main__":
    main()