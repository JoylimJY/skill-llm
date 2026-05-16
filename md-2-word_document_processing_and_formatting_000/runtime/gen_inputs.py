import os
import random
import textwrap

random.seed(42)

WORKSPACE = "/workspace"

# ── directory scaffold ──────────────────────────────────────────────────────
dirs = [
    "scripts",
    "docs/archive",
    "docs/drafts",
    "reports/q1",
    "reports/q2",
    "data/raw",
    "data/processed",
    "assets/templates",
    "assets/logos",
    "meetings/2024",
]
for d in dirs:
    os.makedirs(os.path.join(WORKSPACE, d), exist_ok=True)

# ── distractor files ────────────────────────────────────────────────────────
distractor_files = {
    "docs/archive/old_report.md": "# Old Report\nThis is an outdated version.\n",
    "docs/drafts/notes.txt": "rough notes for Q2 review\n- check compound stability\n- confirm trial dates\n",
    "docs/drafts/summary.md": "## Draft Summary\nWork in progress. Not ready for distribution.\n",
    "reports/q1/q1_results.csv": "compound,yield,purity\nABC-101,72%,98.1%\nABC-102,68%,97.4%\n",
    "reports/q2/q2_plan.txt": "Objectives for Q2:\n1. Complete Phase I trials\n2. Submit regulatory filing\n",
    "data/raw/assay_data.json": '{"batch": "B-2024-07", "results": [{"id": "T01", "value": 4.2}]}\n',
    "data/processed/clean_data.csv": "id,score\nT01,4.2\nT02,3.9\n",
    "assets/templates/cover_template.txt": "CONFIDENTIAL\nPrepared by: Research Division\nDate: {date}\n",
    "assets/logos/placeholder.txt": "logo binary placeholder\n",
    "meetings/2024/minutes_jan.txt": "Attendees: Dr. Smith, Dr. Lee\nAgenda: Phase I review\n",
    "docs/archive/methodology_v1.md": "# Methodology v1\nThis methodology was superseded in March 2024.\n",
    "reports/q1/q1_summary.txt": "Q1 Summary: All targets met. See CSV for details.\n",
}

for rel_path, content in distractor_files.items():
    full_path = os.path.join(WORKSPACE, rel_path)
    with open(full_path, "w", encoding="utf-8") as f:
        f.write(content)

# ── the actual md_to_word conversion script ──────────────────────────────────
# This IS the proprietary script referenced in SKILL.md — it must exist.
md_to_word_script = textwrap.dedent('''\
    #!/usr/bin/env python3
    """
    Markdown to Word Converter
    Converts Markdown files to formatted Word documents (.docx).
    """

    import sys
    import re
    from pathlib import Path
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement


    def set_cell_background(cell, color_hex):
        """Set table cell background color."""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color_hex)
        tcPr.append(shd)


    def add_heading(doc, text, level):
        """Add a heading with specific formatting."""
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = True
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if level == 1:
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        elif level == 2:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
        elif level == 3:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
        para.paragraph_format.space_before = Pt(12)
        para.paragraph_format.space_after = Pt(6)
        return para


    def parse_inline(run_text, para):
        """Parse inline bold formatting."""
        parts = re.split(r"(\\*\\*[^*]+\\*\\*)", run_text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = para.add_run(part[2:-2])
                run.bold = True
            else:
                if part:
                    para.add_run(part)


    def convert_md_to_word(input_path, output_path):
        """Main conversion function."""
        with open(input_path, "r", encoding="utf-8") as f:
            lines = f.readlines()

        doc = Document()

        # Set default font
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        i = 0
        while i < len(lines):
            line = lines[i].rstrip("\\n")

            # Heading detection
            if line.startswith("### "):
                add_heading(doc, line[4:], 3)
                i += 1
                continue
            if line.startswith("## "):
                add_heading(doc, line[3:], 2)
                i += 1
                continue
            if line.startswith("# "):
                add_heading(doc, line[2:], 1)
                i += 1
                continue

            # Horizontal rule
            if re.match(r"^-{3,}$", line) or re.match(r"^\\*{3,}$", line):
                para = doc.add_paragraph()
                para.paragraph_format.space_before = Pt(6)
                para.paragraph_format.space_after = Pt(6)
                i += 1
                continue

            # Blockquote
            if line.startswith("> "):
                para = doc.add_paragraph()
                run = para.add_run(line[2:])
                run.italic = True
                para.paragraph_format.left_indent = Inches(0.5)
                i += 1
                continue

            # Bullet list
            if re.match(r"^[-\\*] ", line):
                para = doc.add_paragraph(style="List Bullet")
                parse_inline(line[2:], para)
                i += 1
                continue

            # Numbered list
            if re.match(r"^\\d+\\. ", line):
                para = doc.add_paragraph(style="List Number")
                text = re.sub(r"^\\d+\\. ", "", line)
                parse_inline(text, para)
                i += 1
                continue

            # Table detection
            if "|" in line:
                table_lines = []
                while i < len(lines) and "|" in lines[i]:
                    table_lines.append(lines[i].rstrip("\\n"))
                    i += 1
                # Filter separator rows
                rows = [r for r in table_lines if not re.match(r"^[|\\s\\-:]+$", r)]
                if rows:
                    parsed = []
                    for row in rows:
                        cells = [c.strip() for c in row.strip("|").split("|")]
                        parsed.append(cells)
                    max_cols = max(len(r) for r in parsed)
                    table = doc.add_table(rows=len(parsed), cols=max_cols)
                    table.style = "Table Grid"
                    for row_idx, row_data in enumerate(parsed):
                        for col_idx, cell_text in enumerate(row_data):
                            cell = table.cell(row_idx, col_idx)
                            cell.text = cell_text
                            if row_idx == 0:
                                set_cell_background(cell, "2E74B5")
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.bold = True
                                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                continue

            # Normal paragraph
            if line.strip():
                para = doc.add_paragraph()
                parse_inline(line, para)

            i += 1

        doc.save(output_path)
        print(f"Converted: {input_path} -> {output_path}")


    if __name__ == "__main__":
        if len(sys.argv) < 2:
            print("Usage: python md_to_word.py <input.md> [output.docx]")
            sys.exit(1)

        input_file = sys.argv[1]
        if len(sys.argv) >= 3:
            output_file = sys.argv[2]
        else:
            output_file = str(Path(input_file).with_suffix(".docx"))

        convert_md_to_word(input_file, output_file)
''')

with open(os.path.join(WORKSPACE, "scripts/md_to_word.py"), "w", encoding="utf-8") as f:
    f.write(md_to_word_script)

# ── the primary input: a rich Markdown research brief ─────────────────────
research_brief_md = textwrap.dedent("""\
    # Compound ABC-201 Clinical Research Brief

    **Prepared by:** Translational Research Division  
    **Date:** 2024-Q3  
    **Status:** Final for Stakeholder Review

    ---

    ## Executive Summary

    Compound ABC-201 has demonstrated **exceptional efficacy** in preclinical models, achieving a 94% target inhibition rate at a dose of 10 mg/kg. This brief summarizes the findings from the Phase I readiness assessment and outlines the path forward to first-in-human trials.

    > All data presented herein is derived from validated assay platforms and has undergone independent quality review.

    ---

    ## Key Findings

    ### Pharmacokinetics

    The compound exhibits a favorable pharmacokinetic profile:

    - Half-life (t½): **18.4 hours** in murine models
    - Bioavailability: **82%** oral absorption
    - Volume of distribution: 3.2 L/kg
    - Protein binding: **>95%**

    ### Safety Profile

    1. No observable adverse effects at doses up to 50 mg/kg
    2. Hepatic enzyme levels remained within normal range across all cohorts
    3. Cardiac safety confirmed via **hERG channel assay** (IC50 > 30 µM)

    ---

    ## Comparative Data Table

    | Parameter | ABC-201 | Reference Compound | Target Threshold |
    |-----------|---------|-------------------|-----------------|
    | IC50 (nM) | 4.2 | 12.7 | < 10 |
    | Selectivity Index | 312 | 145 | > 100 |
    | Oral Bioavailability (%) | 82 | 64 | > 60 |
    | Half-life (h) | 18.4 | 9.2 | > 12 |
    | Protein Binding (%) | 95 | 88 | > 80 |

    ---

    ## Recommended Next Steps

    ### Regulatory Pathway

    - Submit **Pre-IND meeting request** to FDA by Q4 2024
    - Prepare IND-enabling toxicology package (GLP compliant)
    - Engage CRO for **Phase I site selection**

    ### Manufacturing Scale-Up

    1. Transfer synthesis to GMP-compliant facility
    2. Complete **ICH Q3A/Q3B** impurity profiling
    3. Establish reference standard and release specifications

    > Regulatory timelines are contingent on successful completion of the 28-day repeat-dose toxicology study.

    ---

    ## Conclusion

    Compound ABC-201 is **ready for IND filing** pending completion of the regulatory package. The compound's **best-in-class** selectivity profile and **favorable safety window** make it a strong candidate for advancement into Phase I trials. Stakeholder alignment on resource allocation is requested by end of Q3 2024.
""")

with open(os.path.join(WORKSPACE, "docs/drafts/compound_abc201_brief.md"), "w", encoding="utf-8") as f:
    f.write(research_brief_md)

print("Workspace initialized successfully.")
print(f"Input markdown: {os.path.join(WORKSPACE, 'docs/drafts/compound_abc201_brief.md')}")