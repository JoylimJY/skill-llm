import sys
import os
from pptx import Presentation
from pptx.dml.color import RGBColor
import json

def check_presentation(workspace_dir):
    checks = []
    
    # Check if presentation file exists
    pptx_file = os.path.join(workspace_dir, 'claude_presentation.pptx')
    if not os.path.exists(pptx_file):
        checks.append({'name': 'file_exists', 'passed': False, 'detail': 'claude_presentation.pptx not found'})
        return checks
    
    checks.append({'name': 'file_exists', 'passed': True, 'detail': 'claude_presentation.pptx found'})
    
    try:
        prs = Presentation(pptx_file)
        
        # Check slide count
        slide_count = len(prs.slides)
        checks.append({
            'name': 'slide_count',
            'passed': slide_count == 5,
            'detail': f'Expected 5 slides, found {slide_count}'
        })
        
        # Check slide 1 content
        slide1_content = ''
        for shape in prs.slides[0].shapes:
            if hasattr(shape, 'text'):
                slide1_content += shape.text.lower()
        
        has_title = 'claude ai overview' in slide1_content
        has_subtitle = 'anthropic' in slide1_content and 'assistant' in slide1_content
        checks.append({
            'name': 'slide1_title',
            'passed': has_title,
            'detail': f'Title slide contains required title: {has_title}'
        })
        checks.append({
            'name': 'slide1_subtitle', 
            'passed': has_subtitle,
            'detail': f'Title slide contains subtitle about Anthropic assistant: {has_subtitle}'
        })
        
        # Check for required slide topics
        slide_topics = ['what is claude', 'key features', 'use cases', 'getting started']
        topic_found = [False] * 4
        
        for i, slide in enumerate(prs.slides[1:5], 1):
            slide_text = ''
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    slide_text += shape.text.lower()
            
            if i == 1 and 'what is claude' in slide_text:
                topic_found[0] = True
            elif i == 2 and 'key features' in slide_text or 'features' in slide_text:
                topic_found[1] = True
            elif i == 3 and 'use cases' in slide_text:
                topic_found[2] = True
            elif i == 4 and 'getting started' in slide_text:
                topic_found[3] = True
        
        for i, topic in enumerate(slide_topics):
            checks.append({
                'name': f'slide_topic_{topic.replace(" ", "_")}',
                'passed': topic_found[i],
                'detail': f'Found slide with topic "{topic}": {topic_found[i]}'
            })
        
        # Check for brand colors (at least one Anthropic color should be used)
        anthropic_colors = [
            (20, 20, 19),    # #141413 dark
            (250, 249, 245), # #faf9f5 light
            (217, 119, 87),  # #d97757 orange
            (106, 155, 204), # #6a9bcc blue
            (120, 140, 93)   # #788c5d green
        ]
        
        brand_color_used = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'fill') and hasattr(shape.fill, 'fore_color'):
                    try:
                        color = shape.fill.fore_color.rgb
                        if color and (color.red, color.green, color.blue) in anthropic_colors:
                            brand_color_used = True
                            break
                    except:
                        pass
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            try:
                                color = run.font.color.rgb
                                if color and (color.red, color.green, color.blue) in anthropic_colors:
                                    brand_color_used = True
                                    break
                            except:
                                pass
            if brand_color_used:
                break
        
        checks.append({
            'name': 'brand_colors',
            'passed': brand_color_used,
            'detail': f'At least one Anthropic brand color used: {brand_color_used}'
        })
        
        # Check for font application (basic check)
        font_applied = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            try:
                                font_name = run.font.name
                                if font_name and (font_name.lower() in ['poppins', 'lora'] or 'poppins' in font_name.lower() or 'lora' in font_name.lower()):
                                    font_applied = True
                                    break
                            except:
                                pass
                if font_applied:
                    break
            if font_applied:
                break
        
        checks.append({
            'name': 'typography',
            'passed': font_applied,
            'detail': f'Anthropic brand fonts (Poppins/Lora) applied: {font_applied}'
        })
        
    except Exception as e:
        checks.append({'name': 'presentation_readable', 'passed': False, 'detail': f'Error reading presentation: {str(e)}'})
    
    return checks

def main():
    if len(sys.argv) != 2:
        print('Usage: eval_script.py <workspace_directory>')
        sys.exit(1)
    
    workspace_dir = sys.argv[1]
    checks = check_presentation(workspace_dir)
    
    score = sum(1 for c in checks if c['passed']) / len(checks)
    passed = score >= 0.8
    
    result = {
        'passed': passed,
        'score': score,
        'checks': checks
    }
    
    print(json.dumps(result))

if __name__ == '__main__':
    main()