import json
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
import subprocess
import os

def main(workspace_dir):
    checks = []
    
    # Check if the Excel file exists with correct name
    excel_path = Path(workspace_dir) / 'techcorp_dcf_model.xlsx'
    if not excel_path.exists():
        checks.append({"name": "File exists", "passed": False, "detail": "techcorp_dcf_model.xlsx not found"})
        result = {"passed": False, "score": 0.0, "checks": checks}
        print(json.dumps(result))
        return
    
    checks.append({"name": "File exists", "passed": True, "detail": "Excel file created with correct name"})
    
    try:
        # Load workbook
        wb = load_workbook(str(excel_path), data_only=False)
        sheet_names = [name.lower() for name in wb.sheetnames]
        
        # Check required sheets exist
        required_sheets = ['assumptions', 'p&l projection', 'cash flow', 'valuation summary']
        sheets_found = 0
        for req_sheet in required_sheets:
            found = any(req_sheet in sheet_name for sheet_name in sheet_names)
            checks.append({"name": f"Sheet {req_sheet}", "passed": found, "detail": f"Found: {found}"})
            if found:
                sheets_found += 1
        
        # Try to find actual sheet names
        assumptions_sheet = None
        pl_sheet = None
        cf_sheet = None
        val_sheet = None
        
        for sheet_name in wb.sheetnames:
            lower_name = sheet_name.lower()
            if 'assumption' in lower_name:
                assumptions_sheet = wb[sheet_name]
            elif 'p&l' in lower_name or 'projection' in lower_name:
                pl_sheet = wb[sheet_name]
            elif 'cash' in lower_name and 'flow' in lower_name:
                cf_sheet = wb[sheet_name]
            elif 'valuation' in lower_name or 'summary' in lower_name:
                val_sheet = wb[sheet_name]
        
        # Check assumptions sheet content
        assumptions_found = 0
        if assumptions_sheet:
            # Look for key assumptions in the sheet
            assumption_keywords = ['revenue growth', 'ebitda margin', 'tax rate', 'discount rate', 'terminal growth']
            for row in assumptions_sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        cell_text = cell.value.lower()
                        for keyword in assumption_keywords:
                            if keyword in cell_text:
                                assumptions_found += 1
                                break
        
        checks.append({"name": "Assumptions content", "passed": assumptions_found >= 3, "detail": f"Found {assumptions_found} key assumptions"})
        
        # Check P&L sheet content
        pl_content_found = 0
        if pl_sheet:
            pl_keywords = ['revenue', 'ebitda', 'ebit', 'net income', '2024', '2025', '2026', '2027', '2028']
            for row in pl_sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        cell_text = str(cell.value).lower()
                        for keyword in pl_keywords:
                            if keyword in cell_text:
                                pl_content_found += 1
                                break
        
        checks.append({"name": "P&L content", "passed": pl_content_found >= 4, "detail": f"Found {pl_content_found} P&L elements"})
        
        # Check for formulas (should have many)
        formula_count = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                        formula_count += 1
        
        checks.append({"name": "Contains formulas", "passed": formula_count >= 20, "detail": f"Found {formula_count} formulas"})
        
        # Check for proper formatting (blue text for inputs)
        blue_cells_found = 0
        if assumptions_sheet:
            for row in assumptions_sheet.iter_rows():
                for cell in row:
                    if cell.font and cell.font.color and hasattr(cell.font.color, 'rgb'):
                        if cell.font.color.rgb and '0000ff' in str(cell.font.color.rgb).lower():
                            blue_cells_found += 1
        
        checks.append({"name": "Blue text formatting", "passed": blue_cells_found >= 3, "detail": f"Found {blue_cells_found} blue-formatted cells"})
        
        # Check cash flow sheet
        cf_content_found = 0
        if cf_sheet:
            cf_keywords = ['operating cash flow', 'capex', 'free cash flow', 'terminal value']
            for row in cf_sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        cell_text = str(cell.value).lower()
                        for keyword in cf_keywords:
                            if keyword in cell_text:
                                cf_content_found += 1
                                break
        
        checks.append({"name": "Cash Flow content", "passed": cf_content_found >= 2, "detail": f"Found {cf_content_found} CF elements"})
        
        # Check valuation sheet
        val_content_found = 0
        if val_sheet:
            val_keywords = ['enterprise value', 'equity value', 'net cash']
            for row in val_sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        cell_text = str(cell.value).lower()
                        for keyword in val_keywords:
                            if keyword in cell_text:
                                val_content_found += 1
                                break
        
        checks.append({"name": "Valuation content", "passed": val_content_found >= 2, "detail": f"Found {val_content_found} valuation elements"})
        
        wb.close()
        
        # Try to recalculate formulas and check for errors
        try:
            result = subprocess.run([sys.executable, 'scripts/recalc.py', str(excel_path)], 
                                  capture_output=True, text=True, timeout=60)
            if result.returncode == 0:
                recalc_output = json.loads(result.stdout)
                total_errors = recalc_output.get('total_errors', 0)
                checks.append({"name": "Zero formula errors", "passed": total_errors == 0, "detail": f"Found {total_errors} formula errors"})
            else:
                checks.append({"name": "Zero formula errors", "passed": False, "detail": "Could not recalculate formulas"})
        except Exception as e:
            checks.append({"name": "Zero formula errors", "passed": False, "detail": f"Recalculation failed: {str(e)}"})
        
    except Exception as e:
        checks.append({"name": "File processing", "passed": False, "detail": f"Error: {str(e)}"})
    
    # Calculate score
    passed_checks = sum(1 for check in checks if check['passed'])
    total_checks = len(checks)
    score = passed_checks / total_checks if total_checks > 0 else 0.0
    
    result = {
        "passed": score >= 0.8,
        "score": score,
        "checks": checks
    }
    
    print(json.dumps(result))

if __name__ == '__main__':
    main(sys.argv[1])