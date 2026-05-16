import os
import json
import random
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

random.seed(42)

workspace = Path("/workspace")

# --- Create distractor directory structure ---
dirs = [
    "archive/2022/q1",
    "archive/2022/q4",
    "archive/2023/q2",
    "drafts/internal",
    "drafts/external",
    "templates/old",
    "templates/approved",
    "reports/monthly",
    "reports/annual",
    "misc",
]
for d in dirs:
    (workspace / d).mkdir(parents=True, exist_ok=True)

# --- Distractor files ---
distractors = {
    "archive/2022/q1/summary.txt": "这是2022年第一季度的工作总结草稿，尚未格式化。",
    "archive/2022/q4/year_end.md": "# 年终报告\n\n本报告仅供参考。",
    "archive/2023/q2/meeting_notes.txt": "会议纪要：参会人员包括各部门负责人。",
    "drafts/internal/proposal_v1.docx": None,  # will create as docx
    "drafts/external/letter_draft.txt": "尊敬的领导：\n\n请批示。",
    "templates/old/template_2020.json": json.dumps({"font": "宋体", "size": 12}),
    "templates/old/README.md": "# 旧模板说明\n\n此目录包含已废弃的模板文件。",
    "reports/monthly/oct_report.txt": "十月工作报告内容待填写。",
    "reports/annual/2023_annual.md": "## 年度工作报告\n\n### 一、总体情况\n\n工作正常推进。",
    "misc/notes.txt": "杂项笔记：格式参考旧版文件。",
    "misc/contact_list.txt": "联系人列表：张三、李四、王五",
}

for rel_path, content in distractors.items():
    full_path = workspace / rel_path
    if rel_path.endswith(".docx"):
        doc = Document()
        doc.add_paragraph("草稿内容，仅供内部参考。")
        doc.save(str(full_path))
    else:
        full_path.write_text(content, encoding="utf-8")

# ============================================================
# CORE ARTIFACT 1: The approved template Word document
# This has a specific proprietary style that the agent must extract
# ============================================================

def set_paragraph_format(para, font_name, font_size_pt, bold=False,
                          align=WD_ALIGN_PARAGRAPH.LEFT,
                          first_line_chars=0,
                          space_before_pt=0, space_after_pt=0,
                          line_spacing_rule=None, line_spacing_val=None):
    """Helper to set paragraph format on a docx paragraph."""
    pf = para.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before_pt)
    pf.space_after = Pt(space_after_pt)

    if first_line_chars > 0:
        # Set first line indent via XML (character-based)
        pPr = para._p.get_or_add_pPr()
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            pPr.append(ind)
        ind.set(qn('w:firstLineChars'), str(first_line_chars * 100))
        ind.set(qn('w:firstLine'), str(first_line_chars * 160))

    if line_spacing_val:
        pf.line_spacing = line_spacing_val
    
    for run in para.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size_pt)
        run.font.bold = bold
        # Set east asia font
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), font_name)

template_doc = Document()

# Set page margins (GB/T 9704-2012 standard)
section = template_doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(3.7)
section.bottom_margin = Cm(3.5)
section.left_margin = Cm(2.8)
section.right_margin = Cm(2.6)

# Title paragraph
title_para = template_doc.add_paragraph("关于加强信息安全管理工作的通知")
set_paragraph_format(
    title_para,
    font_name="方正小标宋体",
    font_size_pt=22,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    first_line_chars=0,
    space_before_pt=0,
    space_after_pt=0,
    line_spacing_val=Pt(22 * 1.5)
)

# Document number
wenhao_para = template_doc.add_paragraph("信安办发〔2024〕15号")
set_paragraph_format(
    wenhao_para,
    font_name="仿宋_GB2312",
    font_size_pt=16,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    first_line_chars=0,
    space_before_pt=0,
    space_after_pt=0,
    line_spacing_val=Pt(16 * 1.5)
)

# Level-1 heading
h1_para = template_doc.add_paragraph("一、工作背景")
set_paragraph_format(
    h1_para,
    font_name="黑体",
    font_size_pt=16,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    first_line_chars=0,
    space_before_pt=0,
    space_after_pt=0,
    line_spacing_val=Pt(16 * 1.5)
)

# Level-2 heading
h2_para = template_doc.add_paragraph("（一）现状分析")
set_paragraph_format(
    h2_para,
    font_name="楷体_GB2312",
    font_size_pt=16,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    first_line_chars=0,
    space_before_pt=0,
    space_after_pt=0,
    line_spacing_val=Pt(16 * 1.5)
)

# Body paragraph
body_para = template_doc.add_paragraph("当前信息安全形势严峻，各单位须严格落实相关规定，确保数据安全和系统稳定运行。")
set_paragraph_format(
    body_para,
    font_name="仿宋_GB2312",
    font_size_pt=16,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_chars=2,
    space_before_pt=0,
    space_after_pt=0,
    line_spacing_val=Pt(16 * 1.5)
)

template_path = workspace / "templates/approved/official_template.docx"
template_doc.save(str(template_path))

# ============================================================
# CORE ARTIFACT 2: The messy Markdown report to be reformatted
# using the template's style
# ============================================================

messy_md_content = """# 关于推进数字化转型工作的报告

## 一、工作进展

### （一）基础设施建设

各地区已完成网络基础设施升级改造，覆盖率达到95%以上。主要工作包括：

- 完成骨干网络扩容
- 部署新一代防火墙设备
- 建立统一身份认证平台

### （二）应用系统迁移

完成核心业务系统向云平台迁移，迁移率达到80%。

## 二、存在的问题

### （一）资金缺口

部分地区数字化改造资金缺口较大，需要上级统筹协调解决。

### （二）人才短缺

专业技术人才储备不足，培训体系尚不完善。

## 三、下步工作计划

重点推进以下三项工作：继续完善基础设施建设、加快人才培养步伐、强化安全保障措施。
"""

md_report_path = workspace / "drafts/internal/digital_transform_report.md"
md_report_path.write_text(messy_md_content, encoding="utf-8")

# ============================================================
# CORE ARTIFACT 3: A plain text briefing to be converted to
# official format Markdown
# ============================================================

briefing_txt = """2024年度安全检查情况通报

安监办发〔2024〕22号

一、检查概况

本次安全检查覆盖全系统共计56家单位，检查时间为2024年10月至11月。检查组由专家和业务骨干组成，采取现场检查和资料核查相结合的方式开展工作。

（一）检查范围

检查范围涵盖消防安全、网络安全、生产安全三个方面，重点对高风险区域进行了全面排查。

（二）检查方式

采取"四不两直"方式开展检查，即不发通知、不打招呼、不听汇报、不用陪同接待，直奔基层、直插现场。

二、主要发现

各单位总体安全状况良好，但仍存在部分薄弱环节，需要引起高度重视并切实加以整改。

（一）消防安全方面

部分单位消防设施维护不到位，存在安全隐患。

（二）网络安全方面

个别单位网络安全防护措施不完善，数据备份机制不健全。

三、整改要求

各单位须在2024年12月31日前完成整改，并将整改情况书面报告报送安监办。
"""

txt_briefing_path = workspace / "reports/annual/safety_check_briefing.txt"
txt_briefing_path.write_text(briefing_txt, encoding="utf-8")

# --- Extra distractors in templates/approved ---
(workspace / "templates/approved/cover_page_template.txt").write_text(
    "封面模板：此文件不含格式信息，仅供参考。", encoding="utf-8"
)
(workspace / "templates/approved/format_notes.txt").write_text(
    "格式说明：请参照已批准的Word模板文件。注意页边距和字体设置。", encoding="utf-8"
)

print("Workspace generated successfully.")
print(f"Template: {template_path}")
print(f"Markdown report: {md_report_path}")
print(f"Briefing txt: {txt_briefing_path}")