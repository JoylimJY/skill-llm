import sys
import json
from pathlib import Path

def run_checks(workspace: str):
    checks = []
    workspace = Path(workspace)

    # -------------------------------------------------------
    # Helper: find a file by glob
    # -------------------------------------------------------
    def find_file(pattern):
        results = list(workspace.rglob(pattern))
        return results[0] if results else None

    # -------------------------------------------------------
    # CHECK GROUP A: Template format extraction
    # The agent must have produced a JSON format file from the template
    # -------------------------------------------------------
    format_json_file = find_file("*.json")
    # Filter out template/old json distractors
    format_json_candidates = [
        f for f in workspace.rglob("*.json")
        if "template_2020" not in f.name and "format" in f.name.lower()
    ]
    # Also accept any JSON that has paragraph_styles key
    valid_format_jsons = []
    for f in workspace.rglob("*.json"):
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
            if "paragraph_styles" in data and "page_setup" in data:
                valid_format_jsons.append((f, data))
        except:
            pass

    has_format_json = len(valid_format_jsons) > 0
    checks.append({
        "name": "format_json_extracted",
        "passed": has_format_json,
        "detail": f"Found {len(valid_format_jsons)} valid format JSON file(s) with page_setup and paragraph_styles keys." if has_format_json else "No valid format JSON extracted from template."
    })

    if valid_format_jsons:
        _, fdata = valid_format_jsons[0]
        # Check page margins match official template (3.7, 3.5, 2.8, 2.6)
        ps = fdata.get("page_setup", {})
        margins_ok = (
            abs(ps.get("top_margin_cm", 0) - 3.7) < 0.15 and
            abs(ps.get("bottom_margin_cm", 0) - 3.5) < 0.15 and
            abs(ps.get("left_margin_cm", 0) - 2.8) < 0.15 and
            abs(ps.get("right_margin_cm", 0) - 2.6) < 0.15
        )
        checks.append({
            "name": "format_json_margins_correct",
            "passed": margins_ok,
            "detail": f"Page margins in JSON: top={ps.get('top_margin_cm')}, bottom={ps.get('bottom_margin_cm')}, left={ps.get('left_margin_cm')}, right={ps.get('right_margin_cm')}. Expected: 3.7/3.5/2.8/2.6"
        })

        styles = fdata.get("paragraph_styles", {})
        has_heading1 = "heading1" in styles
        has_body = "body" in styles
        checks.append({
            "name": "format_json_has_paragraph_styles",
            "passed": has_heading1 and has_body,
            "detail": f"paragraph_styles keys: {list(styles.keys())}"
        })
    else:
        checks.append({"name": "format_json_margins_correct", "passed": False, "detail": "No format JSON found."})
        checks.append({"name": "format_json_has_paragraph_styles", "passed": False, "detail": "No format JSON found."})

    # -------------------------------------------------------
    # CHECK GROUP B: Markdown report reformatted to Word
    # The agent must produce a .docx output from the markdown report
    # Expected filename variations: anything docx that is NOT a template or distractor
    # -------------------------------------------------------
    try:
        from docx import Document
        from docx.shared import Cm, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        # Find docx outputs that are NOT the original template
        template_path = workspace / "templates/approved/official_template.docx"
        draft_docx = workspace / "drafts/internal/proposal_v1.docx"

        output_docxes = [
            f for f in workspace.rglob("*.docx")
            if f.resolve() != template_path.resolve()
            and f.resolve() != draft_docx.resolve()
        ]

        # Among those, find one that likely came from the markdown report
        # (has title "关于推进数字化转型工作的报告" or similar content)
        md_output_docx = None
        for f in output_docxes:
            try:
                doc = Document(str(f))
                texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                # Look for digital transform report content
                full_text = " ".join(texts)
                if "数字化转型" in full_text or "digital_transform" in f.stem.lower() or "转型" in full_text:
                    md_output_docx = f
                    break
            except:
                pass

        if md_output_docx is None and output_docxes:
            # fallback: take newest non-template docx
            md_output_docx = sorted(output_docxes, key=lambda x: x.stat().st_mtime, reverse=True)[0]

        has_md_output = md_output_docx is not None
        checks.append({
            "name": "md_report_converted_to_docx",
            "passed": has_md_output,
            "detail": f"Found output docx from MD report: {md_output_docx}" if has_md_output else "No output docx found for the markdown report."
        })

        if has_md_output:
            doc = Document(str(md_output_docx))
            section = doc.sections[0]

            # Check page margins (from template: 3.7/3.5/2.8/2.6)
            top_ok = abs(section.top_margin.cm - 3.7) < 0.2
            bottom_ok = abs(section.bottom_margin.cm - 3.5) < 0.2
            left_ok = abs(section.left_margin.cm - 2.8) < 0.2
            right_ok = abs(section.right_margin.cm - 2.6) < 0.2
            margins_ok = top_ok and bottom_ok and left_ok and right_ok
            checks.append({
                "name": "md_output_page_margins",
                "passed": margins_ok,
                "detail": f"Output docx margins: top={section.top_margin.cm:.2f}, bottom={section.bottom_margin.cm:.2f}, left={section.left_margin.cm:.2f}, right={section.right_margin.cm:.2f}"
            })

            # Check paragraph styles applied
            paras = [p for p in doc.paragraphs if p.text.strip()]
            
            # Check title paragraph (first paragraph should be 方正小标宋体 or the template's title font, center aligned)
            title_ok = False
            if paras:
                tp = paras[0]
                if tp.paragraph_format.alignment in (WD_ALIGN_PARAGRAPH.CENTER, 1):
                    title_ok = True
                # Also check via font
                for run in tp.runs:
                    rPr = run._r.find(qn('w:rPr'))
                    if rPr is not None:
                        rFonts = rPr.find(qn('w:rFonts'))
                        if rFonts is not None:
                            ea = rFonts.get(qn('w:eastAsia'))
                            if ea and ('小标宋' in ea or '标宋' in ea):
                                title_ok = True

            checks.append({
                "name": "md_output_title_style",
                "passed": title_ok,
                "detail": f"Title paragraph alignment: {paras[0].paragraph_format.alignment if paras else 'N/A'}"
            })

            # Check that at least one body paragraph has first-line indent (2 chars = firstLineChars=200)
            body_indent_ok = False
            for p in paras:
                pPr = p._p.find(qn('w:pPr'))
                if pPr is not None:
                    ind = pPr.find(qn('w:ind'))
                    if ind is not None:
                        flc = ind.get(qn('w:firstLineChars'))
                        if flc and int(flc) >= 200:
                            body_indent_ok = True
                            break

            checks.append({
                "name": "md_output_body_first_line_indent",
                "passed": body_indent_ok,
                "detail": "At least one body paragraph must have firstLineChars >= 200 (2 characters)."
            })

            # Check heading1 font (黑体) 
            heading1_ok = False
            for p in paras:
                text = p.text.strip()
                import re
                if re.match(r'^[一二三四五六七八九十]+[、．.]', text):
                    for run in p.runs:
                        rPr = run._r.find(qn('w:rPr'))
                        if rPr is not None:
                            rFonts = rPr.find(qn('w:rFonts'))
                            if rFonts is not None:
                                ea = rFonts.get(qn('w:eastAsia'))
                                if ea and '黑体' in ea:
                                    heading1_ok = True
                    break

            checks.append({
                "name": "md_output_heading1_font",
                "passed": heading1_ok,
                "detail": "Heading1 paragraphs (一、二、三、) should use 黑体 font."
            })

            # Check body font (仿宋_GB2312)
            body_font_ok = False
            for p in paras:
                text = p.text.strip()
                import re
                if not re.match(r'^[一二三四五六七八九十]+[、．.]', text) and \
                   not re.match(r'^（[一二三四五六七八九十]+）', text) and \
                   p != paras[0]:
                    for run in p.runs:
                        rPr = run._r.find(qn('w:rPr'))
                        if rPr is not None:
                            rFonts = rPr.find(qn('w:rFonts'))
                            if rFonts is not None:
                                ea = rFonts.get(qn('w:eastAsia'))
                                if ea and '仿宋' in ea:
                                    body_font_ok = True
                    if body_font_ok:
                        break

            checks.append({
                "name": "md_output_body_font",
                "passed": body_font_ok,
                "detail": "Body paragraphs should use 仿宋_GB2312 font."
            })

    except Exception as e:
        checks.append({"name": "md_report_converted_to_docx", "passed": False, "detail": f"Exception: {e}"})
        checks.append({"name": "md_output_page_margins", "passed": False, "detail": f"Exception: {e}"})
        checks.append({"name": "md_output_title_style", "passed": False, "detail": str(e)})
        checks.append({"name": "md_output_body_first_line_indent", "passed": False, "detail": str(e)})
        checks.append({"name": "md_output_heading1_font", "passed": False, "detail": str(e)})
        checks.append({"name": "md_output_body_font", "passed": False, "detail": str(e)})

    # -------------------------------------------------------
    # CHECK GROUP C: Plain text briefing converted to official Markdown
    # -------------------------------------------------------
    try:
        # Find markdown output files that are NOT original distractors
        original_mds = {
            str((workspace / "archive/2022/q4/year_end.md").resolve()),
            str((workspace / "archive/2023/q2/meeting_notes.txt").resolve()),
            str((workspace / "drafts/internal/digital_transform_report.md").resolve()),
            str((workspace / "reports/annual/2023_annual.md").resolve()),
            str((workspace / "templates/old/README.md").resolve()),
        }

        output_mds = [
            f for f in workspace.rglob("*.md")
            if str(f.resolve()) not in original_mds
        ]

        # Find one that has content from the briefing
        briefing_md = None
        for f in output_mds:
            try:
                content = f.read_text(encoding="utf-8")
                if "安全检查" in content or "安监办" in content or "safety_check" in f.stem.lower():
                    briefing_md = f
                    break
            except:
                pass

        has_briefing_md = briefing_md is not None
        checks.append({
            "name": "txt_briefing_converted_to_md",
            "passed": has_briefing_md,
            "detail": f"Found output markdown from txt briefing: {briefing_md}" if has_briefing_md else "No output markdown found for the text briefing."
        })

        if has_briefing_md:
            content = briefing_md.read_text(encoding="utf-8")

            # Check that the document title is marked as # heading
            has_title_heading = content.strip().startswith("# ")
            checks.append({
                "name": "md_briefing_has_title_heading",
                "passed": has_title_heading,
                "detail": f"Output markdown should start with '# ' title. First 80 chars: {content[:80]!r}"
            })

            # Check that level-1 headings (一、二、三、) are marked as ## 
            import re
            heading1_in_md = re.search(r'^##\s+[一二三四五六七八九十]+[、．.]', content, re.MULTILINE)
            checks.append({
                "name": "md_briefing_heading1_format",
                "passed": bool(heading1_in_md),
                "detail": f"Level-1 headings (一、二、) should appear as '## ' in output markdown. Found: {bool(heading1_in_md)}"
            })

            # Check that level-2 headings (（一）etc.) are marked as ###
            heading2_in_md = re.search(r'^###\s+（[一二三四五六七八九十]+）', content, re.MULTILINE)
            checks.append({
                "name": "md_briefing_heading2_format",
                "passed": bool(heading2_in_md),
                "detail": f"Level-2 headings （一）（二） should appear as '### ' in output markdown. Found: {bool(heading2_in_md)}"
            })

            # Check content integrity: key phrases should be present
            has_key_content = "安全检查" in content and "整改" in content
            checks.append({
                "name": "md_briefing_content_preserved",
                "passed": has_key_content,
                "detail": f"Key content ('安全检查', '整改') should be present in output markdown."
            })

    except Exception as e:
        checks.append({"name": "txt_briefing_converted_to_md", "passed": False, "detail": f"Exception: {e}"})
        checks.append({"name": "md_briefing_has_title_heading", "passed": False, "detail": str(e)})
        checks.append({"name": "md_briefing_heading1_format", "passed": False, "detail": str(e)})
        checks.append({"name": "md_briefing_heading2_format", "passed": False, "detail": str(e)})
        checks.append({"name": "md_briefing_content_preserved", "passed": False, "detail": str(e)})

    # -------------------------------------------------------
    # Final scoring
    # -------------------------------------------------------
    total = len(checks)
    passed_count = sum(1 for c in checks if c["passed"])
    score = round(passed_count / total, 3) if total > 0 else 0.0
    overall_passed = score >= 0.75

    return {
        "passed": overall_passed,
        "score": score,
        "checks": checks
    }


if __name__ == "__main__":
    workspace_dir = sys.argv[1] if len(sys.argv) > 1 else "/workspace"
    result = run_checks(workspace_dir)
    print(json.dumps(result, ensure_ascii=False, indent=2))