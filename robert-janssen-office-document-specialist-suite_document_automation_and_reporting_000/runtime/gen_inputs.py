import os
import random
from pathlib import Path
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

random.seed(42)

workspace = Path("/workspace")

# ── directory structure with distractors ──────────────────────────────────────
dirs = [
    "raw_data",
    "raw_data/archive",
    "raw_data/archive/2022",
    "raw_data/archive/2023",
    "raw_data/temp",
    "templates",
    "templates/old",
    "output",
    "scripts",
    "scripts/utils",
    "logs",
]
for d in dirs:
    (workspace / d).mkdir(parents=True, exist_ok=True)

# ── distractor files ──────────────────────────────────────────────────────────
distractor_texts = {
    "raw_data/archive/2022/q4_summary.txt": "Q4 2022 archived. Do not use.",
    "raw_data/archive/2023/q1_notes.txt": "Preliminary notes - superseded.",
    "raw_data/archive/2023/q2_notes.txt": "Q2 raw dump - incomplete.",
    "raw_data/temp/scratch.csv": "region,sales\nNorth,999\nSouth,888",
    "templates/old/report_v1.txt": "Old report template - deprecated.",
    "templates/old/deck_v1.txt": "Old deck template - deprecated.",
    "templates/style_guide.txt": "Use corporate blue #003366 for headers.",
    "scripts/utils/helper.py": "# helper utilities (not in use)\npass\n",
    "scripts/data_cleaner.py": "# placeholder cleaner\npass\n",
    "logs/pipeline.log": "2024-01-10 INFO pipeline started\n2024-01-10 ERROR missing input - aborted\n",
    "raw_data/temp/old_export.txt": "Stale export from legacy system.",
}
for rel_path, content in distractor_texts.items():
    (workspace / rel_path).write_text(content)

# ── MAIN INPUT: messy Q1-2024 sales Excel workbook ───────────────────────────
# Intentionally messy: merged header cells, blank rows, a comment row,
# inconsistent capitalisation in region names, and a totals row that the
# agent must IGNORE (or detect as non-data).
wb = Workbook()
ws = wb.active
ws.title = "Q1_Sales"

# Row 1: merged banner title
ws.merge_cells("A1:F1")
banner = ws["A1"]
banner.value = "ACME Corp – Q1 2024 Regional Sales Export (CONFIDENTIAL)"
banner.font = Font(bold=True, size=14)
banner.alignment = Alignment(horizontal="center")

# Row 2: blank spacer
ws.append([None])

# Row 3: column headers (some with trailing spaces / mixed case)
headers = ["Region ", "Salesperson", "Product Category", "Units Sold", "Unit Price (USD)", "Returns"]
ws.append(headers)
header_row = ws[3]
for cell in header_row:
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="003366")
    cell.font = Font(bold=True, color="FFFFFF")

# Row 4: blank spacer
ws.append([None])

# Data rows (rows 5-24): 20 transactions across 4 regions
regions = ["North", "SOUTH", "East", "west"]   # intentionally inconsistent case
categories = ["Electronics", "Apparel", "Home Goods", "Sporting Goods", "Toys"]

random.seed(42)
transactions = []
for i in range(20):
    region = regions[i % 4]
    person = f"Agent_{chr(65 + i)}"
    cat = random.choice(categories)
    units = random.randint(50, 500)
    price = round(random.uniform(10.0, 200.0), 2)
    returns = random.randint(0, int(units * 0.1))
    transactions.append([region, person, cat, units, price, returns])

for t in transactions:
    ws.append(t)

# Row 25: a comment/metadata row the agent must skip
ws.append(["NOTE: Data extracted from ERP on 2024-04-01. Verify before use.", None, None, None, None, None])

# Row 26: blank
ws.append([None])

# Row 27: a "TOTALS" sentinel row (pre-computed incorrectly — agent must recompute)
ws.append(["TOTALS", None, None, 9999, None, 9999])

# Column widths
for col_idx, _ in enumerate(headers, start=1):
    ws.column_dimensions[get_column_letter(col_idx)].width = 22

# Second sheet: lookup table (distractor)
ws2 = wb.create_sheet("Region Lookup")
ws2.append(["Code", "Full Name", "Manager"])
ws2.append(["N", "North Region", "Alice Johnson"])
ws2.append(["S", "South Region", "Bob Martinez"])
ws2.append(["E", "East Region", "Carol Lee"])
ws2.append(["W", "West Region", "David Kim"])

wb.save(workspace / "raw_data" / "q1_2024_sales_export.xlsx")

print("Workspace generated successfully.")
print(f"Transactions written: {len(transactions)}")

# Print expected aggregates so benchmark builder can verify (not visible to agent)
from collections import defaultdict
region_units = defaultdict(int)
region_revenue = defaultdict(float)
region_returns = defaultdict(int)

for t in transactions:
    r = t[0].strip().capitalize()
    units = t[3]
    price = t[4]
    returns = t[5]
    region_units[r] += units
    region_revenue[r] += units * price
    region_returns[r] += returns

print("\nExpected aggregates (for eval reference):")
total_units = 0
total_revenue = 0.0
total_returns = 0
for r in sorted(region_units.keys()):
    u = region_units[r]
    rev = region_revenue[r]
    ret = region_returns[r]
    print(f"  {r}: units={u}, revenue={rev:.2f}, returns={ret}")
    total_units += u
    total_revenue += rev
    total_returns += ret
print(f"  TOTAL: units={total_units}, revenue={total_revenue:.2f}, returns={total_returns}")