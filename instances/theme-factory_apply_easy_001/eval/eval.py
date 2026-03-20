#!/usr/bin/env python3
import sys
import os
import json
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_presentation_styling(workspace_dir):
    checks = []
    
    # Check if original presentation exists
    original_path = os.path.join(workspace_dir, 'sales_presentation.pptx')
    if not os.path.exists(original_path):
        checks.append({"name": "original_presentation_exists", "passed": False, "detail": "Original presentation file not found"})
        return checks, False
    
    checks.append({"name": "original_presentation_exists", "passed": True, "detail": "Original presentation found"})
    
    # Load presentation
    try:
        prs = Presentation(original_path)
        checks.append({"name": "presentation_loadable", "passed": True, "detail": "Presentation can be loaded"})
    except Exception as e:
        checks.append({"name": "presentation_loadable", "passed": False, "detail": f"Cannot load presentation: {str(e)}"})
        return checks, False
    
    # Check if marker content is preserved
    marker_found = False
    styling_applied = False
    
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, 'text'):
                text = shape.text
                if 'MARKER_TITLE_2024' in text or 'MARKER_SUBTITLE' in text or 'MARKER_METRICS' in text:
                    marker_found = True
                    
                # Check if styling has been applied (look for non-default fonts or colors)
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.name and run.font.name not in ['Calibri', 'Arial', None]:
                                styling_applied = True
                            if hasattr(run.font, 'color') and run.font.color and hasattr(run.font.color, 'rgb'):
                                styling_applied = True
    
    checks.append({"name": "marker_content_preserved", "passed": marker_found, "detail": "Original marker content found" if marker_found else "Marker content missing"})
    checks.append({"name": "styling_applied", "passed": styling_applied, "detail": "Custom styling detected" if styling_applied else "No custom styling detected"})
    
    # Check if themes directory exists
    themes_dir = os.path.join(workspace_dir, 'themes')
    themes_exist = os.path.exists(themes_dir) and len(os.listdir(themes_dir)) > 0
    checks.append({"name": "themes_available", "passed": themes_exist, "detail": "Theme files found" if themes_exist else "No theme files found"})
    
    # Overall pass condition: presentation exists, loads, has markers, and either has styling or themes are available for selection
    overall_pass = marker_found and (styling_applied or themes_exist)
    
    return checks, overall_pass

def main():
    if len(sys.argv) != 2:
        print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "usage", "passed": False, "detail": "Usage: eval_script.py <workspace_dir>"}]}))
        return
    
    workspace_dir = sys.argv[1]
    
    checks, overall_pass = check_presentation_styling(workspace_dir)
    
    # Calculate score based on checks
    passed_checks = sum(1 for check in checks if check['passed'])
    total_checks = len(checks)
    score = passed_checks / total_checks if total_checks > 0 else 0.0
    
    result = {
        "passed": overall_pass,
        "score": score,
        "checks": checks
    }
    
    print(json.dumps(result, indent=2))

if __name__ == '__main__':
    main()