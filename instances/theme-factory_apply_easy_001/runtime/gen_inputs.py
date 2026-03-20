#!/usr/bin/env python3
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Create a basic presentation with marker content
prs = Presentation()

# Title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Q4 Sales Results MARKER_TITLE_2024"
subtitle.text = "Annual Performance Review MARKER_SUBTITLE"

# Content slide 1
bullet_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(bullet_slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Key Metrics MARKER_METRICS"
content.text = "• Revenue increased 15% MARKER_REVENUE\n• Customer satisfaction: 92% MARKER_SATISFACTION\n• Market expansion: 3 new regions MARKER_EXPANSION"

# Content slide 2
slide = prs.slides.add_slide(bullet_slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Future Goals MARKER_GOALS"
content.text = "• Target 20% growth MARKER_TARGET\n• Launch new product line MARKER_PRODUCT\n• Expand team by 25% MARKER_TEAM"

prs.save('sales_presentation.pptx')

# Create theme files directory and sample themes
os.makedirs('themes', exist_ok=True)

# Ocean Depths theme
with open('themes/ocean_depths.txt', 'w') as f:
    f.write("""Theme: Ocean Depths
Primary Color: #1e3a5f
Secondary Color: #4a90a4
Accent Color: #87ceeb
Background Color: #f0f8ff
Text Color: #2c3e50
Header Font: Montserrat
Body Font: Open Sans
""")

# Sunset Boulevard theme
with open('themes/sunset_boulevard.txt', 'w') as f:
    f.write("""Theme: Sunset Boulevard
Primary Color: #ff6b35
Secondary Color: #f7931e
Accent Color: #ffdc00
Background Color: #fff8f0
Text Color: #8b4513
Header Font: Playfair Display
Body Font: Source Sans Pro
""")

# Modern Minimalist theme
with open('themes/modern_minimalist.txt', 'w') as f:
    f.write("""Theme: Modern Minimalist
Primary Color: #2c3e50
Secondary Color: #7f8c8d
Accent Color: #3498db
Background Color: #ffffff
Text Color: #34495e
Header Font: Roboto
Body Font: Lato
""")

print("Generated sales_presentation.pptx and theme files")