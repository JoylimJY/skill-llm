import sys
import json
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
import subprocess
import os

def eval_financial_model(workspace_dir):
    workspace = Path(workspace_dir)
    checks = []
    score = 0.0
    
    # Check if financial model exists
    model_files = list(workspace.glob('*.xlsx'))
    model_files = [f for f in model_files if 'funding_rounds' not in f.name]
    
    if not model_files:
        return {'passed': False, 'score': 0.0, 'checks': [{'name': 'File exists', 'passed': False, 'detail': 'No Excel financial model found'}]}
    
    model_file = model_files[0]
    checks.append({'name': 'File exists', 'passed': True, 'detail': f'Found {model_file.name}'})
    score += 10
    
    try:
        # Load workbook
        wb = load_workbook(str(model_file), data_only=False)
        wb_data = load_workbook(str(model_file), data_only=True)
        
        # Check multiple sheets
        required_sheets = ['Historical', 'Projections', 'Assumptions']
        sheet_score = 0
        for sheet in required_sheets:
            if any(sheet.lower() in s.lower() for s in wb.sheetnames):
                sheet_score += 1
        
        if sheet_score >= 2:
            checks.append({'name': 'Multiple sheets', 'passed': True, 'detail': f'Found {len(wb.sheetnames)} sheets'})
            score += 15
        else:
            checks.append({'name': 'Multiple sheets', 'passed': False, 'detail': f'Only {len(wb.sheetnames)} sheets found'})
        
        # Check for formulas
        formula_count = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                        formula_count += 1
        
        if formula_count >= 20:
            checks.append({'name': 'Formula usage', 'passed': True, 'detail': f'Found {formula_count} formulas'})
            score += 20
        else:
            checks.append({'name': 'Formula usage', 'passed': False, 'detail': f'Only {formula_count} formulas found'})
        
        # Check for historical data integration
        historical_found = False
        revenue_values = [500000, 850000, 1200000]
        
        for sheet_name in wb.sheetnames:
            ws_data = wb_data[sheet_name]
            for row in ws_data.iter_rows():
                for cell in row:
                    if cell.value in revenue_values:
                        historical_found = True
                        break
        
        if historical_found:
            checks.append({'name': 'Historical data integration', 'passed': True, 'detail': 'Historical revenue data found'})
            score += 15
        else:
            checks.append({'name': 'Historical data integration', 'passed': False, 'detail': 'Historical data not properly integrated'})
        
        # Check for projection years (2024-2028)
        projection_years = ['2024', '2025', '2026', '2027', '2028']
        years_found = 0
        
        for sheet_name in wb.sheetnames:
            ws_data = wb_data[sheet_name]
            for row in ws_data.iter_rows(max_row=10):
                for cell in row:
                    if cell.value and str(cell.value) in projection_years:
                        years_found += 1
                        break
        
        if years_found >= 3:
            checks.append({'name': 'Projection periods', 'passed': True, 'detail': f'Found {years_found} projection years'})
            score += 15
        else:
            checks.append({'name': 'Projection periods', 'passed': False, 'detail': f'Only {years_found} projection years found'})
        
        # Check for professional formatting (colors)
        color_usage = False
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(max_row=50, max_col=20):
                for cell in row:
                    if cell.font.color and str(cell.font.color.rgb) in ['FF0000FF', 'FF000000', 'FF008000']:
                        color_usage = True
                        break
        
        if color_usage:
            checks.append({'name': 'Professional formatting', 'passed': True, 'detail': 'Color coding applied'})
            score += 10
        else:
            checks.append({'name': 'Professional formatting', 'passed': False, 'detail': 'No professional color coding found'})
        
        # Check for funding data integration
        funding_values = [1000000, 5000000, 15000000]
        funding_found = False
        
        for sheet_name in wb.sheetnames:
            ws_data = wb_data[sheet_name]
            for row in ws_data.iter_rows():
                for cell in row:
                    if cell.value in funding_values:
                        funding_found = True
                        break
        
        if funding_found:
            checks.append({'name': 'Funding data integration', 'passed': True, 'detail': 'Funding round data integrated'})
            score += 10
        else:
            checks.append({'name': 'Funding data integration', 'passed': False, 'detail': 'Funding data not integrated'})
        
        # Check for error-free formulas
        excel_errors = ['#VALUE!', '#DIV/0!', '#REF!', '#NAME?', '#NULL!', '#NUM!', '#N/A']
        error_count = 0
        
        for sheet_name in wb_data.sheetnames:
            ws = wb_data[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        for err in excel_errors:
                            if err in cell.value:
                                error_count += 1
        
        if error_count == 0:
            checks.append({'name': 'Error-free formulas', 'passed': True, 'detail': 'No Excel errors found'})
            score += 15
        else:
            checks.append({'name': 'Error-free formulas', 'passed': False, 'detail': f'Found {error_count} Excel errors'})
        
        wb.close()
        wb_data.close()
        
    except Exception as e:
        checks.append({'name': 'File analysis', 'passed': False, 'detail': f'Error analyzing file: {str(e)}'})
    
    passed = score >= 70
    return {'passed': passed, 'score': score, 'checks': checks}

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(json.dumps({'passed': False, 'score': 0.0, 'checks': [{'name': 'Usage', 'passed': False, 'detail': 'Usage: eval_script.py <workspace_dir>'}]}))
        sys.exit(1)
    
    result = eval_financial_model(sys.argv[1])
    print(json.dumps(result))