#!/usr/bin/env python3
import sys
import os
import json
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

def check_presentation(workspace_dir):
    checks = []
    passed = True
    score = 0.0
    max_score = 100.0
    
    pptx_files = [f for f in os.listdir(workspace_dir) if f.endswith('.pptx')]
    
    if not pptx_files:
        checks.append({'name': 'File Creation', 'passed': False, 'detail': 'No PowerPoint file found'})
        return {'passed': False, 'score': 0.0, 'checks': checks}
    
    pptx_file = os.path.join(workspace_dir, pptx_files[0])
    
    try:
        prs = Presentation(pptx_file)
        
        # Check slide count (12.5 points)
        expected_slides = 8
        actual_slides = len(prs.slides)
        if actual_slides >= expected_slides:
            checks.append({'name': 'Slide Count', 'passed': True, 'detail': f'Found {actual_slides} slides (expected at least {expected_slides})'})
            score += 12.5
        else:
            checks.append({'name': 'Slide Count', 'passed': False, 'detail': f'Only {actual_slides} slides found, expected at least {expected_slides}'})
            passed = False
        
        # Load requirements to get verification markers
        req_file = os.path.join(workspace_dir, 'presentation_requirements.json')
        if os.path.exists(req_file):
            with open(req_file, 'r') as f:
                requirements = json.load(f)
            verification_markers = requirements['verification_markers']
        else:
            verification_markers = ['VERIFICATION_MARKER_OVERVIEW', 'VERIFICATION_MARKER_MARKET', 'VERIFICATION_MARKER_FINANCE', 'VERIFICATION_MARKER_PRODUCT', 'VERIFICATION_MARKER_TEAM', 'VERIFICATION_MARKER_PARTNERSHIPS', 'VERIFICATION_MARKER_CHALLENGES', 'VERIFICATION_MARKER_FUTURE']
        
        # Check verification markers in content (25 points)
        markers_found = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    for marker in verification_markers:
                        if marker in shape.text:
                            markers_found += 1
                            break
        
        marker_score = (markers_found / len(verification_markers)) * 25
        score += marker_score
        if markers_found >= len(verification_markers) * 0.75:
            checks.append({'name': 'Content Verification', 'passed': True, 'detail': f'Found {markers_found}/{len(verification_markers)} verification markers'})
        else:
            checks.append({'name': 'Content Verification', 'passed': False, 'detail': f'Only found {markers_found}/{len(verification_markers)} verification markers'})
            passed = False
        
        # Check brand colors usage (25 points)
        brand_colors = {
            (20, 20, 19),    # #141413
            (250, 249, 245), # #faf9f5
            (176, 174, 165), # #b0aea5
            (232, 230, 220), # #e8e6dc
            (217, 119, 87),  # #d97757
            (106, 155, 204), # #6a9bcc
            (120, 140, 93)   # #788c5d
        }
        
        colors_found = set()
        color_applications = 0
        
        for slide in prs.slides:
            for shape in slide.shapes:
                # Check fill colors
                if hasattr(shape, 'fill') and shape.fill.type is not None:
                    try:
                        if hasattr(shape.fill, 'fore_color') and hasattr(shape.fill.fore_color, 'rgb'):
                            rgb = shape.fill.fore_color.rgb
                            color_tuple = (rgb.r, rgb.g, rgb.b)
                            if color_tuple in brand_colors:
                                colors_found.add(color_tuple)
                                color_applications += 1
                    except:
                        pass
                
                # Check text colors
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            try:
                                if hasattr(run.font, 'color') and hasattr(run.font.color, 'rgb'):
                                    rgb = run.font.color.rgb
                                    color_tuple = (rgb.r, rgb.g, rgb.b)
                                    if color_tuple in brand_colors:
                                        colors_found.add(color_tuple)
                                        color_applications += 1
                            except:
                                pass
        
        color_variety_score = min(len(colors_found) / 4, 1.0) * 15  # Up to 15 points for color variety
        color_usage_score = min(color_applications / 20, 1.0) * 10   # Up to 10 points for color usage frequency
        total_color_score = color_variety_score + color_usage_score
        score += total_color_score
        
        if len(colors_found) >= 3 and color_applications >= 10:
            checks.append({'name': 'Brand Colors', 'passed': True, 'detail': f'Found {len(colors_found)} brand colors with {color_applications} applications'})
        else:
            checks.append({'name': 'Brand Colors', 'passed': False, 'detail': f'Insufficient brand color usage: {len(colors_found)} colors, {color_applications} applications'})
            passed = False
        
        # Check typography - font usage (25 points)
        font_applications = {'headings': 0, 'body': 0}
        target_fonts = {'Poppins', 'Lora', 'Arial', 'Georgia'}
        fonts_found = set()
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            try:
                                font_name = run.font.name
                                font_size = run.font.size
                                if font_name and font_name in target_fonts:
                                    fonts_found.add(font_name)
                                    if font_size and font_size >= Pt(24):  # Heading size
                                        if font_name in ['Poppins', 'Arial']:
                                            font_applications['headings'] += 1
                                    else:  # Body text size
                                        if font_name in ['Lora', 'Georgia']:
                                            font_applications['body'] += 1
                            except:
                                pass
        
        typography_score = 0
        if len(fonts_found) >= 2:
            typography_score += 10  # Font variety
        if font_applications['headings'] >= 5:
            typography_score += 7.5  # Proper heading fonts
        if font_applications['body'] >= 10:
            typography_score += 7.5  # Proper body fonts
        
        score += typography_score
        
        if len(fonts_found) >= 2 and (font_applications['headings'] >= 3 or font_applications['body'] >= 5):
            checks.append({'name': 'Typography', 'passed': True, 'detail': f'Found {len(fonts_found)} brand fonts with proper application'})
        else:
            checks.append({'name': 'Typography', 'passed': False, 'detail': f'Insufficient typography compliance: {len(fonts_found)} fonts found'})
            passed = False
        
        # Check content structure - multiple content types (12.5 points)
        shapes_per_slide = []
        text_shapes = 0
        other_shapes = 0
        
        for slide in prs.slides:
            slide_shapes = len(slide.shapes)
            shapes_per_slide.append(slide_shapes)
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    text_shapes += 1
                else:
                    other_shapes += 1
        
        avg_shapes = sum(shapes_per_slide) / len(shapes_per_slide) if shapes_per_slide else 0
        
        if avg_shapes >= 3 and text_shapes >= 15:
            checks.append({'name': 'Content Structure', 'passed': True, 'detail': f'Good content variety: avg {avg_shapes:.1f} shapes/slide, {text_shapes} text elements'})
            score += 12.5
        else:
            checks.append({'name': 'Content Structure', 'passed': False, 'detail': f'Limited content structure: avg {avg_shapes:.1f} shapes/slide, {text_shapes} text elements'})
            passed = False
    
    except Exception as e:
        checks.append({'name': 'File Processing', 'passed': False, 'detail': f'Error processing presentation: {str(e)}'})
        return {'passed': False, 'score': 0.0, 'checks': checks}
    
    final_score = min(score, max_score)
    final_passed = passed and final_score >= 75.0
    
    return {
        'passed': final_passed,
        'score': final_score,
        'checks': checks
    }

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(json.dumps({'passed': False, 'score': 0.0, 'checks': [{'name': 'Usage', 'passed': False, 'detail': 'Usage: eval_script.py <workspace_dir>'}]}))
        sys.exit(1)
    
    result = check_presentation(sys.argv[1])
    print(json.dumps(result))