import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Create a sample presentation with known content
pres = Presentation()
pres.slide_width = Inches(10)
pres.slide_height = Inches(7.5)

# Slide 1: Title slide with marker content
slide_layout = pres.slide_layouts[0]  # Title slide layout
slide = pres.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Q4 Business Review MARKER_TITLE_2024"
subtitle.text = "Strategic Analysis and Future Planning MARKER_SUBTITLE_PLANNING"

# Slide 2: Content slide with bullet points
slide_layout = pres.slide_layouts[1]  # Title and content layout
slide = pres.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Key Performance Indicators MARKER_KPI_SECTION"
text_frame = content.text_frame
text_frame.text = "Revenue Growth MARKER_REVENUE_42_PERCENT"

p = text_frame.add_paragraph()
p.text = "Customer Satisfaction MARKER_SATISFACTION_94_SCORE"
p.level = 0

p = text_frame.add_paragraph()
p.text = "Market Expansion MARKER_EXPANSION_5_REGIONS"
p.level = 0

p = text_frame.add_paragraph()
p.text = "Team Growth MARKER_TEAM_28_MEMBERS"
p.level = 1  # Sub-bullet

# Slide 3: Another content slide
slide_layout = pres.slide_layouts[1]
slide = pres.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Next Quarter Goals MARKER_GOALS_Q1_2025"
text_frame = content.text_frame
text_frame.text = "Launch new product line MARKER_PRODUCT_INNOVATION"

p = text_frame.add_paragraph()
p.text = "Increase market share by 15% MARKER_MARKET_SHARE_TARGET"
p.level = 0

p = text_frame.add_paragraph()
p.text = "Hire 10 additional team members MARKER_HIRING_PLAN_10"
p.level = 0

# Save the presentation
pres.save('quarterly_review.pptx')

print("Generated quarterly_review.pptx with marker content for verification")