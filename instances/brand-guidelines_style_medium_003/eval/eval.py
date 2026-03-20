#!/usr/bin/env python3
import sys
import os
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_brand_styling(workspace_dir):
    checks = []
    score = 0.0
    max_score = 100.0
    
    presentation_path = os.path.join(workspace_dir, 'presentation.pptx')
    
    # Check if presentation exists
    if not os.path.exists(presentation_path):
        checks.append({
            "name": "presentation_exists",
            "passed": False,
            "detail": "presentation.pptx not found"
        })
        return {"passed": False, "score": 0.0, "checks": checks}
    
    checks.append({
        "name": "presentation_exists",
        "passed": True,
        "detail": "presentation.pptx found"
    })
    score += 20
    
    try:
        prs = Presentation(presentation_path)
        
        # Define Anthropic brand colors
        brand_colors = {
            'dark': (20, 20, 19),      # #141413
            'light': (250, 249, 245),  # #faf9f5
            'mid_gray': (176, 174, 165), # #b0aea5
            'light_gray': (232, 230, 220), # #e8e6dc
            'orange': (217, 119, 87),  # #d97757
            'blue': (106, 155, 204),   # #6a9bcc
            'green': (120, 140, 93)    # #788c5d
        }
        
        def color_matches_brand(rgb_color, tolerance=5):
            """Check if a color matches any brand color within tolerance"""
            if rgb_color is None:
                return False
            r, g, b = rgb_color.r, rgb_color.g, rgb_color.b
            for color_name, (br, bg, bb) in brand_colors.items():
                if (abs(r - br) <= tolerance and 
                    abs(g - bg) <= tolerance and 
                    abs(b - bb) <= tolerance):
                    return True, color_name
            return False, None
        
        brand_colors_found = 0
        total_text_elements = 0
        proper_fonts_found = 0
        total_font_elements = 0
        
        # Check slides for brand compliance
        for slide_idx, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame') and shape.text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            total_text_elements += 1
                            
                            # Check font
                            font_name = run.font.name
                            font_size = run.font.size
                            
                            if font_name:
                                total_font_elements += 1
                                # Check if appropriate fonts are used
                                if font_size and font_size.pt >= 24:  # Headings
                                    if font_name in ['Poppins', 'Arial']:
                                        proper_fonts_found += 1
                                else:  # Body text
                                    if font_name in ['Lora', 'Georgia']:
                                        proper_fonts_found += 1
                            
                            # Check text color
                            if run.font.color and run.font.color.rgb:
                                is_brand, color_name = color_matches_brand(run.font.color.rgb)
                                if is_brand:
                                    brand_colors_found += 1
                
                # Check shape colors (non-text elements)
                if hasattr(shape, 'fill') and shape.fill.type == 1:  # Solid fill
                    if shape.fill.fore_color and shape.fill.fore_color.rgb:
                        is_brand, color_name = color_matches_brand(shape.fill.fore_color.rgb)
                        if is_brand and color_name in ['orange', 'blue', 'green']:
                            brand_colors_found += 1
        
        # Check brand colors usage
        brand_color_score = min(30, (brand_colors_found / max(1, total_text_elements)) * 30)
        if brand_colors_found > 0:
            checks.append({
                "name": "brand_colors_applied",
                "passed": True,
                "detail": f"Found {brand_colors_found} brand colors applied"
            })
            score += brand_color_score
        else:
            checks.append({
                "name": "brand_colors_applied",
                "passed": False,
                "detail": "No Anthropic brand colors detected"
            })
        
        # Check typography
        font_score = min(30, (proper_fonts_found / max(1, total_font_elements)) * 30)
        if proper_fonts_found > 0:
            checks.append({
                "name": "brand_typography_applied",
                "passed": True,
                "detail": f"Found {proper_fonts_found} elements with proper brand fonts"
            })
            score += font_score
        else:
            checks.append({
                "name": "brand_typography_applied",
                "passed": False,
                "detail": "No proper brand typography detected"
            })
        
        # Check if original content is preserved
        slide_count = len(prs.slides)
        content_preserved = slide_count >= 3
        if content_preserved:
            checks.append({
                "name": "content_preserved",
                "passed": True,
                "detail": f"Original content preserved ({slide_count} slides)"
            })
            score += 20
        else:
            checks.append({
                "name": "content_preserved",
                "passed": False,
                "detail": "Original slide content may have been lost"
            })
        
    except Exception as e:
        checks.append({
            "name": "presentation_processing",
            "passed": False,
            "detail": f"Error processing presentation: {str(e)}"
        })
        return {"passed": False, "score": 0.0, "checks": checks}
    
    final_score = score / max_score
    passed = final_score >= 0.6  # Pass if 60% or better
    
    return {
        "passed": passed,
        "score": final_score,
        "checks": checks
    }

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: eval_script.py <workspace_directory>")
        sys.exit(1)
    
    result = check_brand_styling(sys.argv[1])
    print(result)