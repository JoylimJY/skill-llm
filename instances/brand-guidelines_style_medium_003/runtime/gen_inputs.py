#!/usr/bin/env python3
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create a presentation with generic formatting
prs = Presentation()

# Slide 1: Title slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Q4 Product Roadmap"
subtitle.text = "Strategic Initiatives and Feature Releases"

# Apply generic formatting
for paragraph in title.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(44)
        run.font.color.rgb = RGBColor(0, 0, 0)

for paragraph in subtitle.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(64, 64, 64)

# Slide 2: Content slide with bullet points
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Key Features"
content.text = "AI Model Improvements\nUser Interface Enhancements\nPerformance Optimizations\nSecurity Updates"

# Apply generic formatting to title
for paragraph in title.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(32)
        run.font.color.rgb = RGBColor(0, 0, 0)

# Apply generic formatting to content
for paragraph in content.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(80, 80, 80)

# Slide 3: Another content slide
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Timeline Overview"
content.text = "October: Research Phase\nNovember: Development Sprint\nDecember: Testing and QA\nJanuary: Release Preparation"

# Apply generic formatting
for paragraph in title.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(28)
        run.font.color.rgb = RGBColor(0, 0, 0)

for paragraph in content.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(100, 100, 100)

# Add a shape for testing accent color application
shape = slide.shapes.add_shape(
    1,  # Rectangle
    Inches(6), Inches(3),
    Inches(2), Inches(1)
)
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(128, 128, 128)  # Generic gray

# Save the presentation
prs.save('presentation.pptx')
print('Created presentation.pptx with generic formatting')