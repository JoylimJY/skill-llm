#!/usr/bin/env python3
"""
Generate the sandbox workspace for the LP Quarterly Briefing task.
Creates a realistic PPTX template with placeholder text and a content brief JSON.
"""

import os
import json
import random
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import zipfile
import shutil

random.seed(42)

WORKSPACE = Path("/workspace")

# ── Distractor files to simulate a real messy workspace ──────────────────────
distractor_structure = {
    "archive/2023/Q3/draft_deck.pptx.bak": b"PK\x03\x04" + b"\x00" * 100,
    "archive/2023/Q4/notes.txt": b"Old Q4 notes - ignore\nDraft only\n",
    "archive/2024/Q1/raw_data.csv": b"fund,value\nFund A,12000000\nFund B,8500000\n",
    "archive/2024/Q1/portfolio_summary.xlsx.bak": b"PK\x03\x04" + b"\x00" * 80,
    "internal/compliance/disclosure_template.txt": b"CONFIDENTIAL - FOR LP USE ONLY\nThis material contains forward-looking statements.\n",
    "internal/compliance/sign_off_checklist.md": b"# Sign-off Checklist\n- [ ] Legal review\n- [ ] CFO approval\n",
    "internal/design/brand_colors.json": json.dumps({
        "primary": "#1B2A4A", "secondary": "#C8A951", "accent": "#FFFFFF"
    }).encode(),
    "internal/design/font_guide.txt": b"Header: Georgia 40pt\nBody: Calibri 14pt\n",
    "data/portfolio/holdings_Q2_2024.csv": (
        b"company,sector,invested,current_value,irr\n"
        b"NovaTech,SaaS,5000000,14200000,0.42\n"
        b"GreenFleet,CleanTech,3000000,7100000,0.31\n"
        b"MedPulse,HealthTech,4000000,6800000,0.19\n"
    ),
    "data/portfolio/exits_history.json": json.dumps([
        {"company": "DataBridge", "exit_year": 2022, "multiple": 3.2},
        {"company": "LogiChain", "exit_year": 2023, "multiple": 4.1},
    ]).encode(),
    "data/benchmarks/cambridge_irr_2024.txt": b"Cambridge Associates Benchmark IRR (Vintage 2019): 18.2%\n",
    "scripts/format_report.py": b"# Legacy report formatter - deprecated\nimport sys\nprint('deprecated')\n",
    "scripts/export_csv.sh": b"#!/bin/bash\necho 'CSV export stub'\n",
}

for rel_path, content in distractor_structure.items():
    full_path = WORKSPACE / rel_path
    full_path.parent.mkdir(parents=True, exist_ok=True)
    full_path.write_bytes(content)

print("✓ Distractor files created")

# ── Content Brief JSON ────────────────────────────────────────────────────────
content_brief = {
    "fund_name": "Arcturus Ventures Fund III",
    "reporting_period": "Q2 2024",
    "report_date": "July 15, 2024",
    "fund_metrics": {
        "aum": "$312M",
        "vintage_year": 2019,
        "fund_irr": "27.4%",
        "benchmark_irr": "18.2%",
        "tvpi": "1.8x",
        "dpi": "0.6x",
        "rvpi": "1.2x"
    },
    "portfolio_highlights": [
        {
            "company": "NovaTech",
            "sector": "Enterprise SaaS",
            "status": "Strong performer — ARR grew 68% YoY to $22M",
            "current_multiple": "2.84x"
        },
        {
            "company": "GreenFleet",
            "sector": "CleanTech Mobility",
            "status": "Fleet contracts signed with 3 major logistics providers",
            "current_multiple": "2.37x"
        },
        {
            "company": "MedPulse",
            "sector": "HealthTech",
            "status": "FDA 510(k) clearance received in June 2024",
            "current_multiple": "1.70x"
        }
    ],
    "capital_deployment": {
        "committed": "$280M",
        "deployed": "$231M",
        "reserved": "$49M",
        "new_investments_this_quarter": 1,
        "new_company": "ClearRoute (autonomous logistics routing)"
    },
    "key_risks": [
        "Rising interest rates compressing SaaS multiples",
        "GreenFleet supply chain delays (ETA resolution Q3 2024)",
        "Regulatory uncertainty in HealthTech reimbursement"
    ],
    "upcoming_milestones": [
        "NovaTech Series C fundraise — expected Q3 2024",
        "GreenFleet fleet deployment — 500 units by end of Q3",
        "MedPulse commercial launch — Q4 2024"
    ],
    "lp_note": "The fund continues to outperform the Cambridge Associates benchmark by 920bps. Capital calls for the new ClearRoute investment ($12M) will be issued in August 2024."
}

brief_path = WORKSPACE / "content_brief.json"
brief_path.write_text(json.dumps(content_brief, indent=2))
print("✓ Content brief created")

# ── Create PPTX Template with placeholder text ────────────────────────────────
# We'll build a template PPTX with 5 slides containing placeholder text
# that the agent must replace. We use python-pptx to create it.

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)

blank_layout = prs.slide_layouts[6]  # blank

def add_rect(slide, l, t, w, h, fill_hex):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    r, g, b = int(fill_hex[0:2], 16), int(fill_hex[2:4], 16), int(fill_hex[4:6], 16)
    shape.fill.fore_color.rgb = RGBColor(r, g, b)
    shape.line.fill.background()
    return shape

def add_textbox(slide, l, t, w, h, text, size_pt, bold=False, color_hex="FFFFFF", align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    r, g, b = int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16)
    run.font.color.rgb = RGBColor(r, g, b)
    return txBox

# ── Slide 1: Title Slide ──────────────────────────────────────────────────────
slide1 = prs.slides.add_slide(blank_layout)
add_rect(slide1, 0, 0, 10, 5.625, "1B2A4A")
add_textbox(slide1, 0.6, 1.2, 8.5, 1.0,
            "XXXX_FUND_NAME XXXX_FUND_NUMBER",
            40, bold=True, color_hex="C8A951", align=PP_ALIGN.LEFT)
add_textbox(slide1, 0.6, 2.3, 7, 0.6,
            "Limited Partner Quarterly Update — XXXX_REPORTING_PERIOD",
            20, bold=False, color_hex="CADCFC")
add_textbox(slide1, 0.6, 3.1, 5, 0.5,
            "XXXX_REPORT_DATE",
            14, color_hex="FFFFFF")
add_textbox(slide1, 0.6, 4.8, 4, 0.4,
            "CONFIDENTIAL",
            10, color_hex="C8A951")

# ── Slide 2: Fund Performance Overview ────────────────────────────────────────
slide2 = prs.slides.add_slide(blank_layout)
add_rect(slide2, 0, 0, 10, 0.7, "1B2A4A")
add_textbox(slide2, 0.4, 0.1, 9, 0.55,
            "Fund Performance Overview — This slide layout reserved for performance data",
            18, bold=True, color_hex="FFFFFF")
# Stat boxes (placeholders)
positions = [(0.4, 1.1), (2.8, 1.1), (5.2, 1.1), (7.6, 1.1)]
labels_vals = [
    ("Fund IRR", "XXXX_IRR%"),
    ("TVPI", "XXXX_TVPIx"),
    ("DPI", "XXXX_DPIx"),
    ("AUM", "XXXX_AUM"),
]
for (lx, ly), (lbl, val) in zip(positions, labels_vals):
    add_rect(slide2, lx, ly, 2.1, 1.8, "243B5C")
    add_textbox(slide2, lx + 0.1, ly + 0.15, 1.9, 0.9,
                val, 28, bold=True, color_hex="C8A951", align=PP_ALIGN.CENTER)
    add_textbox(slide2, lx + 0.1, ly + 1.0, 1.9, 0.5,
                lbl, 11, color_hex="CADCFC", align=PP_ALIGN.CENTER)

add_textbox(slide2, 0.4, 3.2, 9, 0.5,
            "Benchmark IRR: XXXX_BENCHMARK_IRR% (Cambridge Associates, Vintage XXXX_VINTAGE)",
            14, color_hex="444444")
add_textbox(slide2, 0.4, 3.9, 9, 1.2,
            "XXXX_LP_NOTE",
            13, color_hex="333333")

# ── Slide 3: Portfolio Highlights ─────────────────────────────────────────────
slide3 = prs.slides.add_slide(blank_layout)
add_rect(slide3, 0, 0, 10, 0.7, "1B2A4A")
add_textbox(slide3, 0.4, 0.1, 9, 0.55,
            "Portfolio Highlights — This slide layout reserved for portfolio companies",
            18, bold=True, color_hex="FFFFFF")

# Three company cards
card_x = [0.3, 3.6, 6.9]
for i, cx in enumerate(card_x):
    add_rect(slide3, cx, 0.9, 3.0, 3.8, "F0F4F8")
    add_rect(slide3, cx, 0.9, 0.12, 3.8, "C8A951")
    add_textbox(slide3, cx + 0.25, 1.0, 2.6, 0.5,
                f"XXXX_COMPANY_{i+1}", 16, bold=True, color_hex="1B2A4A")
    add_textbox(slide3, cx + 0.25, 1.55, 2.6, 0.4,
                f"XXXX_SECTOR_{i+1}", 11, color_hex="666666")
    add_textbox(slide3, cx + 0.25, 2.05, 2.6, 1.8,
                f"XXXX_STATUS_{i+1}", 12, color_hex="333333")
    add_textbox(slide3, cx + 0.25, 4.15, 2.6, 0.4,
                f"Multiple: XXXX_MULTIPLE_{i+1}", 12, bold=True, color_hex="1B2A4A")

# ── Slide 4: Capital Deployment & Pipeline ─────────────────────────────────────
slide4 = prs.slides.add_slide(blank_layout)
add_rect(slide4, 0, 0, 10, 0.7, "1B2A4A")
add_textbox(slide4, 0.4, 0.1, 9, 0.55,
            "Capital Deployment — This slide layout reserved for capital summary",
            18, bold=True, color_hex="FFFFFF")

add_textbox(slide4, 0.4, 0.9, 4.3, 0.5,
            "Capital Summary", 16, bold=True, color_hex="1B2A4A")
rows = [
    ("Committed Capital", "XXXX_COMMITTED"),
    ("Deployed Capital", "XXXX_DEPLOYED"),
    ("Reserved Capital", "XXXX_RESERVED"),
]
for j, (label, val) in enumerate(rows):
    y = 1.5 + j * 0.6
    add_rect(slide4, 0.4, y, 4.3, 0.5, "F0F4F8" if j % 2 == 0 else "FFFFFF")
    add_textbox(slide4, 0.5, y + 0.05, 2.5, 0.4, label, 13, color_hex="333333")
    add_textbox(slide4, 3.0, y + 0.05, 1.5, 0.4, val, 13, bold=True, color_hex="1B2A4A")

add_textbox(slide4, 5.0, 0.9, 4.5, 0.5,
            "Key Risks", 16, bold=True, color_hex="1B2A4A")
risks_placeholder = ["XXXX_RISK_1", "XXXX_RISK_2", "XXXX_RISK_3"]
for j, rp in enumerate(risks_placeholder):
    add_textbox(slide4, 5.0, 1.5 + j * 0.7, 4.5, 0.6, f"• {rp}", 12, color_hex="333333")

# ── Slide 5: Upcoming Milestones ───────────────────────────────────────────────
slide5 = prs.slides.add_slide(blank_layout)
add_rect(slide5, 0, 0, 10, 0.7, "1B2A4A")
add_textbox(slide5, 0.4, 0.1, 9, 0.55,
            "Upcoming Milestones — This slide layout reserved for milestones",
            18, bold=True, color_hex="FFFFFF")

milestones_ph = ["XXXX_MILESTONE_1", "XXXX_MILESTONE_2", "XXXX_MILESTONE_3"]
for j, mph in enumerate(milestones_ph):
    y = 1.0 + j * 1.2
    add_rect(slide5, 0.4, y, 0.5, 0.5, "C8A951")
    add_textbox(slide5, 0.5, y + 0.05, 0.3, 0.4, str(j + 1), 16, bold=True, color_hex="1B2A4A", align=PP_ALIGN.CENTER)
    add_textbox(slide5, 1.1, y, 8.3, 0.5, mph, 14, color_hex="333333")

add_textbox(slide5, 0.4, 4.8, 9.0, 0.5,
            "CONFIDENTIAL — For Authorized LP Use Only",
            10, color_hex="888888", align=PP_ALIGN.CENTER)

# Save template
template_path = WORKSPACE / "lp_briefing_template.pptx"
prs.save(str(template_path))
print(f"✓ Template PPTX created: {template_path}")

# Verify
assert template_path.exists(), "Template creation failed"
assert template_path.stat().st_size > 10000, "Template too small"

print("\n=== Workspace ready ===")
print(f"Template: {template_path}")
print(f"Brief:    {WORKSPACE / 'content_brief.json'}")
print(f"Distractor files: {len(distractor_structure)}")