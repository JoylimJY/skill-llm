#!/usr/bin/env python3
import sys
import os
import json
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_presentation(workspace_dir):
    checks = []
    
    # Check if PPTX file exists
    pptx_files = [f for f in os.listdir(workspace_dir) if f.endswith('.pptx')]
    if not pptx_files:
        checks.append({"name": "file_exists", "passed": False, "detail": "No .pptx file found"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    checks.append({"name": "file_exists", "passed": True, "detail": f"Found presentation file: {pptx_files[0]}"})
    
    try:
        prs = Presentation(os.path.join(workspace_dir, pptx_files[0]))
        
        # Check slide count (should be 4: title + 3 content slides)
        slide_count = len(prs.slides)
        expected_slides = 4
        slide_count_ok = slide_count >= 3  # At least title + 2 content slides
        checks.append({"name": "slide_count", "passed": slide_count_ok, "detail": f"Found {slide_count} slides (expected at least 3)"})
        
        # Check for Anthropic brand colors
        brand_colors_found = False
        anthropic_colors = {
            (20, 20, 19),    # Dark #141413
            (250, 249, 245), # Light #faf9f5
            (176, 174, 165), # Mid Gray #b0aea5
            (232, 230, 220), # Light Gray #e8e6dc
            (217, 119, 87),  # Orange #d97757
            (106, 155, 204), # Blue #6a9bcc
            (120, 140, 93)   # Green #788c5d
        }
        
        color_found = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'fill') and shape.fill.type == 1:  # Solid fill
                    try:
                        color = shape.fill.fore_color.rgb
                        rgb_tuple = (color.red, color.green, color.blue)
                        if rgb_tuple in anthropic_colors:
                            color_found = True
                            break
                    except:
                        continue
            if color_found:
                break
        
        checks.append({"name": "brand_colors", "passed": color_found, "detail": "Found Anthropic brand colors" if color_found else "No Anthropic brand colors detected"})
        
        # Check for content from the data file
        content_found = False
        slide_text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    slide_text += shape.text.lower()
        
        # Look for key terms from the input data
        key_terms = ["q1", "revenue", "growth", "metrics"]
        terms_found = sum(1 for term in key_terms if term in slide_text)
        content_found = terms_found >= 2
        
        checks.append({"name": "content_relevance", "passed": content_found, "detail": f"Found {terms_found}/4 expected key terms in presentation"})
        
        # Calculate score
        passed_checks = sum(1 for check in checks if check["passed"])
        total_checks = len(checks)
        score = passed_checks / total_checks
        
        return {
            "passed": score >= 0.75,
            "score": score,
            "checks": checks
        }
        
    except Exception as e:
        checks.append({"name": "file_readable", "passed": False, "detail": f"Error reading presentation: {str(e)}"})
        return {"passed": False, "score": 0.0, "checks": checks}

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(json.dumps({"error": "Usage: eval_script.py <workspace_dir>"}))
        sys.exit(1)
    
    result = check_presentation(sys.argv[1])
    print(json.dumps(result, indent=2))