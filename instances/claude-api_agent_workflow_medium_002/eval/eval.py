#!/usr/bin/env python3
import sys
import os
import json
from pathlib import Path
import pandas as pd
from docx import Document

def check_file_exists(filepath):
    return os.path.exists(filepath)

def check_docx_content(filepath):
    try:
        doc = Document(filepath)
        full_text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
        
        # Check for key sections and content
        checks = {
            'has_title': any('comparison' in p.text.lower() or 'report' in p.text.lower() for p in doc.paragraphs),
            'mentions_papers': ('machine learning' in full_text.lower() and 'natural language' in full_text.lower() and 'climate' in full_text.lower()),
            'has_authors': ('chen' in full_text.lower() and 'thompson' in full_text.lower() and 'gonzalez' in full_text.lower()),
            'has_findings': ('94%' in full_text or '89%' in full_text or '15%' in full_text),
            'has_methodology': ('convolutional' in full_text.lower() or 'bert' in full_text.lower() or 'bayesian' in full_text.lower()),
            'has_limitations': 'limitations' in full_text.lower(),
            'has_comparison': ('compare' in full_text.lower() or 'comparison' in full_text.lower()),
            'sufficient_length': len(full_text) > 500
        }
        
        return checks, full_text
    except Exception as e:
        return {}, str(e)

def check_csv_content(filepath):
    try:
        df = pd.read_csv(filepath)
        
        checks = {
            'has_required_columns': all(col in df.columns for col in ['title', 'authors', 'findings', 'methodology', 'limitations']),
            'has_three_rows': len(df) == 3,
            'has_paper_titles': any('Machine Learning' in str(title) for title in df.get('title', [])),
            'has_author_data': any('Chen' in str(authors) for authors in df.get('authors', [])),
            'has_findings_data': any('94%' in str(findings) for findings in df.get('findings', [])),
            'no_empty_cells': not df.isnull().any().any()
        }
        
        return checks, df.to_string()
    except Exception as e:
        return {}, str(e)

def main():
    if len(sys.argv) != 2:
        print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "usage", "passed": False, "detail": "Usage: eval_script.py <workspace_dir>"}]}))
        return
    
    workspace_dir = sys.argv[1]
    checks_results = []
    
    # Look for Word document
    docx_files = list(Path(workspace_dir).glob('*.docx'))
    if not docx_files:
        checks_results.append({"name": "docx_exists", "passed": False, "detail": "No .docx file found"})
        docx_score = 0.0
    else:
        docx_file = docx_files[0]
        checks_results.append({"name": "docx_exists", "passed": True, "detail": f"Found {docx_file.name}"})
        
        docx_checks, content = check_docx_content(docx_file)
        for check_name, passed in docx_checks.items():
            checks_results.append({"name": f"docx_{check_name}", "passed": passed, "detail": f"DOCX {check_name}: {'✓' if passed else '✗'}"})
        
        docx_score = sum(docx_checks.values()) / len(docx_checks) if docx_checks else 0.0
    
    # Look for CSV file
    csv_files = list(Path(workspace_dir).glob('*.csv'))
    if not csv_files:
        checks_results.append({"name": "csv_exists", "passed": False, "detail": "No .csv file found"})
        csv_score = 0.0
    else:
        csv_file = csv_files[0]
        checks_results.append({"name": "csv_exists", "passed": True, "detail": f"Found {csv_file.name}"})
        
        csv_checks, content = check_csv_content(csv_file)
        for check_name, passed in csv_checks.items():
            checks_results.append({"name": f"csv_{check_name}", "passed": passed, "detail": f"CSV {check_name}: {'✓' if passed else '✗'}"})
        
        csv_score = sum(csv_checks.values()) / len(csv_checks) if csv_checks else 0.0
    
    # Calculate overall score
    overall_score = (docx_score + csv_score) / 2
    passed = overall_score >= 0.7
    
    result = {
        "passed": passed,
        "score": round(overall_score, 2),
        "checks": checks_results
    }
    
    print(json.dumps(result, indent=2))

if __name__ == "__main__":
    main()