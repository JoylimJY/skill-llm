#!/usr/bin/env python3
import sys
import os
import json
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_brand_styling(workspace_dir):
    checks = []
    passed_count = 0
    
    # Check if the styled presentation exists
    styled_file = os.path.join(workspace_dir, 'quarterly_report.pptx')
    if not os.path.exists(styled_file):
        checks.append({
            "name": "output_file_exists",
            "passed": False,
            "detail": "quarterly_report.pptx not found"
        })
        return {"passed": False, "score": 0.0, "checks": checks}
    
    checks.append({
        "name": "output_file_exists",
        "passed": True,
        "detail": "quarterly_report.pptx found"
    })
    passed_count += 1
    
    try:
        # Load the presentation
        prs = Presentation(styled_file)
        
        # Check if presentation has expected number of slides
        if len(prs.slides) >= 3:
            checks.append({
                "name": "slide_count",
                "passed": True,
                "detail": f"Found {len(prs.slides)} slides"
            })
            passed_count += 1
        else:
            checks.append({
                "name": "slide_count",
                "passed": False,
                "detail": f"Expected at least 3 slides, found {len(prs.slides)}"
            })
        
        # Check for brand colors usage (Dark, Light, Orange, Blue, Green)
        brand_colors = {
            'dark': (20, 20, 19),     # #141413
            'light': (250, 249, 245), # #faf9f5
            'orange': (217, 119, 87), # #d97757
            'blue': (106, 155, 204),  # #6a9bcc
            'green': (120, 140, 93),  # #788c5d
            'mid_gray': (176, 174, 165), # #b0aea5
            'light_gray': (232, 230, 220) # #e8e6dc
        }
        
        brand_colors_found = False
        font_styling_found = False
        
        # Check slides for styling changes
        for i, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    # Check for font changes in text
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if hasattr(run.font, 'name') and run.font.name:
                                if 'poppins' in run.font.name.lower() or 'lora' in run.font.name.lower():
                                    font_styling_found = True
                
                # Check for brand color usage
                if hasattr(shape, 'fill'):
                    try:
                        if shape.fill.type == 1:  # Solid fill
                            color = shape.fill.fore_color.rgb
                            color_tuple = (color.red, color.green, color.blue)
                            for brand_color_name, brand_color_rgb in brand_colors.items():
                                if abs(color_tuple[0] - brand_color_rgb[0]) <= 5 and \
                                   abs(color_tuple[1] - brand_color_rgb[1]) <= 5 and \
                                   abs(color_tuple[2] - brand_color_rgb[2]) <= 5:
                                    brand_colors_found = True
                                    break
                    except:
                        pass
        
        if brand_colors_found:
            checks.append({
                "name": "brand_colors_applied",
                "passed": True,
                "detail": "Anthropic brand colors detected in presentation"
            })
            passed_count += 1
        else:
            checks.append({
                "name": "brand_colors_applied",
                "passed": False,
                "detail": "No Anthropic brand colors detected"
            })
        
        # Check for content preservation
        content_preserved = False
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text:
                    if 'Q4 Metrics Report' in shape.text or 'Key Performance' in shape.text:
                        content_preserved = True
                        break
        
        if content_preserved:
            checks.append({
                "name": "content_preserved",
                "passed": True,
                "detail": "Original content preserved"
            })
            passed_count += 1
        else:
            checks.append({
                "name": "content_preserved",
                "passed": False,
                "detail": "Original content not found"
            })
            
    except Exception as e:
        checks.append({
            "name": "file_processing",
            "passed": False,
            "detail": f"Error processing presentation: {str(e)}"
        })
    
    total_checks = len(checks)
    score = passed_count / total_checks if total_checks > 0 else 0.0
    passed = score >= 0.6  # Pass if at least 60% of checks pass
    
    return {
        "passed": passed,
        "score": score,
        "checks": checks
    }

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "args", "passed": False, "detail": "Usage: eval_script.py <workspace_dir>"}]}))
        sys.exit(1)
    
    result = check_brand_styling(sys.argv[1])
    print(json.dumps(result))