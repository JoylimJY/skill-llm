#!/usr/bin/env python3
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Create a basic PowerPoint presentation with known content
prs = Presentation()

# Slide 1: Title slide
slide_layout = prs.slide_layouts[0]  # Title slide layout
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]

title.text = "Q4 Metrics Report"
subtitle.text = "Performance Analysis and Key Insights"

# Slide 2: Content slide with bullet points
slide_layout = prs.slide_layouts[1]  # Title and content layout
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Key Performance Indicators"
content.text = "Revenue Growth\nCustomer Acquisition\nMarket Expansion\nProduct Development"

# Slide 3: Another content slide
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]

title.text = "Strategic Initiatives"
content.text = "Digital Transformation\nOperational Efficiency\nTeam Development\nSustainability Goals"

# Add a shape for testing accent colors
shape = slide.shapes.add_shape(
    1,  # Rectangle
    Inches(6), Inches(2),
    Inches(2), Inches(1)
)

# Save the presentation
prs.save('quarterly_report.pptx')
print("Generated quarterly_report.pptx with 3 slides")