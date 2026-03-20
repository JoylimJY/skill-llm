#!/usr/bin/env python3
import sys
import json
from pathlib import Path
import subprocess

def main(workspace_dir):
    workspace = Path(workspace_dir)
    target_file = workspace / 'monthly_budget.xlsx'
    
    checks = []
    
    # Check if file exists
    if not target_file.exists():
        checks.append({
            'name': 'file_exists',
            'passed': False,
            'detail': 'monthly_budget.xlsx file not found'
        })
        return {'passed': False, 'score': 0.0, 'checks': checks}
    
    checks.append({
        'name': 'file_exists',
        'passed': True,
        'detail': 'monthly_budget.xlsx file created successfully'
    })
    
    try:
        # Import openpyxl to read the file
        from openpyxl import load_workbook
        
        wb = load_workbook(str(target_file))
        ws = wb.active
        
        # Check for required headers
        required_headers = ['Category', 'Budgeted Amount', 'Actual Amount', 'Difference']
        header_row = [cell.value for cell in ws[1] if cell.value]
        
        headers_found = 0
        for header in required_headers:
            if any(header.lower() in str(h).lower() for h in header_row if h):
                headers_found += 1
        
        checks.append({
            'name': 'required_headers',
            'passed': headers_found >= 3,
            'detail': f'Found {headers_found}/4 required headers (Category, Budgeted Amount, Actual Amount, Difference)'
        })
        
        # Check for sample categories
        sample_categories = ['groceries', 'rent', 'utilities', 'entertainment']
        categories_found = 0
        
        for row in ws.iter_rows(min_row=2, max_row=10):
            cell_value = row[0].value
            if cell_value:
                for category in sample_categories:
                    if category.lower() in str(cell_value).lower():
                        categories_found += 1
                        break
        
        checks.append({
            'name': 'sample_categories',
            'passed': categories_found >= 3,
            'detail': f'Found {categories_found}/4 sample categories (Groceries, Rent, Utilities, Entertainment)'
        })
        
        # Check for formulas (total row with SUM)
        formula_found = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                    if 'SUM' in cell.value.upper():
                        formula_found = True
                        break
            if formula_found:
                break
        
        checks.append({
            'name': 'formulas_present',
            'passed': formula_found,
            'detail': 'Found SUM formula for totals' if formula_found else 'No SUM formulas found for totals'
        })
        
        # Check for basic formatting (bold headers or colored cells)
        formatting_found = False
        for cell in ws[1]:
            if cell.font and cell.font.bold:
                formatting_found = True
                break
            if cell.fill and cell.fill.start_color.index != '00000000':
                formatting_found = True
                break
        
        checks.append({
            'name': 'basic_formatting',
            'passed': formatting_found,
            'detail': 'Professional formatting applied' if formatting_found else 'No special formatting detected'
        })
        
        wb.close()
        
        # Calculate score
        passed_checks = sum(1 for check in checks if check['passed'])
        total_checks = len(checks)
        score = passed_checks / total_checks
        
        return {
            'passed': score >= 0.6,  # Pass if at least 60% of checks pass
            'score': score,
            'checks': checks
        }
        
    except Exception as e:
        checks.append({
            'name': 'file_readable',
            'passed': False,
            'detail': f'Error reading Excel file: {str(e)}'
        })
        
        return {
            'passed': False,
            'score': 0.0,
            'checks': checks
        }

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(json.dumps({'error': 'Usage: python eval_script.py <workspace_dir>'}))
        sys.exit(1)
    
    result = main(sys.argv[1])
    print(json.dumps(result, indent=2))