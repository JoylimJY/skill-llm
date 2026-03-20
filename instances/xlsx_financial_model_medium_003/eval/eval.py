#!/usr/bin/env python3

import json
import sys
import os
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

def evaluate_dcf_model(workspace_dir):
    checks = []
    score = 0.0
    
    # Find Excel file
    excel_files = list(Path(workspace_dir).glob('*.xlsx'))
    if not excel_files:
        checks.append({"name": "Excel file exists", "passed": False, "detail": "No .xlsx file found"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    excel_file = excel_files[0]
    checks.append({"name": "Excel file exists", "passed": True, "detail": f"Found {excel_file.name}"})
    score += 10
    
    try:
        # Load workbook with formulas
        wb_formulas = load_workbook(excel_file, data_only=False)
        # Load workbook with values
        wb_values = load_workbook(excel_file, data_only=True)
        
        sheet_names = wb_values.sheetnames
        main_sheet = wb_values.active
        
        # Check 1: Revenue projection (2024: 50M, growing 15%)
        revenue_found = False
        revenue_values = []
        for row in main_sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, (int, float)):
                    if abs(cell.value - 50) < 1:  # 2024 base revenue
                        revenue_found = True
                        # Check next few cells for growth pattern
                        row_num, col_num = cell.row, cell.column
                        for i in range(5):
                            next_cell = main_sheet.cell(row_num, col_num + i)
                            if next_cell.value and isinstance(next_cell.value, (int, float)):
                                revenue_values.append(next_cell.value)
                        break
            if revenue_found:
                break
        
        if revenue_found and len(revenue_values) >= 3:
            # Check growth pattern (approximately 15%)
            growth_check = True
            for i in range(1, min(3, len(revenue_values))):
                if revenue_values[i] > 0 and revenue_values[i-1] > 0:
                    growth_rate = (revenue_values[i] / revenue_values[i-1]) - 1
                    if abs(growth_rate - 0.15) > 0.02:  # Allow 2% tolerance
                        growth_check = False
                        break
            
            if growth_check:
                checks.append({"name": "Revenue growth pattern", "passed": True, "detail": "15% growth pattern detected"})
                score += 15
            else:
                checks.append({"name": "Revenue growth pattern", "passed": False, "detail": "Growth pattern doesn't match 15%"})
        else:
            checks.append({"name": "Revenue projection", "passed": False, "detail": "Base revenue of $50M not found"})
        
        # Check 2: Formulas present (not hardcoded values)
        formula_count = 0
        for sheet_name in wb_formulas.sheetnames:
            ws = wb_formulas[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                        formula_count += 1
        
        if formula_count >= 20:
            checks.append({"name": "Formula usage", "passed": True, "detail": f"Found {formula_count} formulas"})
            score += 20
        else:
            checks.append({"name": "Formula usage", "passed": False, "detail": f"Only {formula_count} formulas found, expected 20+"})
        
        # Check 3: Sensitivity analysis table
        sensitivity_found = False
        for sheet_name in wb_values.sheetnames:
            ws = wb_values[sheet_name]
            wacc_values = []
            terminal_values = []
            
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, (int, float)):
                        val = cell.value
                        if 0.09 <= val <= 0.16:  # WACC range 9-16%
                            wacc_values.append(val)
                        elif 0.01 <= val <= 0.06:  # Terminal growth 1-6%
                            terminal_values.append(val)
            
            if len(wacc_values) >= 3 and len(terminal_values) >= 3:
                sensitivity_found = True
                break
        
        if sensitivity_found:
            checks.append({"name": "Sensitivity analysis", "passed": True, "detail": "WACC and terminal growth ranges found"})
            score += 20
        else:
            checks.append({"name": "Sensitivity analysis", "passed": False, "detail": "Sensitivity table not detected"})
        
        # Check 4: Key DCF components present
        dcf_components = 0
        search_terms = ['EBITDA', 'FCF', 'Terminal', 'NPV', 'DCF', 'WACC']
        for sheet_name in wb_values.sheetnames:
            ws = wb_values[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        for term in search_terms:
                            if term.lower() in cell.value.lower():
                                dcf_components += 1
                                break
        
        if dcf_components >= 4:
            checks.append({"name": "DCF components", "passed": True, "detail": f"Found {dcf_components} DCF-related terms"})
            score += 15
        else:
            checks.append({"name": "DCF components", "passed": False, "detail": f"Only {dcf_components} DCF terms found"})
        
        # Check 5: Formatting (blue inputs)
        blue_cells = 0
        main_sheet_formulas = wb_formulas.active
        for row in main_sheet_formulas.iter_rows():
            for cell in row:
                if cell.font and cell.font.color:
                    if hasattr(cell.font.color, 'rgb') and cell.font.color.rgb:
                        color = cell.font.color.rgb
                        if color and 'FF0000FF' in str(color).upper():  # Blue color
                            blue_cells += 1
        
        if blue_cells >= 5:
            checks.append({"name": "Blue input formatting", "passed": True, "detail": f"Found {blue_cells} blue-formatted cells"})
            score += 10
        else:
            checks.append({"name": "Blue input formatting", "passed": False, "detail": "Insufficient blue formatting for inputs"})
        
        # Check 6: Zero formatting (dashes)
        dash_formatting = False
        for sheet_name in wb_formulas.sheetnames:
            ws = wb_formulas[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.number_format and '-' in cell.number_format:
                        dash_formatting = True
                        break
                if dash_formatting:
                    break
        
        if dash_formatting:
            checks.append({"name": "Zero dash formatting", "passed": True, "detail": "Dash formatting for zeros detected"})
            score += 10
        else:
            checks.append({"name": "Zero dash formatting", "passed": False, "detail": "No dash formatting for zeros found"})
        
        wb_formulas.close()
        wb_values.close()
        
    except Exception as e:
        checks.append({"name": "File analysis", "passed": False, "detail": f"Error analyzing Excel file: {str(e)}"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    final_score = min(100.0, score)
    passed = final_score >= 70.0
    
    return {
        "passed": passed,
        "score": final_score,
        "checks": checks
    }

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(json.dumps({"error": "Usage: eval_script.py <workspace_dir>"}))
        sys.exit(1)
    
    workspace_dir = sys.argv[1]
    result = evaluate_dcf_model(workspace_dir)
    print(json.dumps(result, indent=2))