import os
import json
from docx import Document
from docx.shared import Inches

# Create a technical document about AI safety
doc = Document()
doc.add_heading('AI Safety in Production Systems', 0)
doc.add_paragraph('Executive Summary: This document outlines critical considerations for deploying AI systems safely in production environments.')

doc.add_heading('Key Concepts', level=1)
doc.add_paragraph('1. Model Alignment: Ensuring AI systems behave according to intended objectives')
doc.add_paragraph('2. Robustness: Systems must handle unexpected inputs gracefully')
doc.add_paragraph('3. Interpretability: Understanding how decisions are made')

doc.add_heading('Key Findings', level=1)
doc.add_paragraph('FINDING_MARKER_1: 73% of production AI failures stem from insufficient input validation')
doc.add_paragraph('FINDING_MARKER_2: Organizations with formal AI governance show 45% fewer incidents')
doc.add_paragraph('FINDING_MARKER_3: Continuous monitoring reduces response time to issues by 60%')

doc.add_heading('Recommendations', level=1)
doc.add_paragraph('RECOMMENDATION_MARKER_1: Implement comprehensive testing frameworks before deployment')
doc.add_paragraph('RECOMMENDATION_MARKER_2: Establish clear escalation procedures for AI system anomalies')
doc.add_paragraph('RECOMMENDATION_MARKER_3: Invest in explainable AI tools for critical decision paths')

doc.save('technical_doc.docx')

# Create a markdown version as well
with open('technical_doc.md', 'w') as f:
    f.write('''# AI Safety in Production Systems

Executive Summary: This document outlines critical considerations for deploying AI systems safely in production environments.

## Key Concepts

1. Model Alignment: Ensuring AI systems behave according to intended objectives
2. Robustness: Systems must handle unexpected inputs gracefully
3. Interpretability: Understanding how decisions are made

## Key Findings

FINDING_MARKER_1: 73% of production AI failures stem from insufficient input validation
FINDING_MARKER_2: Organizations with formal AI governance show 45% fewer incidents
FINDING_MARKER_3: Continuous monitoring reduces response time to issues by 60%

## Recommendations

RECOMMENDATION_MARKER_1: Implement comprehensive testing frameworks before deployment
RECOMMENDATION_MARKER_2: Establish clear escalation procedures for AI system anomalies
RECOMMENDATION_MARKER_3: Invest in explainable AI tools for critical decision paths
''')

# Create expected structure for validation
expected_output = {
    'concepts_found': 3,
    'findings_found': 3,
    'recommendations_found': 3,
    'marker_concepts': ['Model Alignment', 'Robustness', 'Interpretability'],
    'marker_findings': ['FINDING_MARKER_1', 'FINDING_MARKER_2', 'FINDING_MARKER_3'],
    'marker_recommendations': ['RECOMMENDATION_MARKER_1', 'RECOMMENDATION_MARKER_2', 'RECOMMENDATION_MARKER_3']
}

with open('expected_structure.json', 'w') as f:
    json.dump(expected_output, f, indent=2)

print('Generated technical documents with embedded markers for validation')