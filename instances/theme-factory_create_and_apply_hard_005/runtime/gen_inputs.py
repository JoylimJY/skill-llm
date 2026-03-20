import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import random

# Set deterministic seed
random.seed(42)

# Create the base presentation
prs = Presentation()

# Slide 1: Title slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Sustainable Energy Expansion Strategy"
subtitle.text = "Q4 Board Review | MARKER_TITLE_SLIDE | TechGreen Solutions"

# Slide 2: Market overview with bullet points
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Market Opportunity Analysis"
content.text = "• Global renewable energy market projected to reach $1.9T by 2030\n• Solar panel efficiency improvements of 23% year-over-year\n• MARKER_CONTENT_BULLET regulatory incentives driving adoption\n• Key competitor analysis shows market gap in enterprise solutions"

# Slide 3: Financial projections slide
slide_layout = prs.slide_layouts[5]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Financial Projections & Investment Requirements"
# Add a text box with financial data
left = Inches(1)
top = Inches(2)
width = Inches(8)
height = Inches(4)
textbox = slide.shapes.add_textbox(left, top, width, height)
text_frame = textbox.text_frame
text_frame.text = "Revenue Projections (Millions USD)\n\n2024: $45M (Current)\n2025: $78M (73% growth)\n2026: $134M (72% growth)\nMARKER_FINANCIAL_DATA\n\nRequired Investment: $25M\nProjected ROI: 340% by 2026\nBreak-even: Q3 2025"

# Slide 4: Strategic initiatives
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Strategic Implementation Roadmap"
content.text = "Phase 1: Technology Development (Q1-Q2 2024)\n• Advanced battery storage systems\n• MARKER_STRATEGY_PHASE smart grid integration platform\n\nPhase 2: Market Entry (Q3-Q4 2024)\n• Pilot partnerships with 3 Fortune 500 companies\n• Regulatory compliance and certifications\n\nPhase 3: Scale Operations (2025)\n• Manufacturing facility expansion\n• Sales team growth to 50+ professionals"

# Slide 5: Risk assessment and mitigation
slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
title.text = "Risk Assessment & Mitigation Strategies"
# Add content as text box
left = Inches(1)
top = Inches(1.5)
width = Inches(8)
height = Inches(5)
textbox = slide.shapes.add_textbox(left, top, width, height)
text_frame = textbox.text_frame
text_frame.text = "Key Risks Identified:\n\n1. Technology adoption slower than projected\n   → Mitigation: Diversified product portfolio\n\n2. Regulatory changes in renewable incentives\n   → Mitigation: Multi-jurisdictional strategy\n\n3. Supply chain disruptions\n   → Mitigation: MARKER_RISK_CONTENT dual-source suppliers\n\n4. Competitive pressure from established players\n   → Mitigation: Patent protection and innovation focus"

# Save the presentation
prs.save('board_presentation.pptx')

# Create theme files directory structure
os.makedirs('themes', exist_ok=True)

# Create theme files (subset for testing)
with open('themes/tech_innovation.txt', 'w') as f:
    f.write("""Theme: Tech Innovation
Description: Bold and modern tech aesthetic with clean lines and vibrant accents

Colors:
- Primary: #1E3A8A (Deep Blue)
- Secondary: #3B82F6 (Bright Blue) 
- Accent: #10B981 (Emerald Green)
- Background: #F8FAFC (Light Gray)
- Text: #1F2937 (Dark Gray)
- Highlight: #F59E0B (Amber)

Fonts:
- Headers: Calibri Bold
- Body: Calibri Regular
- Accent: Calibri Light

Style Notes:
- High contrast for data visibility
- Clean, modern aesthetic
- Professional yet innovative feel""")

with open('themes/forest_canopy.txt', 'w') as f:
    f.write("""Theme: Forest Canopy
Description: Natural and grounded earth tones representing growth and sustainability

Colors:
- Primary: #065F46 (Forest Green)
- Secondary: #059669 (Green)
- Accent: #34D399 (Light Green)
- Background: #F0FDF4 (Very Light Green)
- Text: #1F2937 (Dark Gray)
- Highlight: #92400E (Earth Brown)

Fonts:
- Headers: Calibri Bold
- Body: Calibri Regular
- Accent: Calibri Light

Style Notes:
- Evokes trust and environmental responsibility
- Excellent for sustainability-focused content
- Professional organic feel""")

with open('themes/ocean_depths.txt', 'w') as f:
    f.write("""Theme: Ocean Depths
Description: Professional and calming maritime theme with deep blues and clean whites

Colors:
- Primary: #1E40AF (Ocean Blue)
- Secondary: #3B82F6 (Sky Blue)
- Accent: #06B6D4 (Cyan)
- Background: #F1F5F9 (Light Blue Gray)
- Text: #0F172A (Very Dark Blue)
- Highlight: #0891B2 (Teal)

Fonts:
- Headers: Calibri Bold
- Body: Calibri Regular
- Accent: Calibri Light

Style Notes:
- Conveys stability and depth
- Excellent readability for financial data
- Trustworthy corporate feel""")

print("Generated board_presentation.pptx and theme files successfully")