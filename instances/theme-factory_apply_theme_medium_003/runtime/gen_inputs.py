#!/usr/bin/env python3

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import json

# Create input presentation with marker content
prs = Presentation()

# Title slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Q4 Sales Results - MARKER_TITLE_CONTENT"
subtitle.text = "Annual Performance Review - MARKER_SUBTITLE_CONTENT"

# Content slide 1
content_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Key Metrics Overview - MARKER_METRICS_TITLE"
content.text = "Revenue Growth: 15%\nCustomer Acquisition: 2,500 new clients\nMarket Share: 23%\nMARKER_METRICS_DATA"

# Content slide 2
slide = prs.slides.add_slide(content_slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Regional Performance - MARKER_REGIONAL_TITLE"
content.text = "North America: $2.3M\nEurope: $1.8M\nAsia Pacific: $1.2M\nMARKER_REGIONAL_DATA"

prs.save('sales_presentation.pptx')

# Create theme files directory
os.makedirs('themes', exist_ok=True)

# Ocean Depths theme
ocean_theme = {
    "name": "Ocean Depths",
    "description": "Professional and calming maritime theme",
    "colors": {
        "primary": "#1e3a5f",
        "secondary": "#2e5984",
        "accent": "#4a90a4",
        "background": "#f0f8ff",
        "text": "#1a1a1a",
        "light_accent": "#87ceeb"
    },
    "fonts": {
        "header": "Calibri",
        "body": "Arial"
    }
}

# Tech Innovation theme
tech_theme = {
    "name": "Tech Innovation",
    "description": "Bold and modern tech aesthetic",
    "colors": {
        "primary": "#0066cc",
        "secondary": "#004499",
        "accent": "#00aaff",
        "background": "#ffffff",
        "text": "#333333",
        "light_accent": "#cce6ff"
    },
    "fonts": {
        "header": "Arial",
        "body": "Calibri"
    }
}

# Modern Minimalist theme
minimal_theme = {
    "name": "Modern Minimalist",
    "description": "Clean and contemporary grayscale",
    "colors": {
        "primary": "#2c2c2c",
        "secondary": "#4a4a4a",
        "accent": "#666666",
        "background": "#ffffff",
        "text": "#1a1a1a",
        "light_accent": "#e6e6e6"
    },
    "fonts": {
        "header": "Calibri",
        "body": "Arial"
    }
}

# Save theme files
with open('themes/ocean_depths.json', 'w') as f:
    json.dump(ocean_theme, f, indent=2)

with open('themes/tech_innovation.json', 'w') as f:
    json.dump(tech_theme, f, indent=2)

with open('themes/modern_minimalist.json', 'w') as f:
    json.dump(minimal_theme, f, indent=2)

print("Generated sales_presentation.pptx and theme files")