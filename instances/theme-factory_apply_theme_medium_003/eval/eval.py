#!/usr/bin/env python3

import sys
import os
import json
from pptx import Presentation
from pptx.dml.color import RGBColor

def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))

def check_theme_applied(workspace_dir):
    checks = []
    
    # Check if presentation file exists
    pptx_path = os.path.join(workspace_dir, 'sales_presentation.pptx')
    if not os.path.exists(pptx_path):
        checks.append({
            "name": "presentation_exists",
            "passed": False,
            "detail": "sales_presentation.pptx not found"
        })
        return checks, False
    
    try:
        prs = Presentation(pptx_path)
        
        # Check that original marker content is preserved
        title_slide = prs.slides[0]
        title_text = title_slide.shapes.title.text
        marker_preserved = "MARKER_TITLE_CONTENT" in title_text
        
        checks.append({
            "name": "content_preserved",
            "passed": marker_preserved,
            "detail": f"Title text: {title_text[:50]}..."
        })
        
        # Check if any theme colors have been applied
        theme_colors_found = False
        professional_colors = [
            (30, 58, 95),   # Ocean Depths primary
            (0, 102, 204),  # Tech Innovation primary
            (44, 44, 44),   # Modern Minimalist primary
        ]
        
        for slide in prs.slides:
            try:
                if hasattr(slide, 'background') and slide.background.fill:
                    bg_color = slide.background.fill.fore_color
                    if hasattr(bg_color, 'rgb'):
                        rgb = (bg_color.rgb.red, bg_color.rgb.green, bg_color.rgb.blue)
                        if rgb in professional_colors:
                            theme_colors_found = True
                            break
                            
                # Check text colors in shapes
                for shape in slide.shapes:
                    if hasattr(shape, 'text_frame'):
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if hasattr(run.font, 'color') and run.font.color.rgb:
                                    rgb = (run.font.color.rgb.red, run.font.color.rgb.green, run.font.color.rgb.blue)
                                    if rgb in professional_colors:
                                        theme_colors_found = True
                                        break
            except Exception:
                continue
                
        checks.append({
            "name": "theme_colors_applied",
            "passed": theme_colors_found,
            "detail": "Professional theme colors detected" if theme_colors_found else "No theme colors found"
        })
        
        # Check if fonts have been modified (look for professional fonts)
        professional_fonts = ["Calibri", "Arial", "Segoe UI", "Helvetica"]
        font_applied = False
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if hasattr(run.font, 'name') and run.font.name in professional_fonts:
                                font_applied = True
                                break
                                
        checks.append({
            "name": "professional_fonts_applied",
            "passed": font_applied,
            "detail": "Professional fonts detected" if font_applied else "No professional fonts found"
        })
        
        # Check slide count is preserved
        slide_count_correct = len(prs.slides) == 3
        checks.append({
            "name": "slide_count_preserved",
            "passed": slide_count_correct,
            "detail": f"Found {len(prs.slides)} slides, expected 3"
        })
        
    except Exception as e:
        checks.append({
            "name": "presentation_readable",
            "passed": False,
            "detail": f"Error reading presentation: {str(e)}"
        })
        return checks, False
    
    passed_count = sum(1 for check in checks if check["passed"])
    total_checks = len(checks)
    overall_passed = passed_count >= 3  # Need at least 3 out of 4 checks to pass
    
    return checks, overall_passed

def main():
    if len(sys.argv) != 2:
        print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "usage", "passed": False, "detail": "Usage: eval_script.py <workspace_dir>"}]}))
        return
    
    workspace_dir = sys.argv[1]
    checks, overall_passed = check_theme_applied(workspace_dir)
    
    score = sum(1 for check in checks if check["passed"]) / len(checks) if checks else 0.0
    
    result = {
        "passed": overall_passed,
        "score": score,
        "checks": checks
    }
    
    print(json.dumps(result))

if __name__ == "__main__":
    main()