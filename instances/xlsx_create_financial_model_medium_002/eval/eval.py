import sys
import os
import json
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

def eval_saas_model(workspace_dir):
    workspace = Path(workspace_dir)
    model_file = workspace / 'saas_model.xlsx'
    markers_file = workspace / 'model_markers.json'
    
    checks = []
    
    # Load expected markers
    try:
        with open(markers_file) as f:
            markers = json.load(f)
    except:
        return {'passed': False, 'score': 0.0, 'checks': [{'name': 'load_markers', 'passed': False, 'detail': 'Could not load markers file'}]}
    
    # Check if file exists
    if not model_file.exists():
        checks.append({'name': 'file_exists', 'passed': False, 'detail': 'saas_model.xlsx not found'})
        return {'passed': False, 'score': 0.0, 'checks': checks}
    
    checks.append({'name': 'file_exists', 'passed': True, 'detail': 'Excel file created'})
    
    try:
        wb = load_workbook(str(model_file), data_only=False)
        
        # Check required sheets
        required_sheets = ['Assumptions', 'Revenue', 'P&L', 'Cash Flow', 'Unit Economics']
        sheet_names = [name.lower() for name in wb.sheetnames]
        
        for sheet in required_sheets:
            found = any(sheet.lower() in name for name in sheet_names)
            checks.append({'name': f'sheet_{sheet.lower()}', 'passed': found, 'detail': f'{sheet} sheet present: {found}'})
        
        # Check assumptions sheet for key values
        assumptions_sheet = None
        for sheet_name in wb.sheetnames:
            if 'assumption' in sheet_name.lower():
                assumptions_sheet = wb[sheet_name]
                break
        
        if assumptions_sheet:
            # Look for CAC value (150)
            cac_found = False
            churn_found = False
            arpu_found = False
            
            for row in assumptions_sheet.iter_rows():
                for cell in row:
                    if cell.value == 150:
                        cac_found = True
                    elif cell.value == 0.025:
                        churn_found = True
                    elif cell.value == 89:
                        arpu_found = True
            
            checks.append({'name': 'cac_assumption', 'passed': cac_found, 'detail': f'CAC value 150 found: {cac_found}'})
            checks.append({'name': 'churn_assumption', 'passed': churn_found, 'detail': f'Churn rate 0.025 found: {churn_found}'})
            checks.append({'name': 'arpu_assumption', 'passed': arpu_found, 'detail': f'ARPU value 89 found: {arpu_found}'})
        
        # Check for formulas (not hardcoded values)
        formula_count = 0
        blue_text_count = 0
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                        formula_count += 1
                    
                    # Check for blue text (inputs)
                    if hasattr(cell, 'font') and cell.font and cell.font.color:
                        if hasattr(cell.font.color, 'rgb') and cell.font.color.rgb == 'FF0000FF':
                            blue_text_count += 1
        
        checks.append({'name': 'has_formulas', 'passed': formula_count >= 20, 'detail': f'Found {formula_count} formulas (need 20+)'})
        
        # Check revenue sheet for monthly structure
        revenue_sheet = None
        for sheet_name in wb.sheetnames:
            if 'revenue' in sheet_name.lower():
                revenue_sheet = wb[sheet_name]
                break
        
        monthly_structure = False
        if revenue_sheet:
            # Look for month indicators
            for row in revenue_sheet.iter_rows(max_row=10):
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        if any(month in cell.value.lower() for month in ['jan', 'feb', 'mar', 'month 1', 'month']):
                            monthly_structure = True
                            break
        
        checks.append({'name': 'monthly_structure', 'passed': monthly_structure, 'detail': f'Monthly structure found: {monthly_structure}'})
        
        # Check P&L sheet for key expense categories
        pl_sheet = None
        for sheet_name in wb.sheetnames:
            if 'p&l' in sheet_name.lower() or 'pl' in sheet_name.lower() or 'income' in sheet_name.lower():
                pl_sheet = wb[sheet_name]
                break
        
        expense_categories = ['cogs', 'sales', 'marketing', 'r&d', 'g&a']
        expense_found = {cat: False for cat in expense_categories}
        
        if pl_sheet:
            for row in pl_sheet.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        cell_text = cell.value.lower()
                        for cat in expense_categories:
                            if cat in cell_text:
                                expense_found[cat] = True
        
        for cat, found in expense_found.items():
            checks.append({'name': f'expense_{cat}', 'passed': found, 'detail': f'{cat.upper()} category found: {found}'})
        
        wb.close()
        
    except Exception as e:
        checks.append({'name': 'file_analysis', 'passed': False, 'detail': f'Error analyzing file: {str(e)}'})
        return {'passed': False, 'score': 0.0, 'checks': checks}
    
    # Calculate score
    passed_checks = sum(1 for check in checks if check['passed'])
    total_checks = len(checks)
    score = passed_checks / total_checks if total_checks > 0 else 0.0
    passed = score >= 0.7
    
    return {
        'passed': passed,
        'score': score,
        'checks': checks
    }

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print('Usage: python eval_script.py <workspace_dir>')
        sys.exit(1)
    
    result = eval_saas_model(sys.argv[1])
    print(json.dumps(result, indent=2))