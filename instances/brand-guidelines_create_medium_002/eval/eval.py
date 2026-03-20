import sys
import os
import json
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_presentation(workspace_dir):
    checks = []
    score = 0.0
    
    # Check if presentation file exists
    pptx_files = [f for f in os.listdir(workspace_dir) if f.endswith('.pptx')]
    if not pptx_files:
        checks.append({"name": "file_exists", "passed": False, "detail": "No PowerPoint file found"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    pptx_file = os.path.join(workspace_dir, pptx_files[0])
    checks.append({"name": "file_exists", "passed": True, "detail": f"Found presentation: {pptx_files[0]}"})
    score += 0.2
    
    try:
        prs = Presentation(pptx_file)
        
        # Check slide count
        slide_count = len(prs.slides)
        if slide_count >= 5:
            checks.append({"name": "slide_count", "passed": True, "detail": f"Has {slide_count} slides (expected 5+)"})
            score += 0.2
        else:
            checks.append({"name": "slide_count", "passed": False, "detail": f"Only {slide_count} slides (expected 5+)"})
        
        # Check for Anthropic brand colors
        brand_colors_found = set()
        anthropic_colors = {
            '#141413': 'dark',
            '#faf9f5': 'light', 
            '#b0aea5': 'mid_gray',
            '#e8e6dc': 'light_gray',
            '#d97757': 'orange',
            '#6a9bcc': 'blue',
            '#788c5d': 'green'
        }
        
        def rgb_to_hex(r, g, b):
            return f'#{r:02x}{g:02x}{b:02x}'
        
        for slide in prs.slides:
            # Check text colors
            for shape in slide.shapes:
                if hasattr(shape, 'text_frame'):
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if hasattr(run.font, 'color') and run.font.color.rgb:
                                rgb = run.font.color.rgb
                                hex_color = rgb_to_hex(rgb.r, rgb.g, rgb.b)
                                if hex_color in anthropic_colors:
                                    brand_colors_found.add(anthropic_colors[hex_color])
                
                # Check shape fill colors
                if hasattr(shape, 'fill') and shape.fill.fore_color.rgb:
                    rgb = shape.fill.fore_color.rgb
                    hex_color = rgb_to_hex(rgb.r, rgb.g, rgb.b)
                    if hex_color in anthropic_colors:
                        brand_colors_found.add(anthropic_colors[hex_color])
        
        if len(brand_colors_found) >= 2:
            checks.append({"name": "brand_colors", "passed": True, "detail": f"Found Anthropic brand colors: {', '.join(brand_colors_found)}"})
            score += 0.3
        else:
            checks.append({"name": "brand_colors", "passed": False, "detail": f"Limited brand colors found: {', '.join(brand_colors_found) if brand_colors_found else 'none'}"})
        
        # Check for expected content
        content_found = []
        expected_content = ['roadmap', 'market', 'features', 'architecture', 'steps']
        
        for slide in prs.slides:
            slide_text = ''
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    slide_text += shape.text.lower() + ' '
            
            for content in expected_content:
                if content in slide_text:
                    content_found.append(content)
        
        content_found = list(set(content_found))
        if len(content_found) >= 4:
            checks.append({"name": "content_coverage", "passed": True, "detail": f"Found expected content: {', '.join(content_found)}"})
            score += 0.2
        else:
            checks.append({"name": "content_coverage", "passed": False, "detail": f"Limited content coverage: {', '.join(content_found)}"})
        
        # Check for title slide
        first_slide_text = ''
        if prs.slides:
            for shape in prs.slides[0].shapes:
                if hasattr(shape, 'text'):
                    first_slide_text += shape.text.lower() + ' '
        
        if 'anthropic' in first_slide_text or 'roadmap' in first_slide_text:
            checks.append({"name": "title_slide", "passed": True, "detail": "Title slide contains expected content"})
            score += 0.1
        else:
            checks.append({"name": "title_slide", "passed": False, "detail": "Title slide missing expected content"})
        
    except Exception as e:
        checks.append({"name": "file_processing", "passed": False, "detail": f"Error processing presentation: {str(e)}"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    passed = score >= 0.7
    return {"passed": passed, "score": score, "checks": checks}

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print(json.dumps({"passed": False, "score": 0.0, "checks": [{"name": "usage", "passed": False, "detail": "Usage: eval_script.py <workspace_dir>"}]}))
        sys.exit(1)
    
    result = check_presentation(sys.argv[1])
    print(json.dumps(result))