import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import random

# Set fixed seed for deterministic output
random.seed(42)

# Create a basic presentation with quarterly sales data
prs = Presentation()

# Slide 1 - Title slide
slide_layout = prs.slide_layouts[0]  # Title slide layout
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Q4 2024 Sales Results"
subtitle.text = "Performance Review and Analysis\nMARKER_TITLE_SLIDE"

# Slide 2 - Content slide
slide_layout = prs.slide_layouts[1]  # Title and content layout
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Key Metrics Overview"
content.text = "• Revenue: $2.5M (+15% YoY)\n• Units Sold: 12,500 (+8% YoY)\n• Market Share: 23% (+2% YoY)\n• Customer Satisfaction: 94%\nMARKER_METRICS_SLIDE"

# Slide 3 - Another content slide
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
content = slide.placeholders[1]
title.text = "Regional Performance"
content.text = "• North America: $1.2M\n• Europe: $800K\n• Asia Pacific: $500K\n• Growth opportunities identified\nMARKER_REGIONAL_SLIDE"

# Save the presentation
prs.save('quarterly_sales.pptx')

# Create theme showcase PDF (simplified version for testing)
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.colors import HexColor

c = canvas.Canvas('theme-showcase.pdf', pagesize=letter)
c.setFont('Helvetica-Bold', 24)
c.drawString(50, 750, 'Theme Showcase')
c.setFont('Helvetica', 14)
c.drawString(50, 720, 'Available themes for your presentation:')

themes = [
    'Ocean Depths - Professional maritime theme',
    'Sunset Boulevard - Warm vibrant colors',
    'Forest Canopy - Natural earth tones',
    'Modern Minimalist - Clean contemporary grayscale',
    'Golden Hour - Rich autumnal palette'
]

y_pos = 680
for theme in themes:
    c.drawString(70, y_pos, f'• {theme}')
    y_pos -= 30

c.drawString(50, 400, 'MARKER_THEME_SHOWCASE')
c.save()

# Create themes directory and sample theme files
os.makedirs('themes', exist_ok=True)

# Ocean Depths theme
with open('themes/ocean_depths.txt', 'w') as f:
    f.write("Theme: Ocean Depths\n")
    f.write("Primary Color: #1B4D72\n")
    f.write("Secondary Color: #2E86AB\n")
    f.write("Accent Color: #A23B72\n")
    f.write("Background Color: #F8FBFF\n")
    f.write("Text Color: #1B4D72\n")
    f.write("Header Font: Calibri Bold\n")
    f.write("Body Font: Calibri\n")
    f.write("MARKER_OCEAN_THEME\n")

# Modern Minimalist theme
with open('themes/modern_minimalist.txt', 'w') as f:
    f.write("Theme: Modern Minimalist\n")
    f.write("Primary Color: #2D3748\n")
    f.write("Secondary Color: #4A5568\n")
    f.write("Accent Color: #E53E3E\n")
    f.write("Background Color: #FFFFFF\n")
    f.write("Text Color: #2D3748\n")
    f.write("Header Font: Arial Bold\n")
    f.write("Body Font: Arial\n")
    f.write("MARKER_MINIMALIST_THEME\n")

print("Generated input files: quarterly_sales.pptx, theme-showcase.pdf, and theme files")