#!/usr/bin/env python3

import sys
import os
import json
from pathlib import Path
try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except ImportError:
    print('{"passed": false, "score": 0.0, "checks": [{"name": "dependencies", "passed": false, "detail": "Missing python-pptx dependency"}]}')
    sys.exit(1)

def check_color_match(color_obj, expected_hex):
    """Check if a color object matches expected hex value"""
    if hasattr(color_obj, 'rgb'):
        rgb = color_obj.rgb
        hex_val = f'#{rgb.red:02x}{rgb.green:02x}{rgb.blue:02x}'
        return hex_val.lower() == expected_hex.lower()
    return False

def evaluate_presentation(workspace_path):
    checks = []
    score = 0.0
    
    # Check if presentation file exists
    pptx_files = list(Path(workspace_path).glob('*.pptx'))
    if not pptx_files:
        checks.append({"name": "file_exists", "passed": False, "detail": "No PowerPoint file found"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    pptx_file = pptx_files[0]
    checks.append({"name": "file_exists", "passed": True, "detail": f"Found presentation: {pptx_file.name}"})
    score += 10.0
    
    try:
        prs = Presentation(str(pptx_file))
        
        # Check minimum slide count (5 required)
        slide_count = len(prs.slides)
        if slide_count >= 5:
            checks.append({"name": "slide_count", "passed": True, "detail": f"Has {slide_count} slides (minimum 5 required)"})
            score += 15.0
        else:
            checks.append({"name": "slide_count", "passed": False, "detail": f"Only {slide_count} slides, need minimum 5"})
        
        # Check for Anthropic brand colors
        brand_colors = {
            '#141413': 'Dark',
            '#faf9f5': 'Light', 
            '#b0aea5': 'Mid Gray',
            '#e8e6dc': 'Light Gray',
            '#d97757': 'Orange',
            '#6a9bcc': 'Blue',
            '#788c5d': 'Green'
        }
        
        colors_found = set()
        text_content = ''
        
        for slide_idx, slide in enumerate(prs.slides):
            # Check text content and colors
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    text_content += shape.text.lower() + ' '
                    
                    # Check text formatting for brand fonts
                    if hasattr(shape, 'text_frame'):
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if hasattr(run.font, 'name'):
                                    font_name = run.font.name
                                    if font_name and ('poppins' in font_name.lower() or 'lora' in font_name.lower()):
                                        colors_found.add('brand_typography')
                
                # Check shape colors
                if hasattr(shape, 'fill'):
                    try:
                        if hasattr(shape.fill, 'fore_color') and hasattr(shape.fill.fore_color, 'rgb'):
                            rgb = shape.fill.fore_color.rgb
                            hex_color = f'#{rgb.red:02x}{rgb.green:02x}{rgb.blue:02x}'
                            if hex_color.lower() in [c.lower() for c in brand_colors.keys()]:
                                colors_found.add(hex_color.lower())
                    except:
                        pass
        
        # Check for brand color usage
        brand_colors_used = len([c for c in brand_colors.keys() if c.lower() in colors_found])
        if brand_colors_used >= 4:
            checks.append({"name": "brand_colors", "passed": True, "detail": f"Found {brand_colors_used} brand colors in use"})
            score += 20.0
        else:
            checks.append({"name": "brand_colors", "passed": False, "detail": f"Only found {brand_colors_used} brand colors, expected at least 4"})
        
        # Check for key content requirements
        content_checks = {
            'color': ('color palette' in text_content or 'color' in text_content),
            'typography': ('typography' in text_content or 'font' in text_content or 'poppins' in text_content or 'lora' in text_content),
            'brand': ('brand' in text_content or 'anthropic' in text_content),
            'guidelines': ('guideline' in text_content or 'usage' in text_content or 'application' in text_content),
            'hex_codes': any(c.replace('#', '') in text_content for c in brand_colors.keys())
        }
        
        content_score = sum(content_checks.values())
        if content_score >= 3:
            checks.append({"name": "content_requirements", "passed": True, "detail": f"Found {content_score}/5 required content elements"})
            score += 25.0
        else:
            checks.append({"name": "content_requirements", "passed": False, "detail": f"Only found {content_score}/5 required content elements"})
        
        # Check for comprehensive coverage (title + color + typography + application + examples)
        slide_purposes = []
        for slide in prs.slides:
            slide_text = ''
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    slide_text += shape.text.lower() + ' '
            
            if any(word in slide_text for word in ['title', 'anthropic', 'brand']):
                slide_purposes.append('title')
            elif any(word in slide_text for word in ['color', 'palette']):
                slide_purposes.append('colors')
            elif any(word in slide_text for word in ['typography', 'font', 'text']):
                slide_purposes.append('typography')
            elif any(word in slide_text for word in ['application', 'chart', 'graphic']):
                slide_purposes.append('application')
            elif any(word in slide_text for word in ['example', 'usage', 'guideline']):
                slide_purposes.append('examples')
        
        unique_purposes = len(set(slide_purposes))
        if unique_purposes >= 4:
            checks.append({"name": "slide_coverage", "passed": True, "detail": f"Covers {unique_purposes} different brand aspects"})
            score += 20.0
        else:
            checks.append({"name": "slide_coverage", "passed": False, "detail": f"Only covers {unique_purposes} brand aspects, need at least 4"})
        
        # Professional quality check - multiple shapes/elements per slide on average
        total_shapes = sum(len(slide.shapes) for slide in prs.slides)
        avg_shapes = total_shapes / len(prs.slides) if prs.slides else 0
        
        if avg_shapes >= 3:
            checks.append({"name": "professional_quality", "passed": True, "detail": f"Average {avg_shapes:.1f} elements per slide indicates good design complexity"})
            score += 10.0
        else:
            checks.append({"name": "professional_quality", "passed": False, "detail": f"Only {avg_shapes:.1f} elements per slide, needs more design elements"})
        
    except Exception as e:
        checks.append({"name": "presentation_analysis", "passed": False, "detail": f"Error analyzing presentation: {str(e)}"})
        return {"passed": False, "score": score, "checks": checks}
    
    passed = score >= 70.0  # Require 70% for hard difficulty
    return {"passed": passed, "score": score, "checks": checks}

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print('{"passed": false, "score": 0.0, "checks": [{"name": "usage", "passed": false, "detail": "Usage: eval_script.py <workspace_path>"}]}')
        sys.exit(1)
    
    workspace_path = sys.argv[1]
    result = evaluate_presentation(workspace_path)
    print(json.dumps(result))