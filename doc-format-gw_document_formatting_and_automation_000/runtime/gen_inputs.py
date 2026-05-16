#!/usr/bin/env python3
"""
Generate the sandbox workspace with a messy, unformatted government document
and distractor files.
"""

import os
import random
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

random.seed(42)

workspace = Path(os.environ.get("WORKSPACE", "/workspace"))
workspace.mkdir(parents=True, exist_ok=True)

# ─── Distractor directory structure ────────────────────────────────────────────
dirs = [
    "archive/2023/q1",
    "archive/2023/q2",
    "drafts/internal",
    "drafts/external",
    "templates/official",
    "templates/internal",
    "reports/monthly",
    "reports/annual",
    "correspondence/inbox",
    "correspondence/outbox",
    "misc",
]
for d in dirs:
    (workspace / d).mkdir(parents=True, exist_ok=True)

# Distractor text files
distractors = {
    "archive/2023/q1/summary.txt": "第一季度工作总结归档文件。",
    "archive/2023/q2/notes.txt": "第二季度会议记录。\n参会人员：张三、李四、王五",
    "drafts/internal/memo_draft.txt": "内部备忘录草稿 - 关于人员调整的通知",
    "drafts/external/letter_template.txt": "外部信函模板，请按照规范填写。",
    "templates/official/header_sample.txt": "文件头部样本 - 单位名称\n文件编号：[年份]第XX号",
    "templates/internal/report_format.txt": "内部报告格式说明\n1. 标题居中\n2. 正文仿宋四号",
    "reports/monthly/oct_report.txt": "十月份工作报告\n完成情况：达标\n备注：无",
    "reports/annual/2023_annual.txt": "2023年度工作报告（初稿）\n总体情况良好。",
    "correspondence/inbox/receipt_001.txt": "收文登记：来文单位 XX局，文号 XX〔2024〕1号",
    "correspondence/outbox/sent_log.txt": "发文记录：\n日期 | 文号 | 主题\n2024-01-10 | 001 | 关于XX的通知",
    "misc/contacts.txt": "联系人列表：\n张三 - 办公室主任\n李四 - 文秘科长",
    "misc/todo.txt": "待办事项：\n1. 完成公文排版\n2. 提交材料\n3. 归档",
}

for rel_path, content in distractors.items():
    fpath = workspace / rel_path
    fpath.parent.mkdir(parents=True, exist_ok=True)
    fpath.write_text(content, encoding="utf-8")

# ─── CREATE THE MESSY INPUT DOCUMENT ──────────────────────────────────────────
# This document has:
#   - Wrong fonts (Arial/Times New Roman, wrong sizes)
#   - No consistent heading detection
#   - Wrong alignment
#   - A table with wrong formatting
#   - Multiple heading levels that need proper detection

doc = Document()

# Remove default styles messing up margins — use a fresh section
section = doc.sections[0]
# Set WRONG margins (agent must fix these)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(3.0)
section.right_margin = Cm(3.0)

def add_messy_para(doc, text, font_name="Arial", font_size=12, bold=False,
                   align=WD_ALIGN_PARAGRAPH.LEFT):
    para = doc.add_paragraph()
    para.alignment = align
    run = para.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    return para

# Main title — wrong font, wrong size, wrong alignment
add_messy_para(
    doc,
    "关于加强基层政务公开工作的实施意见",
    font_name="Times New Roman",
    font_size=16,
    bold=True,
    align=WD_ALIGN_PARAGRAPH.LEFT,  # should be CENTER
)

# Body intro — wrong font
add_messy_para(
    doc,
    "为深入贯彻落实党中央、国务院关于推进政务公开的决策部署，进一步提升基层政务公开水平，结合本地区实际，现就加强基层政务公开工作提出如下意见。",
    font_name="Arial",
    font_size=12,
)

# Heading 1 — should be 黑体 14pt
add_messy_para(
    doc,
    "一、总体要求",
    font_name="Calibri",
    font_size=13,
    bold=True,
)

add_messy_para(
    doc,
    "坚持以人民为中心的发展思想，以公开为常态、不公开为例外，全面推进基层政务公开标准化、规范化建设。",
    font_name="Arial",
    font_size=11,
)

# Heading 2 — should be 楷体_GB2312 14pt
add_messy_para(
    doc,
    "（一）指导思想",
    font_name="Times New Roman",
    font_size=12,
    bold=True,
)

add_messy_para(
    doc,
    "深入贯彻习近平新时代中国特色社会主义思想，牢固树立新发展理念，以保障公民知情权、参与权、表达权、监督权为出发点。",
    font_name="Verdana",
    font_size=11,
)

# Heading 3 — should be 仿宋_GB2312 14pt BOLD
add_messy_para(
    doc,
    "1.工作原则",
    font_name="Arial",
    font_size=12,
    bold=False,  # should be bold
)

add_messy_para(
    doc,
    "坚持依法公开、真实公开、及时公开、便民公开的原则，确保公开内容合法合规。",
    font_name="Arial",
    font_size=11,
)

# Heading 4 — should be 仿宋_GB2312 14pt NOT bold
add_messy_para(
    doc,
    "（1）依法公开",
    font_name="Calibri",
    font_size=11,
    bold=True,  # should NOT be bold
)

add_messy_para(
    doc,
    "依据法律法规规定，明确公开范围，规范公开程序，防止违规公开和应公开未公开。",
    font_name="Arial",
    font_size=11,
)

# Heading 1 again
add_messy_para(
    doc,
    "二、主要任务",
    font_name="Georgia",
    font_size=13,
    bold=False,
)

# Heading 2
add_messy_para(
    doc,
    "（二）重点领域公开",
    font_name="Arial",
    font_size=12,
)

add_messy_para(
    doc,
    "聚焦群众最关心的民生事项，重点推进教育、医疗、社保、住房等领域政务公开工作。",
    font_name="Calibri",
    font_size=11,
)

# Heading 3
add_messy_para(
    doc,
    "2.信息公开目录",
    font_name="Times New Roman",
    font_size=12,
)

add_messy_para(
    doc,
    "各部门应建立健全政务公开信息目录，明确公开事项、公开内容、公开时限和公开方式。",
    font_name="Verdana",
    font_size=10,
)

# Add a TABLE (wrong formatting)
table = doc.add_table(rows=4, cols=3)
table.style = 'Table Grid'

headers = ["序号", "公开事项", "责任部门"]
for j, header in enumerate(headers):
    cell = table.cell(0, j)
    para = cell.paragraphs[0]
    run = para.add_run(header)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True

rows_data = [
    ("1", "行政许可事项", "行政审批局"),
    ("2", "财政预算信息", "财政局"),
    ("3", "扶贫政策落实情况", "扶贫开发办公室"),
]
for i, (num, item, dept) in enumerate(rows_data, 1):
    for j, text in enumerate([num, item, dept]):
        cell = table.cell(i, j)
        para = cell.paragraphs[0]
        run = para.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)  # wrong size, should be 10.5pt

# More body content
add_messy_para(
    doc,
    "三、保障措施",
    font_name="Calibri",
    font_size=14,
    bold=False,
)

add_messy_para(
    doc,
    "（三）强化监督考核",
    font_name="Georgia",
    font_size=11,
)

add_messy_para(
    doc,
    "将政务公开工作纳入年度绩效考核体系，建立常态化督查机制，对工作落实不到位的单位予以通报批评。",
    font_name="Arial",
    font_size=10,
)

add_messy_para(
    doc,
    "3.考核指标",
    font_name="Arial",
    font_size=11,
    bold=False,  # should be bold
)

add_messy_para(
    doc,
    "各地区应制定量化考核指标，重点考核公开及时率、公开完整率、群众满意度等核心指标。",
    font_name="Verdana",
    font_size=11,
)

add_messy_para(
    doc,
    "（2）考核结果运用",
    font_name="Times New Roman",
    font_size=12,
    bold=True,  # should NOT be bold
)

add_messy_para(
    doc,
    "考核结果作为领导干部年度考核、职务晋升的重要参考依据。",
    font_name="Arial",
    font_size=11,
)

# Save the messy document
input_doc_path = workspace / "drafts" / "external" / "政务公开实施意见_草稿.docx"
doc.save(str(input_doc_path))

print(f"[+] Created messy input document: {input_doc_path}")
print(f"[+] Created {len(distractors)} distractor files")
print("[+] Workspace setup complete.")