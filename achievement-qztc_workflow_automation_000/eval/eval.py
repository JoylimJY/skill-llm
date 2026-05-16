#!/usr/bin/env python3
"""
Evaluation script for the achievement-qztc task.
Usage: python3 eval_script.py /workspace
"""
import sys
import json
import re
from pathlib import Path

def find_output_file(workspace: Path):
    """Find the generated output docx file."""
    candidates = list(workspace.rglob("课程目标达成情况分析表-数据可视化-23级软工.docx"))
    # Exclude template
    candidates = [c for c in candidates if "模版" not in c.name]
    return candidates[0] if candidates else None


def run_checks(workspace_str: str):
    checks = []
    workspace = Path(workspace_str)

    # ── Check 0: Output file exists ──
    output_file = find_output_file(workspace)
    checks.append({
        "name": "output_file_exists",
        "passed": output_file is not None,
        "detail": f"Found: {output_file}" if output_file else "Output file not found. Expected: 课程目标达成情况分析表-数据可视化-23级软工.docx"
    })
    if output_file is None:
        return checks

    try:
        from docx import Document
        doc = Document(str(output_file))
    except Exception as e:
        checks.append({"name": "can_open_docx", "passed": False, "detail": str(e)})
        return checks

    checks.append({"name": "can_open_docx", "passed": True, "detail": "Document opened successfully."})

    # ── Check 1: Placeholder replacement in document paragraphs ──
    all_para_text = " ".join(p.text for p in doc.paragraphs)
    all_table_text = []
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                all_table_text.append(cell.text)
    all_text = all_para_text + " ".join(all_table_text)

    # No raw placeholders should remain
    remaining_placeholders = re.findall(r'\$\w+\$', all_text)
    checks.append({
        "name": "no_raw_placeholders",
        "passed": len(remaining_placeholders) == 0,
        "detail": f"Remaining placeholders: {remaining_placeholders}" if remaining_placeholders else "All placeholders replaced."
    })

    # grade and major replaced
    grade_present = "23" in all_text
    major_present = "软工" in all_text
    checks.append({
        "name": "grade_and_major_replaced",
        "passed": grade_present and major_present,
        "detail": f"grade('23') found: {grade_present}, major('软工') found: {major_present}"
    })

    # total students = 7 (9 - 2 旷考)
    total_present = "7人" in all_text
    checks.append({
        "name": "total_students_7ren",
        "passed": total_present,
        "detail": f"'7人' found in document: {total_present}. Excerpt: {[t for t in all_table_text if '人' in t][:5]}"
    })

    # ── Check 2: Table 8 structure ──
    try:
        table8 = doc.tables[8]
        t8_rows = len(table8.rows)
        # Expected: 2 header + 7 data + 1 avg = 10 rows
        expected_t8_rows = 10
        checks.append({
            "name": "table8_row_count",
            "passed": t8_rows == expected_t8_rows,
            "detail": f"Table8 rows: {t8_rows}, expected: {expected_t8_rows}"
        })
    except Exception as e:
        checks.append({"name": "table8_row_count", "passed": False, "detail": str(e)})
        table8 = None

    # ── Check 3: Table 8 data rows - correct student IDs ──
    if table8 is not None:
        try:
            expected_students = [
                ('1', '2301001', '陈晓明'),
                ('2', '2301002', '李雨涵'),
                # 2301003 旷考 - skipped
                ('3', '2301004', '赵静雯'),
                # 2301005 旷考 - skipped
                ('4', '2301006', '孙浩然'),
                ('5', '2301007', '周思雨'),
                ('6', '2301008', '吴明杰'),
                ('7', '2301009', '郑晨曦'),
            ]
            all_match = True
            mismatches = []
            for i, (exp_seq, exp_sid, exp_name) in enumerate(expected_students):
                row_idx = i + 2  # after 2 header rows
                if row_idx >= len(table8.rows) - 1:
                    all_match = False
                    mismatches.append(f"Row {row_idx} missing")
                    continue
                row = table8.rows[row_idx]
                actual_seq = row.cells[0].text.strip()
                actual_sid = row.cells[1].text.strip()
                actual_name = row.cells[2].text.strip()
                if actual_seq != exp_seq or actual_sid != exp_sid or actual_name != exp_name:
                    all_match = False
                    mismatches.append(f"Row {row_idx}: got ({actual_seq},{actual_sid},{actual_name}), expected ({exp_seq},{exp_sid},{exp_name})")

            checks.append({
                "name": "table8_student_data_correct",
                "passed": all_match,
                "detail": "All 7 students correct (旷考 excluded)" if all_match else f"Mismatches: {mismatches}"
            })
        except Exception as e:
            checks.append({"name": "table8_student_data_correct", "passed": False, "detail": str(e)})

    # ── Check 4: Percentage format in Table 8 ──
    if table8 is not None:
        try:
            pct_pattern = re.compile(r'^\d+%$')
            score_pattern = re.compile(r'^\d+\.\d{2}$')
            pct_cols = [4, 6, 8, 10]
            score_cols = [3, 5, 7, 9]

            pct_ok = True
            score_ok = True
            pct_errors = []
            score_errors = []

            for row_idx in range(2, len(table8.rows) - 1):
                row = table8.rows[row_idx]
                for col in pct_cols:
                    val = row.cells[col].text.strip()
                    if val and not pct_pattern.match(val):
                        pct_ok = False
                        pct_errors.append(f"row{row_idx}col{col}='{val}'")
                for col in score_cols:
                    val = row.cells[col].text.strip()
                    if val and not score_pattern.match(val):
                        score_ok = False
                        score_errors.append(f"row{row_idx}col{col}='{val}'")

            checks.append({
                "name": "table8_percentage_format",
                "passed": pct_ok,
                "detail": "All achievement values in XX% format" if pct_ok else f"Bad pct values: {pct_errors[:5]}"
            })
            checks.append({
                "name": "table8_score_format",
                "passed": score_ok,
                "detail": "All scores in X.XX format" if score_ok else f"Bad score values: {score_errors[:5]}"
            })
        except Exception as e:
            checks.append({"name": "table8_percentage_format", "passed": False, "detail": str(e)})
            checks.append({"name": "table8_score_format", "passed": False, "detail": str(e)})

    # ── Check 5: Score ranges correct ──
    if table8 is not None:
        try:
            range_ok = True
            range_errors = []
            denominators = {3: 25, 5: 30, 7: 26.5, 9: 18.5}
            ranges = {3: (10, 24), 5: (20, 28), 7: (20, 25), 9: (15, 18)}

            for row_idx in range(2, len(table8.rows) - 1):
                row = table8.rows[row_idx]
                for score_col, (lo, hi) in ranges.items():
                    val_str = row.cells[score_col].text.strip()
                    pct_col = score_col + 1
                    pct_str = row.cells[pct_col].text.strip()
                    if not val_str:
                        continue
                    try:
                        val = float(val_str)
                        # Check score is integer within range (stored as X.00)
                        val_int = int(round(val))
                        if not (lo <= val_int <= hi):
                            range_ok = False
                            range_errors.append(f"row{row_idx} col{score_col}: {val_int} not in [{lo},{hi}]")
                        # Check percentage consistency
                        denom = denominators[score_col]
                        expected_pct = int(round(val / denom * 100))
                        if pct_str:
                            actual_pct = int(pct_str.replace('%', ''))
                            if abs(actual_pct - expected_pct) > 1:
                                range_ok = False
                                range_errors.append(f"row{row_idx} col{pct_col}: pct={actual_pct}% but expected≈{expected_pct}%")
                    except ValueError:
                        pass

            checks.append({
                "name": "table8_score_ranges_and_pct_consistency",
                "passed": range_ok,
                "detail": "Scores within valid ranges and percentages consistent" if range_ok else f"Errors: {range_errors[:5]}"
            })
        except Exception as e:
            checks.append({"name": "table8_score_ranges_and_pct_consistency", "passed": False, "detail": str(e)})

    # ── Check 6: Average row in Table 8 ──
    if table8 is not None:
        try:
            avg_row = table8.rows[-1]
            avg_label = avg_row.cells[0].text.strip()
            avg_pct_cols_ok = True
            avg_errors = []

            label_ok = avg_label == "平均值"
            pct_pattern_avg = re.compile(r'^\d+%$')
            for col in [4, 6, 8, 10]:
                val = avg_row.cells[col].text.strip()
                if not pct_pattern_avg.match(val):
                    avg_pct_cols_ok = False
                    avg_errors.append(f"avg col{col}='{val}'")

            checks.append({
                "name": "table8_avg_row",
                "passed": label_ok and avg_pct_cols_ok,
                "detail": f"Label='{avg_label}', pct_ok={avg_pct_cols_ok}" + (f", errors: {avg_errors}" if avg_errors else "")
            })
        except Exception as e:
            checks.append({"name": "table8_avg_row", "passed": False, "detail": str(e)})

    # ── Check 7: Table 7 updated with averages ──
    try:
        table7 = doc.tables[7]
        t7_cells_to_check = [
            (5, 7),
            (8, 7),
            (11, 7),
            (14, 7),
        ]
        pct_pattern_t7 = re.compile(r'^\d+%$')
        t7_ok = True
        t7_errors = []
        for r, c in t7_cells_to_check:
            val = table7.rows[r].cells[c].text.strip()
            if not pct_pattern_t7.match(val):
                t7_ok = False
                t7_errors.append(f"table7[{r}][{c}]='{val}'")

        checks.append({
            "name": "table7_averages_updated",
            "passed": t7_ok,
            "detail": "All 4 average cells in table7 have percentage values" if t7_ok else f"Issues: {t7_errors}"
        })
    except Exception as e:
        checks.append({"name": "table7_averages_updated", "passed": False, "detail": str(e)})

    # ── Check 8: Table 7 and Table 8 averages are consistent ──
    if table8 is not None:
        try:
            # Get averages from table8 last row
            avg_row = table8.rows[-1]
            t8_avg_v1 = avg_row.cells[4].text.strip().replace('%', '')
            t8_avg_v2 = avg_row.cells[6].text.strip().replace('%', '')
            t8_avg_v3 = avg_row.cells[8].text.strip().replace('%', '')
            t8_avg_v4 = avg_row.cells[10].text.strip().replace('%', '')

            table7 = doc.tables[7]
            t7_avg_v1 = table7.rows[5].cells[7].text.strip().replace('%', '')
            t7_avg_v2 = table7.rows[8].cells[7].text.strip().replace('%', '')
            t7_avg_v3 = table7.rows[11].cells[7].text.strip().replace('%', '')
            t7_avg_v4 = table7.rows[14].cells[7].text.strip().replace('%', '')

            consistent = (
                t8_avg_v1 == t7_avg_v1 and
                t8_avg_v2 == t7_avg_v2 and
                t8_avg_v3 == t7_avg_v3 and
                t8_avg_v4 == t7_avg_v4
            )
            checks.append({
                "name": "table7_table8_averages_consistent",
                "passed": consistent,
                "detail": f"t8=[{t8_avg_v1}%,{t8_avg_v2}%,{t8_avg_v3}%,{t8_avg_v4}%] t7=[{t7_avg_v1}%,{t7_avg_v2}%,{t7_avg_v3}%,{t7_avg_v4}%]"
            })
        except Exception as e:
            checks.append({"name": "table7_table8_averages_consistent", "passed": False, "detail": str(e)})

    return checks


def main():
    workspace_str = sys.argv[1] if len(sys.argv) > 1 else "/workspace"
    checks = run_checks(workspace_str)

    passed_count = sum(1 for c in checks if c["passed"])
    total = len(checks)
    score = passed_count / total if total > 0 else 0.0
    all_passed = all(c["passed"] for c in checks)

    result = {
        "passed": all_passed,
        "score": round(score, 4),
        "checks": checks
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()