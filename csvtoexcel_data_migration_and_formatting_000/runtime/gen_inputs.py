import os
import csv
import io
import random

random.seed(42)

WORKSPACE = "/workspace"

# ── directory structure ──────────────────────────────────────────────────────
dirs = [
    "data/raw/2024Q4",
    "data/raw/archive",
    "data/processed",
    "reports/monthly",
    "reports/annual",
    "scripts",                        # skill scripts live here
    "config",
    "logs",
    "tmp",
    "exports/pending",
    "exports/done",
]
for d in dirs:
    os.makedirs(os.path.join(WORKSPACE, d), exist_ok=True)

# ── distractor files ─────────────────────────────────────────────────────────
distractors = {
    "config/db_config.yaml": "host: localhost\nport: 5432\nname: retail_db\n",
    "config/report_settings.json": '{"theme":"corporate","locale":"zh_CN"}\n',
    "logs/etl_2024-12-01.log": "INFO: pipeline started\nINFO: 3 sources loaded\nWARN: null values in col 3\n",
    "logs/etl_2024-12-02.log": "INFO: pipeline started\nINFO: 2 sources loaded\n",
    "reports/monthly/nov_summary.txt": "November revenue: 4,320,000 CNY\nUnits sold: 18,200\n",
    "reports/annual/2023_annual.txt": "Total revenue 2023: 52M CNY\n",
    "data/raw/archive/q3_sales_backup.csv": "date,amount\n2024-07-01,1000\n2024-07-02,2000\n",
    "data/processed/cleaned_oct.csv": "sku,qty,price\nA001,10,99.9\nA002,5,199.0\n",
    "tmp/scratch_notes.txt": "TODO: verify GBK files from Shenzhen warehouse\n",
    "exports/pending/placeholder.txt": "awaiting conversion\n",
    "exports/done/.gitkeep": "",
}
for rel, content in distractors.items():
    with open(os.path.join(WORKSPACE, rel), "w", encoding="utf-8") as f:
        f.write(content)

# ── THE SKILL SCRIPT (already exists per spec — we only reference it) ────────
# The SKILL.md states "All scripts mentioned in the SKILL.md already exist in the workspace."
# We place the actual csv_to_excel.py script so the agent can find and run it.
csv_to_excel_script = r'''#!/usr/bin/env python3
"""CSV to Excel converter with professional formatting and Chinese support."""

import argparse
import csv
import os
import sys

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Error: openpyxl is required. Run: pip install openpyxl")
    sys.exit(1)


def detect_encoding(filepath):
    encodings = ["utf-8", "gbk", "gb2312", "utf-8-sig", "latin1"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                f.read()
            return enc
        except (UnicodeDecodeError, LookupError):
            continue
    return "latin1"


def detect_dialect(filepath, encoding):
    with open(filepath, "r", encoding=encoding, newline="") as f:
        sample = f.read(4096)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return dialect


def col_width(text):
    """Chinese chars = 2 units, ASCII = 1 unit."""
    w = 0
    for ch in str(text):
        w += 2 if ord(ch) > 127 else 1
    return w


def apply_formatting(ws, max_col):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_align = Alignment(horizontal="center")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border

    # auto column widths
    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        max_w = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_w = max(max_w, col_width(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_w + 2, 50)

    ws.freeze_panes = "A2"


def csv_to_sheet(ws, filepath):
    encoding = detect_encoding(filepath)
    dialect = detect_dialect(filepath, encoding)
    with open(filepath, "r", encoding=encoding, newline="") as f:
        reader = csv.reader(f, dialect=dialect)
        rows = list(reader)
    max_col = 0
    for row in rows:
        ws.append(row)
        max_col = max(max_col, len(row))
    return max_col


def main():
    parser = argparse.ArgumentParser(description="Convert CSV(s) to Excel")
    parser.add_argument("inputs", nargs="+", help="Input CSV file(s)")
    parser.add_argument("output", nargs="?", help="Output Excel file (single-file mode)")
    parser.add_argument("--output", dest="output_flag", help="Output Excel file (multi-file mode)")
    parser.add_argument("--sheet-names", nargs="+", dest="sheet_names", help="Custom sheet names")
    args = parser.parse_args()

    output_path = args.output_flag or args.output
    inputs = args.inputs

    # single-file positional: last arg is output if no --output
    if not output_path:
        if len(inputs) >= 2:
            output_path = inputs[-1]
            inputs = inputs[:-1]
        else:
            parser.error("Specify output file as positional arg or --output")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    sheet_names = args.sheet_names or []

    for i, csv_path in enumerate(inputs):
        if i < len(sheet_names):
            sname = sheet_names[i][:31]
        else:
            sname = os.path.splitext(os.path.basename(csv_path))[0][:31]
        ws = wb.create_sheet(title=sname)
        max_col = csv_to_sheet(ws, csv_path)
        apply_formatting(ws, max_col)

    wb.save(output_path)
    print(f"Saved: {output_path}")


if __name__ == "__main__":
    main()
'''

script_path = os.path.join(WORKSPACE, "scripts", "csv_to_excel.py")
with open(script_path, "w", encoding="utf-8") as f:
    f.write(csv_to_excel_script)
os.chmod(script_path, 0o755)

# ── INPUT CSVs ───────────────────────────────────────────────────────────────

# 1. 销售部门 (Sales) — UTF-8
sales_rows = [
    ["日期", "产品名称", "销售额(元)", "数量", "销售员"],
    ["2024-10-01", "蓝牙耳机", "12800", "32", "张伟"],
    ["2024-10-02", "智能手表", "45600", "19", "李娜"],
    ["2024-10-03", "无线充电器", "6750", "45", "王芳"],
    ["2024-10-07", "蓝牙耳机", "9600", "24", "张伟"],
    ["2024-10-08", "智能手表", "52000", "22", "陈明"],
    ["2024-10-09", "平板电脑支架", "3200", "40", "李娜"],
    ["2024-10-14", "USB集线器", "8100", "54", "王芳"],
    ["2024-10-15", "机械键盘", "28500", "15", "陈明"],
]

sales_path = os.path.join(WORKSPACE, "data/raw/2024Q4", "sales_oct.csv")
with open(sales_path, "w", encoding="utf-8", newline="") as f:
    writer = csv.writer(f)
    writer.writerows(sales_rows)

# 2. 库存部门 (Inventory) — GBK encoded (the adversarial trap)
inventory_rows = [
    ["SKU编号", "商品名称", "库存数量", "仓库位置", "最后盘点日期"],
    ["SKU-001", "蓝牙耳机", "156", "深圳仓A区", "2024-10-31"],
    ["SKU-002", "智能手表", "43", "深圳仓B区", "2024-10-31"],
    ["SKU-003", "无线充电器", "289", "广州仓C区", "2024-10-30"],
    ["SKU-004", "平板电脑支架", "112", "上海仓A区", "2024-10-28"],
    ["SKU-005", "USB集线器", "78", "北京仓D区", "2024-10-29"],
    ["SKU-006", "机械键盘", "55", "深圳仓A区", "2024-10-31"],
    ["SKU-007", "鼠标垫", "430", "广州仓C区", "2024-10-27"],
]

inventory_path = os.path.join(WORKSPACE, "data/raw/2024Q4", "inventory_oct.csv")
with open(inventory_path, "w", encoding="gbk", newline="") as f:
    writer = csv.writer(f)
    writer.writerows(inventory_rows)

# 3. 财务部门 (Finance) — UTF-8-SIG (BOM)
finance_rows = [
    ["费用类别", "金额(元)", "申请部门", "审批状态", "备注"],
    ["差旅费", "15200", "销售部", "已审批", "10月出差上海"],
    ["办公耗材", "3800", "行政部", "已审批", "打印纸墨盒等"],
    ["设备维修", "6500", "技术部", "待审批", "服务器硬盘更换"],
    ["市场推广", "88000", "市场部", "已审批", "双十一促销活动"],
    ["培训费用", "12000", "人力资源部", "已审批", "季度技能培训"],
    ["快递物流", "4200", "仓储部", "已审批", "10月快递费汇总"],
]

finance_path = os.path.join(WORKSPACE, "data/raw/2024Q4", "finance_oct.csv")
with open(finance_path, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f)
    writer.writerows(finance_rows)

print("Workspace initialized.")
print(f"  Sales CSV (UTF-8):    {sales_path}")
print(f"  Inventory CSV (GBK):  {inventory_path}")
print(f"  Finance CSV (UTF-8-SIG): {finance_path}")