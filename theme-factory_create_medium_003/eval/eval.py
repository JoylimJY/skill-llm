import sys
import os
import json
from pptx import Presentation
import re

def find_file(dir_path, keyword, extensions):
    candidates = []
    for f in os.listdir(dir_path):
        if any(f.lower().endswith(ext) for ext in extensions):
            if keyword.lower() in f.lower():
                candidates.append(os.path.join(dir_path, f))
    return candidates

def parse_theme_file(path):
    data = {}
    try:
        with open(path, 'r') as f:
            lines = f.readlines()
        colors = {}
        fonts = {}
        current_section = None
        for line in lines:
            line = line.strip()
            if re.match(r'^colors\s*:\s*$', line, re.I):
                current_section = 'colors'
            elif re.match(r'^fonts\s*:\s*$', line, re.I):
                current_section = 'fonts'
            elif ':' in line and current_section:
                key_val = line.split(':', 1)
                key = key_val[0].strip().lower()
                val = key_val[1].strip().strip('"')
                if current_section == 'colors':
                    colors[key] = val
                elif current_section == 'fonts':
                    fonts[key] = val
        data['colors'] = colors
        data['fonts'] = fonts
    except Exception as e:
        data = {}
    return data

def color_ok(c):
    # Simple hex color validator
    return bool(re.match(r'^#([0-9a-fA-F]{6})$', c))

def lower_included(substrs, text):
    text = text.lower()
    return any(sub.lower() in text for sub in substrs)

def main():
    workspace = sys.argv[1]
    checks = []

    # 1. Check theme-showcase.pdf exists and include all 10 themes
    pdf_paths = find_file(workspace, 'theme-showcase', ['.pdf'])
    theme_names = [
        'Ocean Depths', 'Sunset Boulevard', 'Forest Canopy', 'Modern Minimalist', 'Golden Hour',
        'Arctic Frost', 'Desert Rose', 'Tech Innovation', 'Botanical Garden', 'Midnight Galaxy'
    ]
    found_theme_showcase = False
    if pdf_paths:
        for pdf_path in pdf_paths:
            try:
                with open(pdf_path, 'r', encoding='latin1') as f:
                    content = f.read().lower()
                if all(tn.lower() in content for tn in theme_names):
                    found_theme_showcase = True
                    break
            except:
                continue
    checks.append({
        'name': 'Showcase PDF with all 10 themes',
        'passed': found_theme_showcase,
        'detail': 'Found theme-showcase.pdf and all 10 theme names inside' if found_theme_showcase else 'Missing theme-showcase.pdf or not all theme names found inside'
    })

    # 2. Check the 'Ocean Depths' theme file exists in themes directory and parse correctly
    theme_files = find_file(os.path.join(workspace, 'themes'), 'ocean depths', ['.yaml', '.yml', '.txt'])
    ocean_theme_data = None
    found_ocean_theme = False
    for tf in theme_files:
        data = parse_theme_file(tf)
        if data and 'colors' in data and 'fonts' in data:
            # Check required keys
            colors = data.get('colors', {})
            fonts = data.get('fonts', {})
            color_keys = ['primary', 'secondary', 'accent1', 'accent2']
            font_keys = ['header', 'body']
            if all(k in colors and color_ok(colors[k]) for k in color_keys) and all(k in fonts and fonts[k] for k in font_keys):
                ocean_theme_data = data
                found_ocean_theme = True
                break
    checks.append({
        'name': 'Ocean Depths theme specification file',
        'passed': found_ocean_theme,
        'detail': 'Valid Ocean Depths theme file with colors and fonts found' if found_ocean_theme else 'Ocean Depths theme file missing or malformed'
    })

    # 3. Check that output file 'project_presentation_themed.pptx' exists
    output_files = find_file(workspace, 'project_presentation_themed', ['.pptx'])
    correct_pptx_path = output_files[0] if output_files else None
    output_exists = correct_pptx_path is not None
    checks.append({
        'name': 'Output themed PPTX exists',
        'passed': output_exists,
        'detail': 'The styled presentation file project_presentation_themed.pptx exists' if output_exists else 'Missing project_presentation_themed.pptx output file'
    })

    # If no output file, cannot do further checks
    if not output_exists:
        print(json.dumps({
            'passed': False,
            'score': 0.0,
            'checks': checks
        }))
        return

    # 4. Load output PPTX and verify theme colors/fonts applied somewhat
    try:
        prs = Presentation(correct_pptx_path)
        # We expect 5 slides
        slides_correct = (len(prs.slides) == 5)
        # Check text fonts include those specified
        header_font_substring = ocean_theme_data.get('fonts', {}).get('header', '').lower() if ocean_theme_data else ''
        body_font_substring = ocean_theme_data.get('fonts', {}).get('body', '').lower() if ocean_theme_data else ''

        font_header_found = False
        font_body_found = False

        # Scan all shape text font names for header and body fonts
        all_fonts = set()
        for slide in prs.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        font = run.font
                        name = font.name
                        if name:
                            all_fonts.add(name.lower())

        font_header_found = any(header_font_substring in f for f in all_fonts if header_font_substring) or not header_font_substring
        font_body_found = any(body_font_substring in f for f in all_fonts if body_font_substring) or not body_font_substring
        text_fonts_check = font_header_found and font_body_found

        # Check hex colors likely applied by examining slide background or shape fills for primary and secondary colors
        colors = ocean_theme_data.get('colors', {}) if ocean_theme_data else {}
        primary_color = colors.get('primary', '').lower()
        secondary_color = colors.get('secondary', '').lower()

        # PPTX colors are usually RGB triples - try to approximate by parsing colors from theme or text.
        # Here we do a heuristic: check if primary/secondary hex strings appear in slide XML content
        xml_combined = "".join([slide.element.xml.decode('utf-8', errors='ignore').lower() for slide in prs.slides])
        primary_in_xml = primary_color and primary_color[1:] in xml_combined
        secondary_in_xml = secondary_color and secondary_color[1:] in xml_combined

        colors_applied = primary_in_xml or secondary_in_xml

        checks.extend([
            {
                'name': 'Output presentation has 5 slides',
                'passed': slides_correct,
                'detail': f'Found {len(prs.slides)} slides; expected 5'
            },
            {
                'name': 'Fonts matching Ocean Depths theme applied',
                'passed': text_fonts_check,
                'detail': 'Header and body fonts appear to match theme fonts' if text_fonts_check else 'Fonts do not match expected theme fonts'
            },
            {
                'name': 'Colors from Ocean Depths theme applied',
                'passed': colors_applied,
                'detail': 'Primary or secondary theme colors found in presentation xml data' if colors_applied else 'Theme colors not detected in presentation'
            }
        ])

    except Exception as e:
        checks.append({
            'name': 'Load and verify output PPTX',
            'passed': False,
            'detail': f'Error opening or processing output PPTX: {str(e)}'
        })

    # Compute overall score
    passed_checks = sum(1 for c in checks if c.get('passed'))
    total_checks = len(checks)
    score = passed_checks / total_checks if total_checks > 0 else 0.0
    passed = (score == 1.0)

    print(json.dumps({
        'passed': passed,
        'score': score,
        'checks': checks
    }))

if __name__ == '__main__':
    main()
