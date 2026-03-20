import sys
import os
import json
from docx import Document

def main():
    workspace = sys.argv[1]
    output_path = os.path.join(workspace, "ProjectSummary.docx")

    checks = []
    passed = True

    # Check file existence
    file_exists = os.path.isfile(output_path)
    checks.append({"name": "file_exists", "passed": file_exists, "detail": f'Output file found: {file_exists}'})
    if not file_exists:
        print(json.dumps({"passed": False, "score": 0, "checks": checks}))
        return

    # Load document
    try:
        doc = Document(output_path)
    except Exception as e:
        checks.append({"name": "open_doc", "passed": False, "detail": f'Failed to open DOCX: {e}'})
        print(json.dumps({"passed": False, "score": 0, "checks": checks}))
        return

    # Check first paragraph is heading 1 with correct text
    if len(doc.paragraphs) == 0:
        checks.append({"name": "heading_exists", "passed": False, "detail": 'Document has no paragraphs'})
        passed = False
    else:
        first_para = doc.paragraphs[0]
        # Check style name includes 'Heading 1' (varies by Word version)
        style_name = first_para.style.name.lower()
        text = first_para.text.strip()
        heading_ok = "heading 1" in style_name and text == "Project Summary"
        checks.append({"name": "heading_1", "passed": heading_ok,
                       "detail": f'Heading style="{style_name}", text="{text}"'})
        if not heading_ok:
            passed = False

    # Check second paragraph has the exact marker text
    if len(doc.paragraphs) < 2:
        checks.append({"name": "paragraph_exists", "passed": False, "detail": 'Second paragraph missing'})
        passed = False
    else:
        para2 = doc.paragraphs[1]
        expected_text = "This document summarizes the key project details and milestones."
        para2_ok = para2.text.strip() == expected_text
        checks.append({"name": "paragraph_text", "passed": para2_ok, "detail": f'Paragraph text="{para2.text.strip()}"'})
        if not para2_ok:
            passed = False

    # Score: 1.0 if all passed, else 0
    score = 1.0 if passed else 0

    print(json.dumps({"passed": passed, "score": score, "checks": checks}))

if __name__ == "__main__":
    main()
