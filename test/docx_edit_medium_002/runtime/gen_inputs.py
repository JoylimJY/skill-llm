import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Create a Word doc with some paragraphs including marker text

def create_docx(filename):
    doc = docx.Document()

    # Add heading Introduction
    doc.add_heading('Introduction', level=1)

    # Add paragraph with 'The term is 30 days.' as marker
    para = doc.add_paragraph('The term is 30 days.')

    # Add another paragraph
    doc.add_paragraph('This is a sample paragraph for testing.')

    # Save
    doc.save(filename)

if __name__ == '__main__':
    create_docx('input.docx')
