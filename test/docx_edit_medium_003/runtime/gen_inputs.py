import docx
from docx.enum.style import WD_STYLE_TYPE

# Create a .docx with headings, paragraphs, and multiple "30 days" occurrences

def create_input_doc():
    doc = docx.Document()

    # Add Heading 1
    h1 = doc.add_heading('Introduction', level=1)

    # Paragraph with multiple "30 days" references
    p1 = doc.add_paragraph('The contract term is 30 days. After these 30 days, renewal is automatic.')

    # Heading 2
    h2 = doc.add_heading('Details', level=2)

    # Paragraph with "30 days" again
    p2 = doc.add_paragraph('Clients have 30 days to respond to notices.')

    # Add a normal paragraph
    doc.add_paragraph('This is additional report text for testing.')

    # Save
    doc.save('input.docx')

if __name__ == '__main__':
    create_input_doc()
