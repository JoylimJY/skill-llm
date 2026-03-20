import docx
import random
random.seed(42)

# Create contract.docx with marker content
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Title
h1 = doc.add_heading('Service Contract', level=1)

# Section: Terms and Conditions
h2 = doc.add_heading('Terms and Conditions', level=2)

# Paragraph intro
p = doc.add_paragraph('The following terms apply to this agreement.')

# Numbered list start - warranty and payment terms
lst1 = doc.add_paragraph(style='List Number')
lst1.add_run('The warranty period is 12 months.')

lst2 = doc.add_paragraph(style='List Number')
lst2.add_run('Payment must be made within 30 days of invoice.')

# Add a paragraph after list
doc.add_paragraph('Additional clauses may be added upon mutual agreement.')

# Save the input file

doc.save('contract.docx')

print('contract.docx generated')
