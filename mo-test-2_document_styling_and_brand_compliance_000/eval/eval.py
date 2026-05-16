import sys
import json
import traceback
from pathlib import Path

def check(name, passed, detail):
    return {"name": name, "passed": passed, "detail": detail}

def rgb_close(actual_rgb, expected_hex, tolerance=5):
    """Check if an RGB color matches expected hex within tolerance."""
    try:
        er = int(expected_hex[0:2], 16)
        eg = int(expected_hex[2:4], 16)
        eb = int(expected_hex[4:6], 16)
        return (abs(actual_rgb[0] - er) <= tolerance and
                abs(actual_rgb[1] - eg) <= tolerance and
                abs(actual_rgb[2] - eb) <= tolerance)
    except Exception:
        return False

def get_rgb_tuple(rgb_color):
    try:
        r = (rgb_color >> 16) & 0xFF
        g = (rgb_color >> 8) & 0xFF
        b = rgb_color & 0xFF
        return (r, g, b)
    except Exception:
        return None

def main():
    workspace = sys.argv[1]
    checks = []
    
    # Brand color definitions
    BRAND_DARK = "141413"
    BRAND_LIGHT = "faf9f5"
    BRAND_MID_GRAY = "b0aea5"
    BRAND_LIGHT_GRAY = "e8e6dc"
    BRAND_ORANGE = "d97757"
    BRAND_BLUE = "6a9bcc"
    BRAND_GREEN = "788c5d"
    
    BRAND_COLORS_ALL = [BRAND_DARK, BRAND_LIGHT, BRAND_MID_GRAY, BRAND_LIGHT_GRAY,
                        BRAND_ORANGE, BRAND_BLUE, BRAND_GREEN]
    ACCENT_CYCLE = [BRAND_ORANGE, BRAND_BLUE, BRAND_GREEN]
    
    HEADING_FONTS = {"poppins", "arial"}
    BODY_FONTS = {"lora", "georgia"}
    
    # Find the branded output file
    branded_file = None
    try:
        # Look for a new file or modified file - agent may create a new one or overwrite
        candidates = list(Path(workspace).rglob("*.pptx"))
        if not candidates:
            checks.append(check("output_file_exists", False, "No .pptx file found anywhere in workspace"))
            print(json.dumps({"passed": False, "score": 0.0, "checks": checks}))
            return
        
        # Prefer files with "brand" in name, else take any pptx
        brand_candidates = [f for f in candidates if "brand" in f.name.lower()]
        if brand_candidates:
            branded_file = brand_candidates[0]
        else:
            # Use the most recently modified pptx
            branded_file = max(candidates, key=lambda f: f.stat().st_mtime)
        
        checks.append(check("output_file_exists", True, f"Found branded file: {branded_file}"))
    except Exception as e:
        checks.append(check("output_file_exists", False, f"Error finding file: {e}"))
        print(json.dumps({"passed": False, "score": 0.0, "checks": checks}))
        return

    # Load the presentation
    try:
        from pptx import Presentation
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        prs = Presentation(str(branded_file))
        checks.append(check("file_parseable", True, "Presentation loaded successfully"))
    except Exception as e:
        checks.append(check("file_parseable", False, f"Could not parse pptx: {e}"))
        print(json.dumps({"passed": False, "score": 0.0, "checks": checks}))
        return

    slides = prs.slides
    
    # ---- Check 1: Slide count preserved ----
    slide_count_ok = len(slides) == 4
    checks.append(check("slide_count_preserved", slide_count_ok,
                        f"Expected 4 slides, got {len(slides)}"))

    # ---- Check 2: Heading fonts (24pt+) use Poppins or Arial ----
    heading_font_violations = []
    body_font_violations = []
    
    for slide_idx, slide in enumerate(slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        font_size = run.font.size
                        font_name = run.font.name
                        if font_size is None:
                            # Try paragraph font
                            font_size = para.font.size
                        if font_size is None:
                            continue
                        size_pt = font_size / 12700  # EMU to points
                        if font_name is None:
                            font_name = para.font.name
                        if font_name is None:
                            continue
                        fn_lower = font_name.lower().strip()
                        if size_pt >= 24:
                            # Should be Poppins or Arial
                            if not any(h in fn_lower for h in HEADING_FONTS):
                                heading_font_violations.append(
                                    f"Slide {slide_idx+1}: '{run.text[:30]}' at {size_pt}pt has font '{font_name}'"
                                )
                        else:
                            # Should be Lora or Georgia
                            if not any(b in fn_lower for b in BODY_FONTS):
                                body_font_violations.append(
                                    f"Slide {slide_idx+1}: '{run.text[:30]}' at {size_pt}pt has font '{font_name}'"
                                )
                    except Exception:
                        continue

    heading_ok = len(heading_font_violations) == 0
    checks.append(check("heading_fonts_correct",
                        heading_ok,
                        "All 24pt+ text uses Poppins/Arial" if heading_ok else
                        f"Violations: {'; '.join(heading_font_violations[:5])}"))

    body_ok = len(body_font_violations) == 0
    checks.append(check("body_fonts_correct",
                        body_ok,
                        "All sub-24pt text uses Lora/Georgia" if body_ok else
                        f"Violations: {'; '.join(body_font_violations[:5])}"))

    # ---- Check 3: Text colors use brand palette ----
    text_color_violations = []
    brand_hex_set = set(BRAND_COLORS_ALL)
    
    for slide_idx, slide in enumerate(slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if not run.text.strip():
                        continue
                    try:
                        color = run.font.color
                        if color and color.type is not None:
                            rgb_val = color.rgb
                            rgb_int = int(rgb_val)
                            r = (rgb_int >> 16) & 0xFF
                            g = (rgb_int >> 8) & 0xFF
                            b = rgb_int & 0xFF
                            is_brand = any(rgb_close((r, g, b), hx) for hx in BRAND_COLORS_ALL)
                            if not is_brand:
                                text_color_violations.append(
                                    f"Slide {slide_idx+1}: '{run.text[:20]}' color=#{r:02X}{g:02X}{b:02X}"
                                )
                    except Exception:
                        continue

    text_color_ok = len(text_color_violations) == 0
    checks.append(check("text_colors_on_brand",
                        text_color_ok,
                        "All text colors use brand palette" if text_color_ok else
                        f"Off-brand text colors: {'; '.join(text_color_violations[:5])}"))

    # ---- Check 4: Non-text shapes use accent colors (cycle: orange, blue, green) ----
    shape_accent_colors = []
    
    for slide_idx, slide in enumerate(slides):
        for shape in slide.shapes:
            if shape.has_text_frame:
                continue
            # It's a non-text shape
            try:
                fill = shape.fill
                if fill.type is not None:
                    try:
                        rgb_val = fill.fore_color.rgb
                        rgb_int = int(rgb_val)
                        r = (rgb_int >> 16) & 0xFF
                        g = (rgb_int >> 8) & 0xFF
                        b = rgb_int & 0xFF
                        shape_accent_colors.append((slide_idx + 1, shape.name, r, g, b))
                    except Exception:
                        pass
            except Exception:
                pass

    # Check that all non-text shape colors are from accent palette
    accent_violations = []
    for (slide_num, shape_name, r, g, b) in shape_accent_colors:
        is_accent = any(rgb_close((r, g, b), hx) for hx in ACCENT_CYCLE)
        if not is_accent:
            accent_violations.append(
                f"Slide {slide_num} shape '{shape_name}': #{r:02X}{g:02X}{b:02X} is not an accent color"
            )

    shapes_accent_ok = len(accent_violations) == 0 and len(shape_accent_colors) > 0
    checks.append(check("shapes_use_accent_colors",
                        shapes_accent_ok,
                        f"Found {len(shape_accent_colors)} non-text shapes, all using accent colors" if shapes_accent_ok else
                        (f"No non-text shapes found with fills" if len(shape_accent_colors) == 0 else
                         f"Accent violations: {'; '.join(accent_violations[:5])}")))

    # ---- Check 5: Accent color cycling order (orange → blue → green) ----
    accent_sequence_ok = False
    accent_sequence_detail = ""
    if len(shape_accent_colors) >= 2:
        seen_accents = []
        for (_, _, r, g, b) in shape_accent_colors:
            for i, hx in enumerate(ACCENT_CYCLE):
                if rgb_close((r, g, b), hx):
                    seen_accents.append(i)  # 0=orange, 1=blue, 2=green
                    break
        
        # Verify cycling pattern: each should be next in cycle from previous
        cycle_valid = True
        for i in range(1, len(seen_accents)):
            expected_next = (seen_accents[i-1] + 1) % 3
            if seen_accents[i] != expected_next:
                cycle_valid = False
                break
        
        accent_sequence_ok = cycle_valid and len(seen_accents) == len(shape_accent_colors)
        cycle_names = ["orange", "blue", "green"]
        seen_names = [cycle_names[i] for i in seen_accents]
        accent_sequence_detail = (
            f"Accent cycle correct: {seen_names}" if accent_sequence_ok else
            f"Incorrect cycle order: {seen_names} (expected orange→blue→green rotation)"
        )
    else:
        accent_sequence_detail = f"Too few non-text shapes to verify cycle ({len(shape_accent_colors)} found)"
        accent_sequence_ok = len(shape_accent_colors) >= 1  # partial credit if at least one

    checks.append(check("accent_color_cycle_order", accent_sequence_ok, accent_sequence_detail))

    # ---- Check 6: Slide backgrounds use brand colors ----
    background_violations = []
    for slide_idx, slide in enumerate(slides):
        try:
            bg = slide.background
            fill = bg.fill
            if fill.type is not None:
                try:
                    rgb_val = fill.fore_color.rgb
                    rgb_int = int(rgb_val)
                    r = (rgb_int >> 16) & 0xFF
                    g = (rgb_int >> 8) & 0xFF
                    b = rgb_int & 0xFF
                    is_brand = any(rgb_close((r, g, b), hx) for hx in BRAND_COLORS_ALL)
                    if not is_brand:
                        background_violations.append(
                            f"Slide {slide_idx+1}: background #{r:02X}{g:02X}{b:02X} not in brand palette"
                        )
                except Exception:
                    pass
        except Exception:
            pass

    bg_ok = len(background_violations) == 0
    checks.append(check("slide_backgrounds_on_brand",
                        bg_ok,
                        "All slide backgrounds use brand colors" if bg_ok else
                        f"Off-brand backgrounds: {'; '.join(background_violations)}"))

    # ---- Check 7: Specific brand dark color applied (#141413) ----
    dark_color_used = False
    light_color_used = False
    
    for slide_idx, slide in enumerate(slides):
        # Check backgrounds
        try:
            bg = slide.background.fill
            if bg.type is not None:
                rgb_val = bg.fore_color.rgb
                rgb_int = int(rgb_val)
                r, g, b = (rgb_int >> 16) & 0xFF, (rgb_int >> 8) & 0xFF, rgb_int & 0xFF
                if rgb_close((r, g, b), BRAND_DARK):
                    dark_color_used = True
                if rgb_close((r, g, b), BRAND_LIGHT):
                    light_color_used = True
        except Exception:
            pass
        # Check text colors
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        if run.font.color and run.font.color.type is not None:
                            rgb_val = run.font.color.rgb
                            rgb_int = int(rgb_val)
                            r, g, b = (rgb_int >> 16) & 0xFF, (rgb_int >> 8) & 0xFF, rgb_int & 0xFF
                            if rgb_close((r, g, b), BRAND_DARK):
                                dark_color_used = True
                            if rgb_close((r, g, b), BRAND_LIGHT):
                                light_color_used = True
                    except Exception:
                        pass

    checks.append(check("brand_dark_color_applied", dark_color_used,
                        f"Brand dark (#141413) found in presentation" if dark_color_used else
                        "Brand dark color #141413 not found anywhere in presentation"))
    checks.append(check("brand_light_color_applied", light_color_used,
                        f"Brand light (#faf9f5) found in presentation" if light_color_used else
                        "Brand light color #faf9f5 not found anywhere in presentation"))

    # ---- Compute overall score ----
    total = len(checks)
    passed_count = sum(1 for c in checks if c["passed"])
    score = passed_count / total if total > 0 else 0.0
    
    # Must pass critical checks to overall pass
    critical_checks = [
        "output_file_exists",
        "file_parseable", 
        "heading_fonts_correct",
        "body_fonts_correct",
        "shapes_use_accent_colors",
        "accent_color_cycle_order",
    ]
    critical_passed = all(
        c["passed"] for c in checks if c["name"] in critical_checks
    )
    
    overall_passed = critical_passed and score >= 0.75

    print(json.dumps({
        "passed": overall_passed,
        "score": round(score, 3),
        "checks": checks
    }, indent=2))

if __name__ == "__main__":
    main()