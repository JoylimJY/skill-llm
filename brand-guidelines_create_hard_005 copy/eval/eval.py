#!/usr/bin/env python3
import sys
import os
import json
from pptx import Presentation
from pptx.dml.color import RGBColor

def check_presentation(workspace_dir):
    checks = []
    
    # Find PowerPoint file
    pptx_files = [f for f in os.listdir(workspace_dir) if f.endswith('.pptx')]
    if not pptx_files:
        checks.append({"name": "file_exists", "passed": False, "detail": "No .pptx file found"})
        return {"passed": False, "score": 0.0, "checks": checks}
    
    # Check filename
    target_filename = 'anthropic_ai_safety.pptx'
    filename_found = any(target_filename.lower() in f.lower() for f in pptx_files)
    checks.append({"name": "correct_filename", "passed": filename_found, "detail": f"Found files: {pptx_files}, looking for: {target_filename}"})
    
    # Load presentation
    pptx_file = pptx_files[0]
    try:
        prs = Presentation(os.path.join(workspace_dir, pptx_file))
    except Exception as e:
        checks.append({"name": "file_readable", "passed": False, "detail": f"Cannot read presentation: {e}"})
        return {"passed": False, "score": len([c for c in checks if c['passed']]) / len(checks), "checks": checks}
    
    checks.append({"name": "file_readable", "passed": True, "detail": "Presentation file is readable"})
    
    # Check slide count (should have at least 3 slides)
    slide_count = len(prs.slides)
    checks.append({"name": "slide_count", "passed": slide_count >= 3, "detail": f"Found {slide_count} slides, expected at least 3"})
    
    # Check title slide content
    if slide_count > 0:
        title_slide = prs.slides[0]
        title_text = ""
        subtitle_text = ""
        
        for shape in title_slide.shapes:
            if hasattr(shape, 'text'):
                text = shape.text.lower()
                if 'ai safety research' in text and 'anthropic' in text:
                    title_text = shape.text
                elif 'constitutional ai' in text and 'alignment' in text:
                    subtitle_text = shape.text
        
        title_found = bool(title_text)
        subtitle_found = bool(subtitle_text)
        
        checks.append({"name": "title_slide_title", "passed": title_found, "detail": f"Title found: {title_text if title_text else 'None'}"})
        checks.append({"name": "title_slide_subtitle", "passed": subtitle_found, "detail": f"Subtitle found: {subtitle_text if subtitle_text else 'None'}"})
    else:
        checks.extend([
            {"name": "title_slide_title", "passed": False, "detail": "No slides found"},
            {"name": "title_slide_subtitle", "passed": False, "detail": "No slides found"}
        ])
    
    # Check research areas slide
    research_slide_found = False
    research_content_found = False
    
    if slide_count > 1:
        for slide in prs.slides[1:]:
            slide_text = ""
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    slide_text += shape.text.lower() + " "
            
            if 'research areas' in slide_text:
                research_slide_found = True
                if ('constitutional ai' in slide_text and 
                    'harmlessness training' in slide_text and 
                    'ai alignment' in slide_text):
                    research_content_found = True
                break
    
    checks.append({"name": "research_slide_exists", "passed": research_slide_found, "detail": "Research Areas slide found" if research_slide_found else "Research Areas slide not found"})
    checks.append({"name": "research_content", "passed": research_content_found, "detail": "Required research topics found" if research_content_found else "Missing required research topics"})
    
    # Check conclusion slide
    conclusion_slide_found = False
    conclusion_content_found = False
    
    if slide_count > 2:
        for slide in prs.slides:
            slide_text = ""
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    slide_text += shape.text.lower() + " "
            
            if 'future directions' in slide_text:
                conclusion_slide_found = True
                if 'safe and beneficial ai' in slide_text:
                    conclusion_content_found = True
                break
    
    checks.append({"name": "conclusion_slide_exists", "passed": conclusion_slide_found, "detail": "Future Directions slide found" if conclusion_slide_found else "Future Directions slide not found"})
    checks.append({"name": "conclusion_content", "passed": conclusion_content_found, "detail": "Required conclusion content found" if conclusion_content_found else "Missing required conclusion content"})
    
    # Check for brand colors usage
    brand_colors_used = False
    anthropic_colors = {
        (20, 20, 19),    # Dark #141413
        (250, 249, 245), # Light #faf9f5
        (176, 174, 165), # Mid Gray #b0aea5
        (232, 230, 220), # Light Gray #e8e6dc
        (217, 119, 87),  # Orange #d97757
        (106, 155, 204), # Blue #6a9bcc
        (120, 140, 93)   # Green #788c5d
    }
    
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, 'fill') and shape.fill.type == 1:  # Solid fill
                try:
                    color = shape.fill.fore_color.rgb
                    rgb_tuple = (color.red, color.green, color.blue)
                    if rgb_tuple in anthropic_colors:
                        brand_colors_used = True
                        break
                except:
                    continue
            if hasattr(shape, 'text_frame'):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        try:
                            color = run.font.color.rgb
                            if color:
                                rgb_tuple = (color.red, color.green, color.blue)
                                if rgb_tuple in anthropic_colors:
                                    brand_colors_used = True
                                    break
                        except:
                            continue
        if brand_colors_used:
            break
    
    checks.append({"name": "brand_colors", "passed": brand_colors_used, "detail": "Anthropic brand colors detected" if brand_colors_used else "No Anthropic brand colors detected"})
    
    # Calculate score
    score = sum(1 for c in checks if c['passed']) / len(checks)
    passed = score >= 0.8
    
    return {"passed": passed, "score": score, "checks": checks}

if __name__ == "__main__":
    workspace_dir = sys.argv[1]
    result = check_presentation(workspace_dir)
    print(json.dumps(result))